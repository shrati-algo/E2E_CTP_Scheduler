#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_machine_utilization.py -- machine utilization + components workbook for one run (2026-09-30).

    python make_machine_utilization.py <run_dir> [out.xlsx]

Sheets:
  Notes                 window, definitions, assumptions
  Machine Utilization   one row per machine: area, busy h, PM h, available h, utilization (formula), lots, items
  Daily Utilization     busy h per machine per plant day + utilization % per day (formulas)
  Machine Components    what each machine makes: item, type, lots, qty, UOM, busy h, share of machine time (formula)
"""
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RUN = sys.argv[1]
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(RUN, "MACHINE_UTILIZATION.xlsx")
HERE = os.path.dirname(os.path.abspath(__file__))

fs = pd.read_csv(os.path.join(RUN, "phase5_floor_schedule_updated.csv"), parse_dates=["start_time", "end_time"])
cu = pd.read_csv(os.path.join(RUN, "phase6_curing_revision_updated.csv"), parse_dates=["startTime", "endTime"])
pl = pd.read_csv(os.path.join(HERE, "inputs", "plan.csv"), parse_dates=["startTime", "endTime"])
W0 = (pl["startTime"].min() - pd.Timedelta(hours=7)).normalize() + pd.Timedelta(hours=7)
W1 = pl["endTime"].max()
DAYS = list(pd.date_range(W0, W1 - pd.Timedelta(seconds=1), freq="D"))

dep = fs["department"].astype(str).str.upper()
up = fs[~dep.isin(["BUILDING", "CURING"])].assign(area="Upstream")
bd = fs[dep == "BUILDING"].assign(area="Building (TBM)", process="GT BUILDING")
recs = pd.concat([
    up[["area", "department", "process", "machine", "item", "item_type", "start_time", "end_time", "produce_qty", "UOM", "lot_id"]],
    bd[["area", "department", "process", "machine", "item", "item_type", "start_time", "end_time", "produce_qty", "UOM", "lot_id"]],
    cu.assign(area="Curing (press)", department="CURING", process="CURING", machine=cu["pressNo"].astype(str),
              item=cu["skuCode"].astype(str), item_type="Cured tyre", start_time=cu["startTime"],
              end_time=cu["endTime"], produce_qty=cu["qty"], UOM="NOS", lot_id=cu["block_id"].astype(str))
    [["area", "department", "process", "machine", "item", "item_type", "start_time", "end_time", "produce_qty", "UOM", "lot_id"]],
], ignore_index=True)
recs["machine"] = recs["machine"].astype(str).str.strip()
recs["produce_qty"] = pd.to_numeric(recs["produce_qty"], errors="coerce").fillna(0)


def inwin(s, e, a, b):
    return ((e.clip(upper=b) - s.clip(lower=a)).dt.total_seconds() / 3600).clip(lower=0)


recs["h_win"] = inwin(recs["start_time"], recs["end_time"], W0, W1)
for d in DAYS:
    recs[f"h_{d:%d%b}"] = inwin(recs["start_time"], recs["end_time"], d, d + pd.Timedelta(days=1))

# PM downtime per machine inside the window
pm = defaultdict_h = {}
try:
    dt = pd.read_csv(os.path.join(HERE, "db_pipeline", "raw", "jkplanning_CTP__downtime_reason.csv"),
                     parse_dates=["Scheduled Start Time", "Scheduled End Time"])

    def mk(m):
        m = str(m).strip()
        if m.isdigit() and m.startswith("34"):
            return f"TBMPCR{int(m[2:])}Stage2"
        if m.isdigit() and m.startswith("38"):
            return f"TBM-{int(m[2:]):02d}"
        return m
    dt["m"] = dt["Machine ID"].map(mk)
    dt["h"] = inwin(dt["Scheduled Start Time"], dt["Scheduled End Time"], W0, W1)
    pm = dt.groupby("m")["h"].sum().to_dict()
except Exception:
    pm = {}

WIN_H = (W1 - W0).total_seconds() / 3600
mach = (recs.groupby(["area", "machine"])
        .agg(department=("department", lambda s: ", ".join(sorted(set(map(str, s))))),
             process=("process", lambda s: ", ".join(sorted(set(map(str, s))))),
             busy_h=("h_win", "sum"), lots=("lot_id", "nunique"), items=("item", "nunique"),
             item_types=("item_type", lambda s: ", ".join(sorted(set(map(str, s)))[:4])))
        .reset_index())
area_order = {"Upstream": 0, "Building (TBM)": 1, "Curing (press)": 2}
mach = mach.sort_values(["area", "busy_h"], key=lambda c: c.map(area_order) if c.name == "area" else -c).reset_index(drop=True)
mach["pm_h"] = mach["machine"].map(pm).fillna(0.0)

comp = (recs.groupby(["area", "machine", "item", "item_type", "UOM"])
        .agg(lots=("lot_id", "nunique"), qty=("produce_qty", "sum"), busy_h=("h_win", "sum"),
             first=("start_time", "min"), last=("end_time", "max")).reset_index())
comp = comp.sort_values(["area", "machine", "busy_h"], key=lambda c: c.map(area_order) if c.name == "area" else c,
                        ascending=[True, True, False]).reset_index(drop=True)
daily = recs.groupby(["area", "machine"])[[f"h_{d:%d%b}" for d in DAYS]].sum().reset_index()
daily = mach[["area", "machine"]].merge(daily, on=["area", "machine"], how="left").fillna(0)

# ---------------------------------------------------------------- workbook
F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F4E78")
H_FONT = Font(name=F, bold=True, color="FFFFFF")
B_FONT = Font(name=F)
BLUE = Font(name=F, color="0000FF")
BOLD = Font(name=F, bold=True)
thin = Side(style="thin", color="BFBFBF")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
AREA_FILL = {"Upstream": "EAF1FB", "Building (TBM)": "FFF2CC", "Curing (press)": "E2EFDA"}
wb = Workbook()


def header(ws, row, cols, widths):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        x = ws.cell(row=row, column=j, value=c)
        x.font, x.fill, x.border = H_FONT, H_FILL, BRD
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


# Notes
ws = wb.active
ws.title = "Notes"
notes = [
    ("CTP machine utilization and components", None),
    ("Run", os.path.basename(os.path.normpath(RUN))),
    ("Window start (plant day 07:00)", W0.to_pydatetime()),
    ("Window end (plan end)", W1.to_pydatetime()),
    ("Window hours", WIN_H),
    ("", None),
    ("Definitions", None),
    ("Busy h", "Production run time inside the window (from the scheduled floor schedule / press blocks). Changeover gaps are NOT counted as busy."),
    ("PM h", "Planned PM/MTC downtime inside the window (jkplanning_CTP.downtime_reason)."),
    ("Available h", "Window hours - PM hours (24 x 7, no breaks / breakdowns: see ASSUMPTIONS.md J18)."),
    ("Utilization", "Busy h / Available h (formula)."),
    ("Share of machine time", "Item busy h / machine busy h (formula)."),
    ("Daily sheet", "Busy h per plant day (07:00-07:00), utilization = busy h / 24."),
    ("Colour", "Blue number = input taken from the schedule; black = formula."),
]
for i, (a, b) in enumerate(notes, 1):
    ws.cell(row=i, column=1, value=a).font = Font(name=F, bold=(b is None and a != ""), size=13 if i == 1 else 10)
    c = ws.cell(row=i, column=2, value=b)
    c.font = BLUE if isinstance(b, (int, float)) else B_FONT
    if isinstance(b, pd.Timestamp) or hasattr(b, "year"):
        c.number_format = "dd-mmm-yyyy hh:mm"
ws["B5"].number_format = "0.0"
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110

# Machine Utilization
ws = wb.create_sheet("Machine Utilization")
cols = ["Area", "Machine", "Department", "Process", "Item types", "Lots / blocks", "Items made",
        "Busy h", "PM h", "Available h", "Utilization", "Idle h"]
header(ws, 1, cols, [16, 16, 24, 22, 38, 12, 11, 11, 9, 12, 12, 10])
r = 2
for _, m in mach.iterrows():
    vals = [m["area"], m["machine"], m["department"], m["process"], m["item_types"], int(m["lots"]),
            int(m["items"]), round(float(m["busy_h"]), 2), round(float(m["pm_h"]), 2)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = BLUE if j in (6, 7, 8, 9) else B_FONT
        c.border = BRD
        c.fill = PatternFill("solid", fgColor=AREA_FILL[m["area"]])
    ws.cell(row=r, column=10, value=f"=Notes!$B$5-I{r}")
    ws.cell(row=r, column=11, value=f"=IF(J{r}>0,H{r}/J{r},0)")
    ws.cell(row=r, column=12, value=f"=J{r}-H{r}")
    for j in (10, 11, 12):
        c = ws.cell(row=r, column=j)
        c.font, c.border = B_FONT, BRD
        c.fill = PatternFill("solid", fgColor=AREA_FILL[m["area"]])
    for j in (8, 9, 10, 12):
        ws.cell(row=r, column=j).number_format = "#,##0.0;-#,##0.0;-"
    ws.cell(row=r, column=11).number_format = "0.0%"
    r += 1
last = r - 1
r += 1
ws.cell(row=r, column=1, value="Summary by area").font = BOLD
r += 1
for j, h in enumerate(["Area", "Machines", "Busy h", "Available h", "Utilization", "Machines >= 85%"], 1):
    c = ws.cell(row=r, column=j, value=h)
    c.font, c.fill, c.border = H_FONT, H_FILL, BRD
for a in ["Upstream", "Building (TBM)", "Curing (press)"]:
    r += 1
    ws.cell(row=r, column=1, value=a).font = B_FONT
    ws.cell(row=r, column=2, value=f'=COUNTIF($A$2:$A${last},A{r})')
    ws.cell(row=r, column=3, value=f'=SUMIF($A$2:$A${last},A{r},$H$2:$H${last})')
    ws.cell(row=r, column=4, value=f'=SUMIF($A$2:$A${last},A{r},$J$2:$J${last})')
    ws.cell(row=r, column=5, value=f'=IF(D{r}>0,C{r}/D{r},0)')
    ws.cell(row=r, column=6, value=f'=COUNTIFS($A$2:$A${last},A{r},$K$2:$K${last},">=0.85")')
    for j in range(1, 7):
        c = ws.cell(row=r, column=j)
        c.border = BRD
        c.font = B_FONT
        c.fill = PatternFill("solid", fgColor=AREA_FILL[a])
    ws.cell(row=r, column=3).number_format = "#,##0.0"
    ws.cell(row=r, column=4).number_format = "#,##0.0"
    ws.cell(row=r, column=5).number_format = "0.0%"
ws.auto_filter.ref = f"A1:L{last}"

# Daily Utilization
ws = wb.create_sheet("Daily Utilization")
dcols = [f"{d:%d-%b}" for d in DAYS]
cols = ["Area", "Machine"] + [f"{d} busy h" for d in dcols] + [f"{d} util" for d in dcols]
header(ws, 1, cols, [16, 16] + [11] * (2 * len(DAYS)))
n = len(DAYS)
for i, (_, m) in enumerate(daily.iterrows(), 2):
    ws.cell(row=i, column=1, value=m["area"]).font = B_FONT
    ws.cell(row=i, column=2, value=m["machine"]).font = B_FONT
    for k, d in enumerate(DAYS):
        c = ws.cell(row=i, column=3 + k, value=round(float(m[f"h_{d:%d%b}"]), 2))
        c.font, c.number_format = BLUE, "0.0;-0.0;-"
        col = get_column_letter(3 + k)
        u = ws.cell(row=i, column=3 + n + k, value=f"={col}{i}/24")
        u.font, u.number_format = B_FONT, "0%;-0%;-"
    for j in range(1, 3 + 2 * n):
        ws.cell(row=i, column=j).border = BRD
        ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=AREA_FILL[m["area"]])
ws.auto_filter.ref = f"A1:{get_column_letter(2 + 2 * n)}{len(daily) + 1}"

# Machine Components
ws = wb.create_sheet("Machine Components")
cols = ["Area", "Machine", "Item / component", "Item type", "UOM", "Lots / blocks", "Quantity",
        "Busy h", "Share of machine time", "First start", "Last finish"]
header(ws, 1, cols, [16, 16, 26, 22, 7, 12, 14, 10, 14, 17, 17])
mrow = {m: i for i, m in enumerate(mach["machine"] + "|" + mach["area"], 2)}
for i, (_, c0) in enumerate(comp.iterrows(), 2):
    vals = [c0["area"], c0["machine"], c0["item"], c0["item_type"], c0["UOM"], int(c0["lots"]),
            round(float(c0["qty"]), 2), round(float(c0["busy_h"]), 2)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = BLUE if j in (6, 7, 8) else B_FONT
    mr = mrow[c0["machine"] + "|" + c0["area"]]
    ws.cell(row=i, column=9, value=f"=IF('Machine Utilization'!$H${mr}>0,H{i}/'Machine Utilization'!$H${mr},0)")
    ws.cell(row=i, column=10, value=pd.Timestamp(c0["first"]).to_pydatetime())
    ws.cell(row=i, column=11, value=pd.Timestamp(c0["last"]).to_pydatetime())
    ws.cell(row=i, column=7).number_format = "#,##0.0"
    ws.cell(row=i, column=8).number_format = "0.0;-0.0;-"
    ws.cell(row=i, column=9).number_format = "0.0%"
    ws.cell(row=i, column=10).number_format = "dd-mmm hh:mm"
    ws.cell(row=i, column=11).number_format = "dd-mmm hh:mm"
    for j in range(1, 12):
        x = ws.cell(row=i, column=j)
        x.border = BRD
        x.fill = PatternFill("solid", fgColor=AREA_FILL[c0["area"]])
        if j == 9:
            x.font = B_FONT
ws.auto_filter.ref = f"A1:K{len(comp) + 1}"

from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)   # Excel recalculates every formula on open
wb.save(OUT)
print(f"wrote {OUT}: {len(mach)} machines, {len(comp)} machine-item rows, {len(DAYS)} days "
      f"({W0:%d-%b %H:%M} -> {W1:%d-%b %H:%M}, {WIN_H:.0f} h)")
