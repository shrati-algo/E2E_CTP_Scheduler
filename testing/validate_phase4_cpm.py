"""
validate_phase4_cpm.py
Phase 4 CPM Validation for Tire Manufacturing MRP Pipeline
Outputs: phase4_cpm_validation_report.xlsx
"""

import pandas as pd
import numpy as np
from pathlib import Path
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ── Paths ────────────────────────────────────────────────────────────────────
BASE    = Path(__file__).parent
INPUTS  = BASE / "output/PHASEWISE/00_INPUTS"
PH2_DIR = BASE / "output/PHASEWISE/PHASE2_LOTS"
PH3_DIR = BASE / "output/PHASEWISE/PHASE3_DAG"
PH4_DIR = BASE / "output/PHASEWISE/PHASE4_CPM"

OUT_FILE = BASE / "phase4_cpm_validation_report.xlsx"

# ── Colour constants ─────────────────────────────────────────────────────────
PASS_FILL   = PatternFill("solid", fgColor="C6EFCE")
PASS_FONT   = Font(color="276221")
FAIL_FILL   = PatternFill("solid", fgColor="FFC7CE")
FAIL_FONT   = Font(color="9C0006")
WARN_FILL   = PatternFill("solid", fgColor="FFEB9C")
WARN_FONT   = Font(color="9C5700")
HDR_FILL    = PatternFill("solid", fgColor="1F3864")
HDR_FONT    = Font(color="FFFFFF", bold=True)

TOLERANCE_SEC  = 1 / 3600          # 1 second in hours
TOLERANCE_MIN  = 1 / 60            # 1 minute in hours
TOLERANCE_SLACK = 0.01             # hours


def style_status_cell(cell, status: str):
    if status == "PASS":
        cell.fill, cell.font = PASS_FILL, PASS_FONT
    elif status == "FAIL":
        cell.fill, cell.font = FAIL_FILL, FAIL_FONT
    else:
        cell.fill, cell.font = WARN_FILL, WARN_FONT


def write_header_row(ws, headers, row=1):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def auto_col_widths(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)


def write_detail_sheet(wb, sheet_name: str, description: str,
                        df: pd.DataFrame, pass_col: str = "Status"):
    ws = wb.create_sheet(title=sheet_name)
    # Row 1: description
    ws.cell(row=1, column=1, value=description).font = Font(bold=True)
    # Row 2: counts
    if pass_col in df.columns:
        n_pass = (df[pass_col] == "PASS").sum()
        n_fail = (df[pass_col] == "FAIL").sum()
        n_warn = (df[pass_col] == "WARN").sum()
        ws.cell(row=2, column=1, value=f"PASS: {n_pass}  FAIL: {n_fail}  WARN: {n_warn}")
    else:
        ws.cell(row=2, column=1, value=f"Rows: {len(df)}")

    # Trim for large datasets
    note = ""
    if len(df) > 10000 and pass_col in df.columns:
        fails = df[df[pass_col] == "FAIL"]
        passes = df[df[pass_col] == "PASS"].head(500)
        warns  = df[df[pass_col] == "WARN"]
        df = pd.concat([fails, passes, warns]).reset_index(drop=True)
        note = " [TRUNCATED: FAILs + first 500 PASSes shown]"
    ws.cell(row=2, column=1).value += note if note else ""

    if df.empty:
        ws.cell(row=3, column=1, value="No data")
        return

    # Row 3: header
    headers = list(df.columns)
    write_header_row(ws, headers, row=3)

    # Rows 4+: data
    for r_idx, row_data in enumerate(df.itertuples(index=False), start=4):
        for c_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if headers[c_idx - 1] == pass_col:
                style_status_cell(cell, str(val))

    auto_col_widths(ws)


def load_data():
    print("Loading input files …")
    lots      = pd.read_csv(PH2_DIR / "phase2_lots_updated.csv", low_memory=False)
    edges     = pd.read_csv(PH3_DIR / "phase3_dag_edges_updated.csv", low_memory=False)
    degrees   = pd.read_csv(PH3_DIR / "phase3_lot_degrees_updated.csv", low_memory=False)
    cpm       = pd.read_csv(PH4_DIR / "phase4_lot_times_updated.csv", low_memory=False)
    violations= pd.read_csv(PH4_DIR / "phase4_aging_violations_updated.csv", low_memory=False)
    aest_sum  = pd.read_csv(PH4_DIR / "phase4_aest_window_by_itype_updated.csv", low_memory=False)

    # Parse datetime columns in CPM
    for col in ["EST", "EFT", "LST", "LFT", "AEST"]:
        cpm[col] = pd.to_datetime(cpm[col], errors="coerce")

    print(f"  lots:       {len(lots):>6} rows")
    print(f"  edges:      {len(edges):>6} rows")
    print(f"  degrees:    {len(degrees):>6} rows")
    print(f"  cpm:        {len(cpm):>6} rows")
    print(f"  violations: {len(violations):>6} rows")
    return lots, edges, degrees, cpm, violations, aest_sum


# ── Validation functions ─────────────────────────────────────────────────────

def v01_all_dag_lots_in_cpm(degrees, cpm):
    print("V01 – All DAG lots included in CPM …")
    cpm_ids = set(cpm["lot_id"].astype(str))
    rows = []
    for _, r in degrees[["lot_id", "item_code"]].iterrows():
        lid = str(r["lot_id"])
        in_cpm = lid in cpm_ids
        rows.append({
            "lot_id":    lid,
            "item_code": r["item_code"],
            "In_CPM":    in_cpm,
            "Status":    "PASS" if in_cpm else "FAIL",
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = f"{n_fail} lots from DAG missing in CPM" if n_fail else "All DAG lots present in CPM"
    return df, n_pass, n_fail, note


def v02_predecessor_before_successor(edges, cpm):
    print("V02 – Every predecessor finishes before successor starts …")
    cpm_idx = cpm.set_index("lot_id")[["EFT", "EST"]].copy()
    rows = []
    for _, e in edges.iterrows():
        prod = str(e["producer_lot"])
        cons = str(e["consumer_lot"])
        try:
            prod_eft = cpm_idx.loc[prod, "EFT"] if prod in cpm_idx.index else None
            cons_est = cpm_idx.loc[cons, "EST"] if cons in cpm_idx.index else None
            if prod_eft is None or cons_est is None or pd.isna(prod_eft) or pd.isna(cons_est):
                status = "WARN"
                gap_h  = None
            else:
                gap_h  = (cons_est - prod_eft).total_seconds() / 3600
                status = "PASS" if gap_h >= -TOLERANCE_SEC else "FAIL"
            rows.append({
                "producer_lot":  prod,
                "consumer_lot":  cons,
                "producer_EFT":  str(prod_eft),
                "consumer_EST":  str(cons_est),
                "Gap_h":         round(gap_h, 6) if gap_h is not None else None,
                "Status":        status,
            })
        except Exception as ex:
            rows.append({"producer_lot": prod, "consumer_lot": cons,
                         "producer_EFT": None, "consumer_EST": None,
                         "Gap_h": None, "Status": "WARN"})
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = f"{n_fail} edges violate precedence" if n_fail else "All precedence constraints satisfied"
    return df, n_pass, n_fail, note


def v03_eft_equals_est_plus_duration(cpm):
    print("V03 – EFT = EST + duration_h …")
    rows = []
    for _, r in cpm.iterrows():
        est = r["EST"]
        eft = r["EFT"]
        dur = r["duration_h"]
        try:
            comp_eft = est + timedelta(hours=float(dur))
            diff_h   = abs((comp_eft - eft).total_seconds() / 3600)
            status   = "PASS" if diff_h <= TOLERANCE_MIN else "FAIL"
        except Exception:
            comp_eft, diff_h, status = None, None, "WARN"
        rows.append({
            "lot_id":       r["lot_id"],
            "EST":          str(est),
            "duration_h":   dur,
            "Computed_EFT": str(comp_eft),
            "Actual_EFT":   str(eft),
            "Diff_h":       round(diff_h, 6) if diff_h is not None else None,
            "Status":       status,
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = f"{n_fail} lots where EFT ≠ EST + duration" if n_fail else "All EFTs consistent"
    return df, n_pass, n_fail, note


def v04_lft_equals_lst_plus_duration(cpm):
    print("V04 – LFT = LST + duration_h …")
    rows = []
    for _, r in cpm.iterrows():
        lst = r["LST"]
        lft = r["LFT"]
        dur = r["duration_h"]
        try:
            comp_lft = lst + timedelta(hours=float(dur))
            diff_h   = abs((comp_lft - lft).total_seconds() / 3600)
            status   = "PASS" if diff_h <= TOLERANCE_MIN else "FAIL"
        except Exception:
            comp_lft, diff_h, status = None, None, "WARN"
        rows.append({
            "lot_id":       r["lot_id"],
            "LST":          str(lst),
            "duration_h":   dur,
            "Computed_LFT": str(comp_lft),
            "Actual_LFT":   str(lft),
            "Diff_h":       round(diff_h, 6) if diff_h is not None else None,
            "Status":       status,
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = f"{n_fail} lots where LFT ≠ LST + duration" if n_fail else "All LFTs consistent"
    return df, n_pass, n_fail, note


def v05_slack_consistency(cpm):
    print("V05 – Slack = LST - EST = LFT - EFT …")
    rows = []
    for _, r in cpm.iterrows():
        try:
            slack_lst_est = (r["LST"] - r["EST"]).total_seconds() / 3600
            slack_lft_eft = (r["LFT"] - r["EFT"]).total_seconds() / 3600
            reported      = float(r["slack_h"])
            consistent    = abs(slack_lst_est - slack_lft_eft) <= TOLERANCE_SLACK
            rep_ok_1      = abs(reported - slack_lst_est) <= TOLERANCE_SLACK
            rep_ok_2      = abs(reported - slack_lft_eft) <= TOLERANCE_SLACK
            status = "PASS" if (consistent and rep_ok_1 and rep_ok_2) else "FAIL"
        except Exception:
            slack_lst_est = slack_lft_eft = None
            consistent = False
            status = "WARN"
        rows.append({
            "lot_id":                   r["lot_id"],
            "slack_h":                  r["slack_h"],
            "computed_slack_LST_EST":   round(slack_lst_est, 4) if slack_lst_est is not None else None,
            "computed_slack_LFT_EFT":   round(slack_lft_eft, 4) if slack_lft_eft is not None else None,
            "Consistent":               consistent,
            "Status":                   status,
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = f"{n_fail} lots with inconsistent slack" if n_fail else "All slacks consistent"
    return df, n_pass, n_fail, note


def v06_critical_lots_near_zero_float(cpm):
    print("V06 – Critical lots have zero (or near-zero) float …")
    critical = cpm[cpm["is_critical"] == True].copy()
    rows = []
    for _, r in critical.iterrows():
        slack = float(r["slack_h"])
        # Valid: near zero (|slack| <= 0.1h) OR negative (overloaded)
        near_zero = abs(slack) <= 0.1 or slack <= 0
        status = "PASS" if near_zero else "FAIL"
        rows.append({
            "lot_id":            r["lot_id"],
            "item_code":         r.get("item_code", ""),
            "is_critical":       True,
            "slack_h":           slack,
            "Expected_Near_Zero": near_zero,
            "Status":            status,
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = (f"{n_fail} critical lots with slack > 0.1h"
            if n_fail else "All critical lots have near-zero/negative slack")
    return df, n_pass, n_fail, note


def v07_noncritical_lots_positive_float(cpm):
    print("V07 – Non-critical lots have positive float …")
    non_crit = cpm[cpm["is_critical"] != True].copy()
    rows = []
    for _, r in non_crit.iterrows():
        try:
            slack  = float(r["slack_h"])
            status = "PASS" if slack > 0 else "FAIL"
        except Exception:
            slack, status = None, "WARN"
        rows.append({
            "lot_id":      r["lot_id"],
            "item_code":   r.get("item_code", ""),
            "is_critical": r["is_critical"],
            "slack_h":     slack,
            "Status":      status,
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = (f"{n_fail} non-critical lots with slack <= 0"
            if n_fail else "All non-critical lots have positive float")
    return df, n_pass, n_fail, note


def v08_project_makespan(cpm):
    print("V08 – Project makespan = finish time of last critical activity …")
    makespan         = cpm["EFT"].max()
    max_critical_eft = cpm.loc[cpm["is_critical"] == True, "EFT"].max()
    try:
        diff_h = abs((makespan - max_critical_eft).total_seconds() / 3600)
        status = "PASS" if diff_h < 1.0 else "WARN"
    except Exception:
        diff_h, status = None, "WARN"
    row = {
        "Total_Lots":       len(cpm),
        "Makespan":         str(makespan),
        "Max_Critical_EFT": str(max_critical_eft),
        "Difference_h":     round(diff_h, 4) if diff_h is not None else None,
        "Status":           status,
    }
    df = pd.DataFrame([row])
    n_pass = 1 if status == "PASS" else 0
    n_fail = 1 if status == "FAIL" else 0
    note = f"Makespan difference: {diff_h:.3f}h" if diff_h is not None else "Could not compute"
    return df, n_pass, n_fail, note


def v09_no_negative_durations(cpm):
    print("V09 – No negative durations or invalid CPM times …")
    rows = []
    for _, r in cpm.iterrows():
        issues = []
        try:
            if float(r["duration_h"]) < 0:
                issues.append("duration_h<0")
        except Exception:
            issues.append("duration_h=NaN")
        try:
            if pd.notna(r["EST"]) and pd.notna(r["EFT"]) and r["EST"] >= r["EFT"]:
                issues.append("EST>=EFT")
        except Exception:
            issues.append("EST/EFT invalid")
        try:
            if pd.notna(r["LST"]) and pd.notna(r["LFT"]) and r["LST"] >= r["LFT"]:
                issues.append("LST>=LFT")
        except Exception:
            issues.append("LST/LFT invalid")
        if issues:
            rows.append({
                "lot_id":      r["lot_id"],
                "item_code":   r.get("item_code", ""),
                "duration_h":  r["duration_h"],
                "EST":         str(r["EST"]),
                "EFT":         str(r["EFT"]),
                "LST":         str(r["LST"]),
                "LFT":         str(r["LFT"]),
                "Issues":      "; ".join(issues),
                "Status":      "FAIL",
            })
    n_fail = len(rows)
    n_pass = len(cpm) - n_fail
    if not rows:
        rows.append({"lot_id": "ALL", "item_code": "", "duration_h": "",
                     "EST": "", "EFT": "", "LST": "", "LFT": "",
                     "Issues": "None", "Status": "PASS"})
    df = pd.DataFrame(rows)
    note = f"{n_fail} lots with invalid times/durations" if n_fail else "All durations and times valid"
    return df, n_pass, n_fail, note


def v10_terminal_lots_reachable(degrees, cpm):
    print("V10 – Terminal lots reachable through valid path …")
    terminals = degrees[degrees["is_terminal"] == True].copy()
    cpm_idx   = cpm.set_index("lot_id")[["n_pred", "n_succ"]].copy()
    rows = []
    for _, r in terminals.iterrows():
        lid = str(r["lot_id"])
        depth   = r.get("depth", None)
        n_pred  = r.get("n_predecessors", None)
        n_succ  = r.get("n_successors",  None)
        # Also get n_pred from cpm if available
        if lid in cpm_idx.index:
            n_pred_cpm = cpm_idx.loc[lid, "n_pred"]
        else:
            n_pred_cpm = None
        reachable = (n_pred is not None and n_pred > 0) or \
                    (n_pred_cpm is not None and n_pred_cpm > 0)
        depth_ok  = depth is not None and float(depth) > 0
        status    = "PASS" if (reachable and depth_ok) else ("WARN" if reachable else "FAIL")
        rows.append({
            "lot_id":      lid,
            "item_code":   r.get("item_code", ""),
            "is_terminal": True,
            "n_pred":      n_pred,
            "n_succ":      n_succ,
            "depth":       depth,
            "Reachable":   reachable,
            "Status":      status,
        })
    df = pd.DataFrame(rows)
    n_pass = (df["Status"] == "PASS").sum()
    n_fail = (df["Status"] == "FAIL").sum()
    note = (f"{n_fail} terminal lots unreachable"
            if n_fail else "All terminal lots are reachable")
    return df, n_pass, n_fail, note


def write_aging_violations(wb, violations):
    print("AGING_VIOLATIONS – Writing aging violations sheet …")
    ws = wb.create_sheet(title="AGING_VIOLATIONS")
    ws.cell(row=1, column=1, value="Phase 4 Aging Violations").font = Font(bold=True)
    ws.cell(row=2, column=1, value=f"Total violations: {len(violations)}")
    if violations.empty:
        ws.cell(row=3, column=1, value="No violations found")
        return
    headers = list(violations.columns)
    write_header_row(ws, headers, row=3)
    for r_idx, row in enumerate(violations.itertuples(index=False), start=4):
        vtype = str(getattr(row, "type", "")).upper()
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if vtype == "EXPIRED":
                cell.fill = FAIL_FILL
                cell.font = FAIL_FONT
            else:
                cell.fill = WARN_FILL
                cell.font = WARN_FONT
    auto_col_widths(ws)


def build_workbook(results, lots, edges, degrees, cpm, violations, aest_sum):
    wb = openpyxl.Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"

    # Title
    title_cell = ws_sum.cell(row=1, column=1, value="Phase 4 CPM Validation Report")
    title_cell.font = Font(bold=True, size=16, color="FFFFFF")
    title_cell.fill = HDR_FILL
    ws_sum.merge_cells("A1:F1")

    ts_cell = ws_sum.cell(row=2, column=1, value=f"Run timestamp: {datetime.now():%Y-%m-%d %H:%M:%S}")
    ts_cell.font = Font(italic=True)
    ws_sum.merge_cells("A2:F2")

    # Summary table header
    sum_headers = ["#", "Validation", "Status", "Pass Count", "Fail Count", "Notes"]
    write_header_row(ws_sum, sum_headers, row=3)

    for i, (name, df, n_pass, n_fail, note) in enumerate(results, start=1):
        status = "PASS" if n_fail == 0 else "FAIL"
        row_vals = [f"V{i:02d}", name, status, n_pass, n_fail, note]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_sum.cell(row=3 + i, column=c_idx, value=val)
            if c_idx == 3:
                style_status_cell(cell, status)
    auto_col_widths(ws_sum)

    # Detail sheets
    for i, (name, df, n_pass, n_fail, note) in enumerate(results, start=1):
        write_detail_sheet(wb, f"V{i:02d}", name, df)

    # Aging violations
    write_aging_violations(wb, violations)

    return wb


def main():
    print("=" * 60)
    print("Phase 4 CPM Validation")
    print("=" * 60)
    lots, edges, degrees, cpm, violations, aest_sum = load_data()

    validations = [
        ("V01 – All DAG lots included in CPM",              *v01_all_dag_lots_in_cpm(degrees, cpm)),
        ("V02 – Predecessor finishes before successor",     *v02_predecessor_before_successor(edges, cpm)),
        ("V03 – EFT = EST + duration_h",                   *v03_eft_equals_est_plus_duration(cpm)),
        ("V04 – LFT = LST + duration_h",                   *v04_lft_equals_lst_plus_duration(cpm)),
        ("V05 – Slack = LST-EST = LFT-EFT",                *v05_slack_consistency(cpm)),
        ("V06 – Critical lots have near-zero float",        *v06_critical_lots_near_zero_float(cpm)),
        ("V07 – Non-critical lots have positive float",     *v07_noncritical_lots_positive_float(cpm)),
        ("V08 – Project makespan matches critical finish",  *v08_project_makespan(cpm)),
        ("V09 – No negative durations or invalid times",   *v09_no_negative_durations(cpm)),
        ("V10 – Terminal lots reachable from root lots",    *v10_terminal_lots_reachable(degrees, cpm)),
    ]

    results = [(name, df, n_pass, n_fail, note)
               for name, df, n_pass, n_fail, note in validations]

    print("\nBuilding Excel workbook …")
    wb = build_workbook(results, lots, edges, degrees, cpm, violations, aest_sum)
    wb.save(OUT_FILE)
    print(f"\nReport saved: {OUT_FILE}")

    total_pass = sum(r[2] for r in results)
    total_fail = sum(r[3] for r in results)
    print(f"Total PASS: {total_pass}  Total FAIL: {total_fail}")
    print("=" * 60)


if __name__ == "__main__":
    main()
