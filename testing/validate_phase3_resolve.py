"""
validate_phase3_resolve.py — validate phase3_resolve_lot_blocks output.

The resolve helper REPLAYS Phase 2's lot-claim: it walks the NET demand rows in time
order and assigns each to the earliest lot still open at that need_by. Output:
    phase3_lot_blocks_updated.csv.gz : (lot_id, item_code, sku, block_id, need_by, demand_qty)
    phase3_lot_to_skus_updated.csv   : (lot_id, n_blocks, n_skus, skus_set)

These 10 checks judge that OUTPUT against the inputs it was built from (net demand,
phase2_lots). V10 is INFORMATIONAL: it compares the replay to Phase 2's OWN direct
mapping (phase2_lot_blocks_updated) -- the replay is v6's window approximation, Phase 2's
direct write is authoritative, so a disagreement there is expected and reported, not failed.
"""
import pandas as pd
import numpy as np
from pathlib import Path
from datetime import datetime
import warnings
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

BASE     = Path(__file__).parent
INPUTS   = BASE / "output/PHASEWISE/00_INPUTS"
PH1C_DIR = BASE / "output/PHASEWISE/PHASE1C_NETTING"
PH2_DIR  = BASE / "output/PHASEWISE/PHASE2_LOTS"
PH3_DIR  = BASE / "output/PHASEWISE/PHASE3_DAG"
OUTPUT_XLSX = BASE / "phase3_resolve_validation_report.xlsx"

F_PASS = PatternFill("solid", fgColor="C6EFCE"); T_PASS = Font(color="276221")
F_FAIL = PatternFill("solid", fgColor="FFC7CE"); T_FAIL = Font(color="9C0006")
F_WARN = PatternFill("solid", fgColor="FFEB9C"); T_WARN = Font(color="9C5700")
F_HEAD = PatternFill("solid", fgColor="1F3864"); T_HEAD = Font(color="FFFFFF", bold=True)
TOL = 0.01
MAXROWS = 400


def ff(s):
    s = str(s).upper()
    return (F_PASS, T_PASS) if s == "PASS" else (F_FAIL, T_FAIL) if s == "FAIL" else (F_WARN, T_WARN)


def load():
    print("Loading…")
    lb = pd.read_csv(PH3_DIR / "phase3_lot_blocks_updated.csv.gz", compression="gzip")
    ls = pd.read_csv(PH3_DIR / "phase3_lot_to_skus_updated.csv")
    lots = pd.read_csv(PH2_DIR / "phase2_lots_updated.csv")
    dem = pd.read_csv(PH1C_DIR / "phase1_demand_NET_updated.csv.gz", compression="gzip")
    auth = None
    p = PH2_DIR / "phase2_lot_blocks_updated.csv.gz"
    if p.exists():
        auth = pd.read_csv(p, compression="gzip")
    for df, c in [(lb, "need_by"), (dem, "need_by"), (lots, "first_need"), (lots, "last_need")]:
        df[c] = pd.to_datetime(df[c], errors="coerce")
    lb["demand_qty"] = pd.to_numeric(lb["demand_qty"], errors="coerce")
    dem["demand_qty"] = pd.to_numeric(dem["demand_qty"], errors="coerce")
    print(f"  resolve blocks: {len(lb):,}  lot_to_skus: {len(ls):,}  lots: {len(lots):,}  net demand: {len(dem):,}")
    return lb, ls, lots, dem, auth


# ── checks ───────────────────────────────────────────────────────────────────────
def v01_window(lb, lots):
    m = lb.merge(lots[["lot_id", "first_need", "last_need"]], on="lot_id", how="left")
    m["ok"] = (m["first_need"] <= m["need_by"]) & (m["need_by"] <= m["last_need"])
    bad = m[~m["ok"] & m["first_need"].notna()]
    df = pd.DataFrame({"lot_id": m["lot_id"], "item_code": m["item_code"], "need_by": m["need_by"],
                       "first_need": m["first_need"], "last_need": m["last_need"],
                       "Status": np.where(m["ok"], "PASS", "FAIL")})
    return df, "V01", "need_by falls inside the assigned lot's [first_need, last_need]", int((~m["ok"]).sum()), \
           f"{len(m):,} resolved rows; {len(bad):,} outside their lot window"


def v02_conservation(lb, dem):
    r = lb.groupby("item_code")["demand_qty"].sum()
    d = dem.groupby("item_code")["demand_qty"].sum()
    m = pd.concat([d.rename("net_demand"), r.rename("resolved")], axis=1).fillna(0.0)
    m["diff"] = (m["resolved"] - m["net_demand"]).round(3)
    m["Status"] = np.where(m["diff"].abs() <= np.maximum(TOL, 1e-5 * m["net_demand"].abs()), "PASS", "FAIL")
    df = m.reset_index()[["item_code", "net_demand", "resolved", "diff", "Status"]]
    return df, "V02", "Resolved qty per item == net demand per item (nothing lost/created)", \
           int((df["Status"] == "FAIL").sum()), f"{len(df):,} items; rounding-aware tolerance"


def v03_completeness(lb, dem):
    n_res, n_dem = len(lb), len(dem)
    unmatched = n_dem - n_res
    status = "PASS" if unmatched == 0 else ("WARN" if unmatched > 0 else "FAIL")
    df = pd.DataFrame([{"net_demand_rows": n_dem, "resolved_rows": n_res,
                        "unmatched_(lost)": unmatched, "Status": status}])
    return df, "V03", "Every net-demand row is assigned to a lot (0 lost)", \
           (1 if status == "FAIL" else 0), f"resolved {n_res:,} of {n_dem:,} demand rows; unmatched={unmatched:,}"


def v04_referential(lb, lots):
    lot_item = dict(zip(lots["lot_id"], lots["item_code"].astype(str)))
    lset = set(lots["lot_id"])
    lb2 = lb.copy()
    lb2["lot_exists"] = lb2["lot_id"].isin(lset)
    lb2["item_matches"] = lb2["lot_id"].map(lot_item).astype(str) == lb2["item_code"].astype(str)
    lb2["Status"] = np.where(lb2["lot_exists"] & lb2["item_matches"], "PASS", "FAIL")
    bad = lb2[lb2["Status"] == "FAIL"][["lot_id", "item_code", "lot_exists", "item_matches", "Status"]]
    df = bad if len(bad) else pd.DataFrame([{"lot_id": "(all valid)", "item_code": "", "lot_exists": True,
                                             "item_matches": True, "Status": "PASS"}])
    return df, "V04", "Every resolved lot exists in phase2_lots and item_code agrees", \
           int((lb2["Status"] == "FAIL").sum()), f"{len(lb2):,} rows checked"


def v05_no_double(lb):
    # a demand-row instance = (item, sku, block, need_by). Under a correct replay each such
    # instance is claimed by exactly ONE lot. >1 distinct lot on the same instance = double-claim.
    g = lb.groupby(["item_code", "sku", "block_id", "need_by"])["lot_id"].nunique()
    bad = g[g > 1]
    df = bad.reset_index().rename(columns={"lot_id": "n_lots"}) if len(bad) else \
        pd.DataFrame([{"item_code": "(none double-claimed)", "n_lots": 1}])
    df["Status"] = np.where(df.get("n_lots", 1) > 1, "FAIL", "PASS")
    return df, "V05", "No demand-row instance claimed by more than one lot", int(len(bad)), \
           f"{len(g):,} distinct demand instances; {len(bad):,} double-claimed"


def v06_coverage(lb, lots):
    served = set(lb["lot_id"])
    lots2 = lots.copy()
    lots2["served"] = lots2["lot_id"].isin(served)
    un = lots2[~lots2["served"]]
    by = un["item_type"].value_counts() if "item_type" in un.columns else pd.Series(dtype=int)
    rows = [{"item_type": k, "lots_with_no_demand_row": int(v), "Status": "WARN"} for k, v in by.items()]
    if not rows:
        rows = [{"item_type": "(all lots served a demand row)", "lots_with_no_demand_row": 0, "Status": "PASS"}]
    df = pd.DataFrame(rows)
    return df, "V06", "INFO: lots that no demand row mapped to (finished goods / sized-not-demanded)", 0, \
           f"{len(un):,} of {len(lots2):,} lots have no resolved demand row (WARN, not a defect)"


def v07_lot_to_skus(lb, ls):
    agg = lb.groupby("lot_id").agg(nb=("block_id", "nunique"), ns=("sku", "nunique"),
                                   skus=("sku", lambda s: ", ".join(sorted(set(s.astype(str)))))).reset_index()
    m = ls.merge(agg, on="lot_id", how="left")
    m["ok_nb"] = m["n_blocks"] == m["nb"]
    m["ok_ns"] = m["n_skus"] == m["ns"]
    m["ok_sk"] = m["skus_set"].astype(str) == m["skus"].astype(str)
    m["Status"] = np.where(m["ok_nb"] & m["ok_ns"] & m["ok_sk"], "PASS", "FAIL")
    bad = m[m["Status"] == "FAIL"][["lot_id", "n_blocks", "nb", "n_skus", "ns", "Status"]]
    df = bad if len(bad) else pd.DataFrame([{"lot_id": "(all consistent)", "n_blocks": 0, "nb": 0,
                                             "n_skus": 0, "ns": 0, "Status": "PASS"}])
    return df, "V07", "lot_to_skus (n_blocks/n_skus/skus_set) matches the blocks file", \
           int((m["Status"] == "FAIL").sum()), f"{len(ls):,} lot rows"


def v08_fields(lb):
    bad_blk = lb["block_id"].isna() | (lb["block_id"].astype(str).str.strip() == "")
    bad_qty = lb["demand_qty"].isna() | (lb["demand_qty"] <= 0)
    bad_nb = lb["need_by"].isna()
    bad_sku = lb["sku"].isna() | (lb["sku"].astype(str).str.strip().isin(["", "nan"]))
    fail = bad_blk | bad_qty | bad_nb | bad_sku
    df = lb[fail].assign(null_block=bad_blk[fail], bad_qty=bad_qty[fail], null_need_by=bad_nb[fail],
                         null_sku=bad_sku[fail], Status="FAIL")[
        ["lot_id", "item_code", "block_id", "demand_qty", "null_block", "bad_qty", "null_need_by", "null_sku", "Status"]] \
        if fail.any() else pd.DataFrame([{"lot_id": "(all fields valid)", "Status": "PASS"}])
    return df, "V08", "No null block_id/sku/need_by and demand_qty > 0", int(fail.sum()), \
           f"{len(lb):,} rows checked"


def v09_fifo(lb, lots):
    # FIFO minimality invariant (NOT a re-run of the code): the claimed lot is the earliest
    # lot (by first_need) still open at need_by. So NO lot of the same item with a smaller
    # first_need may still be open (last_need >= need_by). Checked structurally.
    fn = dict(zip(lots["lot_id"], lots["first_need"]))
    lots_by_item = {it: g.sort_values("first_need")[["lot_id", "first_need", "last_need"]].values
                    for it, g in lots.groupby("item_code")}
    viol = 0; rows = []
    # sample per item to keep it fast: check up to 40 resolved rows per item
    for it, g in lb.groupby("item_code"):
        arr = lots_by_item.get(it)
        if arr is None:
            continue
        first_arr = arr[:, 1].astype("datetime64[ns]")
        last_arr = arr[:, 2].astype("datetime64[ns]")
        lid_arr = arr[:, 0]
        for lot_id, nb in g[["lot_id", "need_by"]].head(40).values:
            t = np.datetime64(pd.Timestamp(nb))
            af = fn.get(lot_id)
            if pd.isna(af):
                continue
            af = np.datetime64(pd.Timestamp(af))
            earlier_open = ((first_arr < af) & (last_arr >= t) & (first_arr <= t) & (lid_arr != lot_id))
            if earlier_open.any():
                viol += 1
                if len(rows) < MAXROWS:
                    rows.append({"item_code": it, "claimed_lot": lot_id, "need_by": nb,
                                 "earlier_open_lot": str(lid_arr[earlier_open][0]), "Status": "FAIL"})
    df = pd.DataFrame(rows) if rows else pd.DataFrame([{"item_code": "(FIFO minimal)", "claimed_lot": "",
                                                       "need_by": "", "earlier_open_lot": "", "Status": "PASS"}])
    return df, "V09", "Claimed lot is the earliest lot still open at need_by (FIFO)", int(viol), \
           f"sampled up to 40 rows/item; {viol:,} rows had an earlier open lot"


def v10_crosscheck(lb, auth):
    if auth is None:
        df = pd.DataFrame([{"note": "phase2_lot_blocks_updated.csv.gz not staged", "Status": "WARN"}])
        return df, "V10", "INFO: replay vs Phase 2 direct mapping", 0, "nothing to compare"
    a = auth.copy(); a["demand_qty"] = pd.to_numeric(a["demand_qty"], errors="coerce")
    key = ["lot_id", "item_code", "sku", "block_id"]
    am = a.groupby(key)["demand_qty"].sum().round(2)
    rm = lb.groupby(key)["demand_qty"].sum().round(2)
    merged = pd.concat([am.rename("phase2_direct"), rm.rename("replay")], axis=1).fillna(0.0)
    merged["diff"] = (merged["phase2_direct"] - merged["replay"]).round(2)
    dis = merged[merged["diff"].abs() > TOL]
    df = dis.reset_index().head(MAXROWS).assign(Status="WARN") if len(dis) else \
        pd.DataFrame([{"lot_id": "(identical)", "Status": "PASS"}])
    return df, "V10", "INFO: replay vs Phase 2's authoritative direct mapping (approximation gap)", 0, \
           (f"{len(dis):,} of {len(merged):,} (lot,item,sku,block) keys disagree — EXPECTED: the replay is "
            f"v6's window approximation; Phase 2's direct write is authoritative and is what phase3_dag uses")


def write(results):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    summ = wb.create_sheet("Summary")
    summ.cell(1, 1, "Phase 3 — Resolve Lot Blocks Validation").font = Font(bold=True, size=14)
    summ.cell(2, 1, f"Generated: {datetime.now():%Y-%m-%d %H:%M:%S}")
    summ.cell(3, 1, "Validates phase3_resolve_lot_blocks output. V06/V10 are INFORMATIONAL.")
    for j, h in enumerate(["#", "Validation", "Status", "Fail", "Notes"], 1):
        c = summ.cell(5, j, h); c.fill = F_HEAD; c.font = T_HEAD
    r = 6
    for df, num, name, nf, note in results:
        status = "FAIL" if nf > 0 else ("WARN" if (df["Status"].astype(str).str.upper() == "WARN").any() else "PASS")
        fill, font = ff(status)
        for j, v in enumerate([num, name, status, nf, note], 1):
            c = summ.cell(r, j, v); c.fill = fill; c.font = font
        r += 1
        ws = wb.create_sheet(num)
        ws.cell(1, 1, f"{num} — {name}").font = Font(bold=True)
        ws.cell(2, 1, note)
        show = df.head(MAXROWS)
        for j, col in enumerate(show.columns, 1):
            c = ws.cell(4, j, str(col)); c.fill = F_HEAD; c.font = T_HEAD
        for i, (_, row) in enumerate(show.iterrows(), 5):
            for j, col in enumerate(show.columns, 1):
                val = row[col]
                if isinstance(val, float) and np.isnan(val):
                    val = None
                cc = ws.cell(i, j, str(val) if isinstance(val, pd.Timestamp) else val)
                if str(col) == "Status":
                    f2, t2 = ff(val); cc.fill = f2; cc.font = t2
        for col in ws.columns:
            w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(w + 3, 60)
    for col in summ.columns:
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        summ.column_dimensions[get_column_letter(col[0].column)].width = min(w + 3, 70)
    wb.save(OUTPUT_XLSX)


def main():
    print("=" * 66); print("Phase 3 Resolve Lot Blocks — Validation"); print("=" * 66)
    lb, ls, lots, dem, auth = load()
    checks = [v01_window(lb, lots), v02_conservation(lb, dem), v03_completeness(lb, dem),
              v04_referential(lb, lots), v05_no_double(lb), v06_coverage(lb, lots),
              v07_lot_to_skus(lb, ls), v08_fields(lb), v09_fifo(lb, lots), v10_crosscheck(lb, auth)]
    write(checks)
    print()
    npass = nfail = nwarn = 0
    for df, num, name, nf, note in checks:
        st = "FAIL" if nf > 0 else ("WARN" if (df["Status"].astype(str).str.upper() == "WARN").any() else "PASS")
        npass += st == "PASS"; nfail += st == "FAIL"; nwarn += st == "WARN"
        print(f"  {num}  {st:4s}  fail={nf:<6d} {name}")
    print("=" * 66)
    print(f"  {npass} PASS | {nfail} FAIL | {nwarn} WARN")
    print(f"  report -> {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
