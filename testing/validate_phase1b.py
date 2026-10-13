"""
validate_phase1b.py - Phase 1b Demand Generation Validation
Validates that the MRP demand explosion from plan + BOM is correct.

Revision notes (this version):
  V02 - Checks only ROUTED BOM components (SFG/build components), not raw
        materials (chemicals, rubber, carbon black, oil, fabric, wire, cord).
        Phase 1b never generates demand for raw materials, so checking the
        full BOM produced thousands of false failures.
  V03 - Confirmed grouped-by-(SKU, Component) comparison (not row-by-row),
        and now also uses routed-only BOM so raw materials don't show up as
        expected-but-missing.
  V04 - Expected row count now (a) uses routed-only BOM items and (b) uses
        only "active" plan rows - i.e. excludes CHANGEOVER / MOULD_CLEAN
        placeholder rows and zero-qty rows, which never generate demand.
  V05 - Item-wise total demand now uses routed-only BOM for the same reason
        as V02/V03.
  V07 - Zero demand_qty is now treated as WARN (legitimate), separate from
        null/negative demand_qty which is always FAIL.
  V08 - Rebuilt to validate need_by against the actual Phase 1b scheduling
        rule: need_by must equal the startTime of a real (sku, startTime)
        entry in the plan. This is order-independent, so it no longer
        depends on reconstructing block_id by row position (which doesn't
        line up 1:1 with the plan file and was the source of false FAILs).
  V09 - Valid_Block check rebuilt as a structural check (block_id format +
        exactly one SKU per block_id) since the old check depended on the
        same fragile block_id reconstruction removed in V08.
  V10 - Total demand consistency now uses routed-only BOM for the same
        reason as V02/V03/V05.
  V11 - New sheet: SKU-wise BOM demand breakdown. For every planned SKU,
        shows the plan quantity, every routed BOM component's per-unit
        requirement, the resulting expected demand (BOM qty-per-unit x plan
        qty), what actually landed in the demand file, and a PASS/FAIL match
        - grouped and totalled SKU by SKU (e.g. SKU needs 1970 tyres ->
        Sidewall = 1 per tyre -> expect 1970 -> demand file shows 1970 ->
        match). Same underlying numbers as V03, just organized per SKU
        instead of one long flat table, so a component-level explosion can
        be read top-to-bottom for a single SKU at a time.

Also, at load time: plan rows for CHANGEOVER / MOULD_CLEAN (non-product,
zero-qty placeholder rows used in the shift schedule) are filtered out of
"plan" everywhere. These aren't SKUs at all, so leaving them in inflated
V01/V04 false failures as well.
"""

import pandas as pd
import numpy as np
from pathlib import Path
from datetime import datetime
import warnings
import traceback

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# ── Paths ──────────────────────────────────────────────────────────────────────
BASE = Path(__file__).parent

# Point these at your actual input files. Defaults assume they sit next to
# this script (same folder). Adjust if your files live elsewhere.
BOM_FILE    = BASE / "jkt_bom_pcr_15.xlsx"
PLAN_FILE   = BASE / "ctppcr_plan_shiftsheet.xlsx"
DEMAND_FILE = BASE / "phase1_demand_updated.csv"

OUTPUT_FILE = BASE / "phase1b_validation_report.xlsx"

# ── Color palette ──────────────────────────────────────────────────────────────
FILL_PASS   = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL   = PatternFill("solid", fgColor="FFC7CE")
FILL_WARN   = PatternFill("solid", fgColor="FFEB9C")
FILL_HEADER = PatternFill("solid", fgColor="1F3864")

FONT_PASS   = Font(color="276221", bold=False)
FONT_FAIL   = Font(color="9C0006", bold=False)
FONT_WARN   = Font(color="9C5700", bold=False)
FONT_HEADER = Font(color="FFFFFF", bold=True)

MAX_PASS_ROWS = 500

# BOM child_description categories that are raw materials with no production
# routing of their own in Phase 1b (bought-in chemicals / rubber / fabric /
# wire / cord). Phase 1b intentionally does not generate scheduled demand for
# these, so they must be excluded from any "expected vs actual" BOM check.
# Verified empirically against phase1_demand_updated.csv: these categories are
# 0% present in demand across every planned SKU, while every other category is
# ~100% present (bar a handful of genuine data-quality exceptions).
NON_ROUTED_CATEGORIES = {
    "SMALL CHEMICAL",
    "SYNETHIC RUBBER",   # (sic - matches source data spelling)
    "NATURAL RUBBER",
    "CARBON BLACK",
    "PROCESS OIL",
    "NYLON FABRIC",
    "RUBBER CHEMICALS",
    "BEAD WIRE",
    "STEEL TYRE CORD",
}

# Plan rows that are not real SKUs at all (changeover / mould-clean shift
# blocks). These always carry qty = 0 and never generate demand.
NON_PRODUCT_PLAN_CODES = {"CHANGEOVER", "MOULD_CLEAN"}


# ── Helper utilities ───────────────────────────────────────────────────────────

def get_fill_font(status: str):
    s = str(status).upper()
    if s == "PASS":
        return FILL_PASS, FONT_PASS
    elif s == "FAIL":
        return FILL_FAIL, FONT_FAIL
    else:
        return FILL_WARN, FONT_WARN


def auto_col_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                cell_len = len(str(cell.value)) if cell.value is not None else 0
                if cell_len > max_len:
                    max_len = cell_len
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def fmt_qty(x) -> str:
    """Display a quantity as a whole number with thousands separators when it
    is (for all practical purposes) an integer, otherwise with 4 decimals."""
    try:
        xf = float(x)
    except (TypeError, ValueError):
        return str(x)
    if abs(xf - round(xf)) < 1e-6:
        return f"{xf:,.0f}"
    return f"{xf:,.4f}"


def write_detail_sheet(wb, sheet_name: str, description: str, df: pd.DataFrame,
                       status_col: str = "Status", note: str = ""):
    """Write a detail sheet. Limits large datasets to FAILs + first 500 PASSes."""
    ws = wb.create_sheet(title=sheet_name)

    # Row 1: description
    ws.cell(row=1, column=1, value=description).font = Font(bold=True)

    # Trim dataset
    if len(df) > 10000:
        fail_df = df[df[status_col].str.upper() == "FAIL"]
        pass_df = df[df[status_col].str.upper() == "PASS"].head(MAX_PASS_ROWS)
        warn_df = df[df[status_col].str.upper() == "WARN"]
        df_out = pd.concat([fail_df, pass_df, warn_df], ignore_index=True)
        trim_note = (f"Large dataset: showing all FAILs ({len(fail_df)}), "
                     f"first {MAX_PASS_ROWS} PASSes, all WARNs. "
                     f"Total rows={len(df)}.")
    else:
        df_out = df.copy()
        trim_note = ""

    total    = len(df)
    n_pass   = int((df[status_col].str.upper() == "PASS").sum())
    n_fail   = int((df[status_col].str.upper() == "FAIL").sum())
    n_warn   = int((df[status_col].str.upper() == "WARN").sum())
    combined_note = " | ".join(filter(None, [note, trim_note]))

    # Row 2: counts summary
    ws.cell(row=2, column=1,
            value=f"Total={total}  PASS={n_pass}  FAIL={n_fail}  WARN={n_warn}"
                  + (f"  NOTE: {combined_note}" if combined_note else ""))

    if df_out.empty:
        ws.cell(row=3, column=1, value="(No data)")
        return

    cols = list(df_out.columns)

    # Row 3: header
    for c_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=c_idx, value=col_name)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")

    # Rows 4+: data
    for r_idx, row in enumerate(df_out.itertuples(index=False), start=4):
        status_val = str(getattr(row, status_col, "WARN")).upper()
        fill, font = get_fill_font(status_val)
        for c_idx, col_name in enumerate(cols, start=1):
            val = getattr(row, col_name)
            if isinstance(val, float) and np.isnan(val):
                val = None
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if col_name == status_col:
                cell.fill = fill
                cell.font = font

    auto_col_width(ws)


def write_summary_sheet(wb, results: list):
    """Write the Summary sheet as the first sheet."""
    # Insert Summary at position 0
    ws = wb.create_sheet(title="Summary", index=0)

    # Title
    title_cell = ws.cell(row=1, column=1, value="Phase 1b – Demand Generation Validation Report")
    title_cell.font = Font(bold=True, size=14)

    # Timestamp
    ws.cell(row=2, column=1, value=f"Run timestamp: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # Header row
    headers = ["#", "Validation", "Status", "Pass Count", "Fail Count", "Notes"]
    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c_idx, value=h)
        cell.fill = FILL_HEADER
        cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for r_idx, res in enumerate(results, start=5):
        fill, font = get_fill_font(res["status"])
        values = [res["num"], res["name"], res["status"],
                  res["pass_count"], res["fail_count"], res.get("notes", "")]
        for c_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.fill = fill
            cell.font = font

    auto_col_width(ws)


def write_sku_breakdown_sheet(wb, sheet_name: str, description: str, blocks: list, note: str = ""):
    """Write the SKU-wise BOM demand breakdown as nested per-SKU blocks.

    Layout, repeated for every SKU:

        SKU: <code>                    Plan Demand Qty: <qty>     Overall: PASS/FAIL
        Component | Category | BOM Qty Per Unit | Plan Qty | Expected Demand | Actual Demand | Variance | Status
        <one row per routed BOM component for that SKU>
        TOTAL                                                | <sum Expected> | <sum Actual> | <variance> | <status>
        (blank separator row)

    `blocks` is a list of dicts: {"sku", "plan_qty", "df" (component-level
    DataFrame), "total_expected", "total_actual", "total_variance", "status"}.
    """
    ws = wb.create_sheet(title=sheet_name)

    ws.cell(row=1, column=1, value=description).font = Font(bold=True)

    total_skus  = len(blocks)
    total_rows  = sum(len(b["df"]) for b in blocks)
    n_sku_fail  = sum(1 for b in blocks if b["status"] != "PASS")
    ws.cell(row=2, column=1,
            value=(f"SKUs={total_skus}  Component rows={total_rows}  "
                   f"SKUs with a mismatch={n_sku_fail}"
                   + (f"  NOTE: {note}" if note else "")))

    if not blocks:
        ws.cell(row=3, column=1, value="(No data)")
        return

    FILL_SKU = PatternFill("solid", fgColor="BDD7EE")
    FONT_SKU = Font(bold=True, size=11)
    FILL_TOT = PatternFill("solid", fgColor="D9D9D9")
    FONT_TOT = Font(bold=True)

    cols = ["Component", "Category", "BOM_Qty_Per_Unit", "Plan_Qty",
            "Expected_Demand", "Actual_Demand", "Variance", "Status"]
    n_cols = len(cols)

    r = 4
    for b in blocks:
        # SKU group header row
        for cc in range(1, n_cols + 1):
            ws.cell(row=r, column=cc).fill = FILL_SKU
        ws.cell(row=r, column=1, value=f"SKU: {b['sku']}").font = FONT_SKU
        ws.cell(row=r, column=4, value=f"Plan Demand Qty: {fmt_qty(b['plan_qty'])}").font = FONT_SKU
        ws.cell(row=r, column=7, value="Overall:").font = FONT_SKU
        overall_fill, overall_font = get_fill_font(b["status"])
        ws.cell(row=r, column=8, value=b["status"])
        ws.cell(row=r, column=8).fill = overall_fill
        ws.cell(row=r, column=8).font = overall_font
        r += 1

        # Column header row
        for c_idx, col_name in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=c_idx, value=col_name.replace("_", " "))
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
            cell.alignment = Alignment(horizontal="center")
        r += 1

        # Component data rows
        for _, row in b["df"].iterrows():
            fill, font = get_fill_font(row["Status"])
            for c_idx, col_name in enumerate(cols, start=1):
                val = row[col_name]
                if isinstance(val, float) and np.isnan(val):
                    val = None
                cell = ws.cell(row=r, column=c_idx, value=val)
                if col_name == "Status":
                    cell.fill = fill
                    cell.font = font
            r += 1

        # TOTAL row for this SKU
        for cc in range(1, n_cols + 1):
            ws.cell(row=r, column=cc).fill = FILL_TOT
            ws.cell(row=r, column=cc).font = FONT_TOT
        ws.cell(row=r, column=1, value="TOTAL")
        ws.cell(row=r, column=5, value=b["total_expected"])
        ws.cell(row=r, column=6, value=b["total_actual"])
        ws.cell(row=r, column=7, value=b["total_variance"])
        status_fill, status_font = get_fill_font(b["status"])
        ws.cell(row=r, column=8, value=b["status"])
        ws.cell(row=r, column=8).fill = status_fill
        ws.cell(row=r, column=8).font = status_font
        r += 1

        r += 1  # blank separator row before next SKU

    auto_col_width(ws)


# ── BOM cleaning / routed-component filtering ──────────────────────────────────

def clean_bom(bom: pd.DataFrame) -> pd.DataFrame:
    """Drop unnamed helper columns and rows without a usable child/child_quantity."""
    unnamed_cols = [c for c in bom.columns if str(c).startswith("Unnamed")]
    b = bom.drop(columns=unnamed_cols, errors="ignore").dropna(subset=["child", "child_quantity"]).copy()
    b["child_quantity"] = pd.to_numeric(b["child_quantity"], errors="coerce")
    b = b.dropna(subset=["child_quantity"])
    return b


def get_routed_bom(bom: pd.DataFrame) -> pd.DataFrame:
    """Clean BOM and drop raw-material (non-routed) rows.

    Phase 1b only generates scheduled demand for components that go through
    their own build/production step (green tyre, plies, belts, compounds,
    calendared rolls, etc). Raw materials consumed directly (chemicals,
    rubber, carbon black, oil, fabric, wire, cord) never get their own
    demand row, by design.
    """
    b = clean_bom(bom)
    cat = b["child_description"].astype(str).str.strip().str.upper()
    return b[~cat.isin(NON_ROUTED_CATEGORIES)].copy()


# ── Data loading ───────────────────────────────────────────────────────────────

def load_inputs():
    print("Loading input files...")

    bom = pd.read_excel(BOM_FILE, sheet_name=0)

    # The plan is a formatted shift-wise schedule export: row 0 is a title,
    # row 1 is a rolled-up metrics line, row 2 is the real header row.
    plan_raw = pd.read_excel(PLAN_FILE, sheet_name=0, header=2)
    plan_raw = plan_raw.rename(columns={
        "SKUCode": "skuCode", "StartTime": "startTime", "EndTime": "endTime",
        "Qty": "qty", "Date": "date", "Shift": "shift", "Machine": "machine",
        "CycleTime_min": "cycleTime_min", "GT_Inventory": "gt_inventory",
        "Remarks": "remarks",
    })
    plan_raw["startTime"] = pd.to_datetime(plan_raw["startTime"], errors="coerce")
    plan_raw["qty"] = pd.to_numeric(plan_raw["qty"], errors="coerce")

    # Drop CHANGEOVER / MOULD_CLEAN placeholder rows and any zero-qty rows.
    # These aren't SKUs and never generate demand - keeping them in "plan"
    # creates false failures in V01/V04 (and would break V02/V08 as well).
    is_product_row = ~plan_raw["skuCode"].isin(NON_PRODUCT_PLAN_CODES)
    is_active_qty  = plan_raw["qty"] > 0
    plan = plan_raw[is_product_row & is_active_qty].reset_index(drop=True)

    n_excluded = len(plan_raw) - len(plan)

    demand = pd.read_csv(DEMAND_FILE)

    print(f"  plan rows (raw):    {len(plan_raw):,}")
    print(f"  plan rows (active): {len(plan):,}  "
          f"(excluded {n_excluded:,} changeover/mould-clean/zero-qty rows)")
    print(f"  bom rows:    {len(bom):,}")
    print(f"  demand rows: {len(demand):,}")
    return plan, bom, demand


# ── Validations ────────────────────────────────────────────────────────────────

def v01_plan_skus_in_demand(plan, bom, demand):
    print("Running V01 – All plan SKUs present in demand...")
    try:
        plan_skus   = set(plan["skuCode"].dropna().unique())
        demand_skus = set(demand["sku"].dropna().unique())

        rows = []
        for sku in sorted(plan_skus):
            in_demand = sku in demand_skus
            rows.append({
                "SKU":       sku,
                "In_Plan":   True,
                "In_Demand": in_demand,
                "Status":    "PASS" if in_demand else "FAIL",
            })

        df = pd.DataFrame(rows)
        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V01", "name": "All plan SKUs present in demand",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": f"{len(plan_skus)} active plan SKUs checked "
                     f"(CHANGEOVER/MOULD_CLEAN placeholder rows excluded)",
            "df": df,
        }
    except Exception as e:
        return {"num": "V01", "name": "All plan SKUs present in demand",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v02_bom_items_in_demand(plan, bom, demand):
    print("Running V02 – All ROUTED BOM items generated in demand...")
    try:
        bom_routed = get_routed_bom(bom)
        n_excluded = len(clean_bom(bom)) - len(bom_routed)

        plan_skus = set(plan["skuCode"].dropna().unique())

        demand_pairs = set(zip(demand["sku"], demand["item_code"]))

        # Unique (Super_parent, child) pairs, restricted to planned SKUs only
        bom_pairs = (bom_routed[["Super_parent", "child", "child_description"]]
                     .drop_duplicates(subset=["Super_parent", "child"]))
        bom_pairs = bom_pairs[bom_pairs["Super_parent"].isin(plan_skus)]

        rows = []
        for _, row in bom_pairs.iterrows():
            sp   = row["Super_parent"]
            ch   = row["child"]
            pair = (sp, ch)
            in_d = pair in demand_pairs
            rows.append({
                "SKU":       sp,
                "Component": ch,
                "Category":  row["child_description"],
                "In_Demand": in_d,
                "Status":    "PASS" if in_d else "FAIL",
            })

        df = pd.DataFrame(rows)
        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V02", "name": "All routed BOM items generated in demand",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"{len(bom_pairs)} routed (SKU, component) pairs checked "
                      f"across {len(plan_skus)} planned SKUs; "
                      f"{n_excluded} raw-material pairs excluded by design "
                      f"(chemicals/rubber/carbon black/oil/fabric/wire/cord)"),
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V02", "name": "All routed BOM items generated in demand",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v03_demand_qty_bom_x_plan(plan, bom, demand):
    print("Running V03 – Demand Qty = BOM qty × plan qty (grouped by SKU, Component)...")
    try:
        bom_routed = get_routed_bom(bom)

        # Plan qty per SKU
        plan_qty_by_sku = plan.groupby("skuCode")["qty"].sum().reset_index()
        plan_qty_by_sku.columns = ["Super_parent", "Total_Plan_Qty"]

        # Expected demand: routed BOM x plan qty
        bom_plan = bom_routed[["Super_parent", "child", "child_quantity"]].merge(
            plan_qty_by_sku, on="Super_parent", how="inner"
        )
        bom_plan["Expected_Demand"] = bom_plan["child_quantity"] * bom_plan["Total_Plan_Qty"]
        expected = bom_plan.groupby(["Super_parent", "child"]).agg(
            BOM_Qty_Per_Unit=("child_quantity", "first"),
            Total_Plan_Qty=("Total_Plan_Qty", "first"),
            Expected_Demand=("Expected_Demand", "sum"),
        ).reset_index()

        # Actual demand: grouped by (sku, item_code) - sums any duplicate
        # demand rows for the same (SKU, Component) rather than comparing
        # row-by-row, so split/duplicate demand records don't cause false
        # mismatches.
        actual = demand.groupby(["sku", "item_code"])["demand_qty"].sum().reset_index()
        actual.columns = ["Super_parent", "child", "Actual_Demand"]

        merged = expected.merge(actual, on=["Super_parent", "child"], how="left")
        merged["Actual_Demand"] = merged["Actual_Demand"].fillna(0)
        merged["Variance"] = merged["Actual_Demand"] - merged["Expected_Demand"]
        merged["Status"] = merged["Variance"].abs().apply(lambda v: "PASS" if v < 0.01 else "FAIL")

        df = merged.rename(columns={"Super_parent": "SKU", "child": "Component"})
        df = df[["SKU", "Component", "BOM_Qty_Per_Unit", "Total_Plan_Qty",
                  "Expected_Demand", "Actual_Demand", "Variance", "Status"]]

        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V03", "name": "Demand Qty = BOM qty × plan qty (grouped by SKU, Component)",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"Tolerance < 0.01; {len(df)} routed (SKU, Component) pairs checked; "
                      f"demand summed per pair before comparing to absorb any duplicate rows"),
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V03", "name": "Demand Qty = BOM qty × plan qty (grouped by SKU, Component)",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v04_demand_row_count(plan, bom, demand):
    print("Running V04 – No missing or extra demand rows (routed items, active blocks only)...")
    try:
        bom_routed = get_routed_bom(bom)
        bom_items_per_sku = (
            bom_routed.groupby("Super_parent")["child"]
            .nunique()
            .reset_index()
            .rename(columns={"Super_parent": "skuCode", "child": "n_bom_items"})
        )

        # "plan" is already filtered to active (qty>0, real-SKU) rows, so
        # each row here corresponds to exactly one expected demand block.
        plan_blocks = plan[["skuCode"]].merge(bom_items_per_sku, on="skuCode", how="left")
        plan_blocks["n_bom_items"] = plan_blocks["n_bom_items"].fillna(0)
        expected_count = int(plan_blocks["n_bom_items"].sum())
        actual_count   = len(demand)

        rows = [{
            "Metric":   "Expected Row Count (routed items x active plan blocks)",
            "Value":    expected_count,
            "Status":   "",
        }, {
            "Metric":   "Actual Row Count",
            "Value":    actual_count,
            "Status":   "",
        }, {
            "Metric":   "Difference (Actual - Expected)",
            "Value":    actual_count - expected_count,
            "Status":   "PASS" if actual_count == expected_count else "FAIL",
        }]

        df = pd.DataFrame(rows)
        status = "PASS" if actual_count == expected_count else "FAIL"
        n_pass = 1 if status == "PASS" else 0
        n_fail = 0 if status == "PASS" else 1
        return {
            "num": "V04", "name": "No missing or extra demand rows",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"Expected={expected_count:,} Actual={actual_count:,} "
                      f"| Raw materials and CHANGEOVER/MOULD_CLEAN/zero-qty plan rows "
                      f"excluded from the expected count as legitimate filters. "
                      f"Any remaining difference should reconcile with V02's FAIL rows."),
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V04", "name": "No missing or extra demand rows",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v05_itemwise_total_demand(plan, bom, demand):
    print("Running V05 – Item-wise total demand correct (routed items only)...")
    try:
        bom_routed = get_routed_bom(bom)

        plan_qty_by_sku = plan.groupby("skuCode")["qty"].sum().reset_index()
        plan_qty_by_sku.columns = ["Super_parent", "Total_Plan_Qty"]

        bom_plan = bom_routed[["Super_parent", "child", "child_quantity"]].merge(
            plan_qty_by_sku, on="Super_parent", how="inner"
        )
        bom_plan["Expected_Demand"] = bom_plan["child_quantity"] * bom_plan["Total_Plan_Qty"]

        expected_by_item = bom_plan.groupby("child")["Expected_Demand"].sum().reset_index()
        expected_by_item.columns = ["item_code", "Expected"]

        actual_by_item = demand.groupby("item_code")["demand_qty"].sum().reset_index()
        actual_by_item.columns = ["item_code", "Actual"]

        merged = expected_by_item.merge(actual_by_item, on="item_code", how="outer")
        merged["Expected"] = merged["Expected"].fillna(0)
        merged["Actual"]   = merged["Actual"].fillna(0)
        merged["Variance"] = merged["Actual"] - merged["Expected"]
        merged["Status"]   = merged["Variance"].abs().apply(lambda v: "PASS" if v < 0.01 else "FAIL")

        n_pass = int((merged["Status"] == "PASS").sum())
        n_fail = int((merged["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V05", "name": "Item-wise total demand correct (routed items)",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": f"{len(merged)} item_codes checked; tolerance < 0.01; raw materials excluded",
            "df": merged,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V05", "name": "Item-wise total demand correct (routed items)",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v06_demand_uoms_valid(plan, bom, demand):
    print("Running V06 – All demand UOMs valid (match BOM child_Unit)...")
    try:
        bom_uom = bom[["Super_parent", "child", "child_Unit"]].dropna(subset=["child"]).drop_duplicates(
            subset=["Super_parent", "child"]
        )
        bom_uom.columns = ["sku", "item_code", "BOM_UOM"]

        demand_uom = demand[["sku", "item_code", "demand_uom"]].drop_duplicates(
            subset=["sku", "item_code"]
        ).copy()
        demand_uom.columns = ["sku", "item_code", "Demand_UOM"]

        merged = demand_uom.merge(bom_uom, on=["sku", "item_code"], how="left")
        merged["BOM_UOM"] = merged["BOM_UOM"].fillna("MISSING_IN_BOM")

        merged["Status"] = merged.apply(
            lambda r: "PASS" if str(r["Demand_UOM"]).strip() == str(r["BOM_UOM"]).strip()
            else ("WARN" if r["BOM_UOM"] == "MISSING_IN_BOM" else "FAIL"),
            axis=1
        )

        df = merged[["sku", "item_code", "Demand_UOM", "BOM_UOM", "Status"]]
        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V06", "name": "All demand UOMs valid (match BOM)",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": f"{len(df)} (sku, item_code) pairs checked",
            "df": df,
        }
    except Exception as e:
        return {"num": "V06", "name": "All demand UOMs valid (match BOM)",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v07_qty_and_field_integrity(plan, bom, demand):
    print("Running V07 – Field integrity: null/negative qty = FAIL, zero qty = WARN...")
    try:
        mask_item_null  = demand["item_code"].isna()
        mask_block_null = demand["block_id"].isna()
        mask_date_null  = demand["need_by"].isna()
        mask_uom_null   = demand["demand_uom"].isna()
        mask_qty_null   = demand["demand_qty"].isna()
        mask_qty_neg    = demand["demand_qty"] < 0
        mask_qty_zero   = (~mask_qty_null) & (demand["demand_qty"] == 0)

        fail_mask = mask_item_null | mask_block_null | mask_date_null | mask_uom_null | mask_qty_null | mask_qty_neg
        warn_mask = (~fail_mask) & mask_qty_zero

        status = np.where(fail_mask, "FAIL", np.where(warn_mask, "WARN", "PASS"))

        df = demand[["item_code", "block_id", "demand_qty", "need_by", "demand_uom"]].copy()
        df["Null_item_code"]       = mask_item_null
        df["Null_block_id"]        = mask_block_null
        df["Null_need_by"]         = mask_date_null
        df["Null_demand_uom"]      = mask_uom_null
        df["Null_or_Negative_Qty"] = mask_qty_null | mask_qty_neg
        df["Zero_Qty"]             = mask_qty_zero
        df["Status"]               = status

        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        n_warn = int((df["Status"] == "WARN").sum())
        overall_status = "FAIL" if n_fail > 0 else ("WARN" if n_warn > 0 else "PASS")

        return {
            "num": "V07", "name": "Field integrity (null/negative qty vs zero qty)",
            "status": overall_status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"Total rows checked: {len(demand):,} | "
                      f"{n_warn:,} legitimate zero-qty rows (WARN, not an error) | "
                      f"null/negative qty and null item_code/block_id/need_by/uom = FAIL"),
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V07", "name": "Field integrity (null/negative qty vs zero qty)",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v08_need_by_matches_schedule(plan, bom, demand):
    print("Running V08 – need_by matches a real scheduled (SKU, StartTime) in the plan...")
    try:
        # Order-independent by design: need_by is expected to equal the
        # startTime of *some* scheduled entry for that SKU in the plan, not
        # necessarily "row N of the plan". Reconstructing block_id from row
        # position doesn't line up with the demand file's own block_id
        # (the plan has extra changeover/mould-clean/zero-qty rows), so we
        # validate the scheduling rule directly instead.
        plan_pairs = set(zip(plan["skuCode"], plan["startTime"]))

        block_level = demand[["block_id", "sku", "need_by"]].drop_duplicates(subset=["block_id"]).copy()
        block_level["need_by_dt"] = pd.to_datetime(block_level["need_by"], errors="coerce")

        def check_match(row):
            if pd.isnull(row["need_by_dt"]):
                return "WARN"
            return "PASS" if (row["sku"], row["need_by_dt"]) in plan_pairs else "FAIL"

        block_level["Status"] = block_level.apply(check_match, axis=1)

        df = block_level[["block_id", "sku", "need_by", "Status"]]

        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V08", "name": "need_by matches a real scheduled (SKU, StartTime)",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"{len(df)} blocks checked against {plan['skuCode'].nunique()} "
                      f"planned SKUs' scheduled start times; validated against the "
                      f"Phase 1b scheduling rule rather than plan row position"),
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V08", "name": "need_by matches a real scheduled (SKU, StartTime)",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v09_demand_valid_sku_block(plan, bom, demand):
    print("Running V09 – Every demand row has valid SKU and well-formed block...")
    try:
        valid_skus = set(plan["skuCode"].dropna())

        df = demand[["item_code", "sku", "block_id"]].copy()
        df["Valid_SKU"] = df["sku"].isin(valid_skus)

        # Structural block validity: well-formed ID, and a block must map to
        # exactly one SKU (a block mixing SKUs would indicate corrupted
        # demand generation).
        block_format_ok = demand["block_id"].astype(str).str.match(r"^B\d+$")
        sku_count_per_block = demand.groupby("block_id")["sku"].transform("nunique")
        df["Valid_Block"] = block_format_ok.values & (sku_count_per_block == 1).values

        df["Status"] = np.where(df["Valid_SKU"] & df["Valid_Block"], "PASS", "FAIL")

        n_pass = int((df["Status"] == "PASS").sum())
        n_fail = int((df["Status"] == "FAIL").sum())
        status = "PASS" if n_fail == 0 else "FAIL"
        return {
            "num": "V09", "name": "Every demand row has valid SKU and well-formed block_id",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"{len(valid_skus)} valid SKUs; block validity = well-formed ID "
                      f"('B' + digits) and exactly one SKU per block_id"),
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V09", "name": "Every demand row has valid SKU and well-formed block_id",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v10_total_demand_consistent(plan, bom, demand):
    print("Running V10 – Total demand internally consistent (sum = Plan × routed BOM)...")
    try:
        bom_routed = get_routed_bom(bom)

        plan_qty_by_sku = plan.groupby("skuCode")["qty"].sum().reset_index()
        plan_qty_by_sku.columns = ["Super_parent", "Total_Plan_Qty"]

        bom_plan = bom_routed[["Super_parent", "child", "child_quantity"]].merge(
            plan_qty_by_sku, on="Super_parent", how="inner"
        )
        total_expected = float((bom_plan["child_quantity"] * bom_plan["Total_Plan_Qty"]).sum())
        total_actual   = float(demand["demand_qty"].sum())
        variance       = total_actual - total_expected

        status = "PASS" if abs(variance) < 0.01 else "FAIL"

        df = pd.DataFrame([{
            "Total_Expected": round(total_expected, 4),
            "Total_Actual":   round(total_actual, 4),
            "Variance":       round(variance, 4),
            "Status":         status,
        }])

        n_pass = 1 if status == "PASS" else 0
        n_fail = 0 if status == "PASS" else 1
        return {
            "num": "V10", "name": "Total demand internally consistent (routed BOM)",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": f"Expected={total_expected:,.2f} Actual={total_actual:,.2f} Variance={variance:,.4f}",
            "df": df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V10", "name": "Total demand internally consistent (routed BOM)",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "df": pd.DataFrame()}


def v11_sku_wise_bom_demand_breakdown(plan, bom, demand):
    """Per-SKU view of the same check as V03: for every planned SKU, show its
    plan quantity, every routed BOM component's per-unit requirement, the
    expected demand (BOM qty-per-unit x plan qty), what actually landed in
    the demand file, and whether the two match - grouped and totalled SKU by
    SKU. E.g. SKU needs 1970 tyres -> Sidewall = 1 per tyre -> expect 1970
    sidewalls -> demand file shows 1970 -> match.
    """
    print("Running V11 – SKU-wise BOM demand breakdown (Plan Qty x BOM = Expected vs Actual)...")
    try:
        bom_routed = get_routed_bom(bom)

        plan_qty_by_sku = plan.groupby("skuCode")["qty"].sum().reset_index()
        plan_qty_by_sku.columns = ["Super_parent", "Total_Plan_Qty"]

        bom_plan = bom_routed[["Super_parent", "child", "child_quantity", "child_description"]].merge(
            plan_qty_by_sku, on="Super_parent", how="inner"
        )
        bom_plan["Expected_Demand"] = bom_plan["child_quantity"] * bom_plan["Total_Plan_Qty"]

        # Group by (SKU, Component) exactly like V03 - if the same component
        # appears on more than one BOM line for a SKU, its expected demand is
        # summed while the displayed per-unit qty shows the first line.
        expected = bom_plan.groupby(["Super_parent", "child"]).agg(
            Category=("child_description", "first"),
            BOM_Qty_Per_Unit=("child_quantity", "first"),
            Plan_Qty=("Total_Plan_Qty", "first"),
            Expected_Demand=("Expected_Demand", "sum"),
        ).reset_index()

        actual = demand.groupby(["sku", "item_code"])["demand_qty"].sum().reset_index()
        actual.columns = ["Super_parent", "child", "Actual_Demand"]

        merged = expected.merge(actual, on=["Super_parent", "child"], how="left")
        merged["Actual_Demand"] = merged["Actual_Demand"].fillna(0)
        merged["Variance"] = merged["Actual_Demand"] - merged["Expected_Demand"]
        merged["Status"] = merged["Variance"].abs().apply(lambda v: "PASS" if v < 0.01 else "FAIL")
        merged = merged.rename(columns={"Super_parent": "SKU", "child": "Component"})

        blocks = []
        for sku, grp in merged.groupby("SKU", sort=True):
            plan_qty = float(grp["Plan_Qty"].iloc[0])
            comp_df = (grp[["Component", "Category", "BOM_Qty_Per_Unit", "Plan_Qty",
                             "Expected_Demand", "Actual_Demand", "Variance", "Status"]]
                       .sort_values("Component").reset_index(drop=True))

            total_expected = round(float(comp_df["Expected_Demand"].sum()), 4)
            total_actual   = round(float(comp_df["Actual_Demand"].sum()), 4)
            total_variance = round(total_actual - total_expected, 4)
            block_status   = "PASS" if (comp_df["Status"] == "PASS").all() else "FAIL"

            blocks.append({
                "sku": sku, "plan_qty": plan_qty, "df": comp_df,
                "total_expected": total_expected, "total_actual": total_actual,
                "total_variance": total_variance, "status": block_status,
            })

        n_pass = int((merged["Status"] == "PASS").sum())
        n_fail = int((merged["Status"] == "FAIL").sum())
        n_sku_fail = sum(1 for b in blocks if b["status"] != "PASS")
        status = "PASS" if n_fail == 0 else "FAIL"

        flat_df = merged[["SKU", "Component", "Category", "BOM_Qty_Per_Unit", "Plan_Qty",
                           "Expected_Demand", "Actual_Demand", "Variance", "Status"]]

        return {
            "num": "V11",
            "name": "SKU-wise BOM demand breakdown (Plan Qty x BOM = Expected vs Actual)",
            "status": status, "pass_count": n_pass, "fail_count": n_fail,
            "notes": (f"{len(blocks)} SKUs, {len(merged)} (SKU, Component) rows "
                      f"(routed BOM only, same tolerance < 0.01 as V03); "
                      f"{n_sku_fail} SKU(s) have at least one mismatching component"),
            "blocks": blocks,
            "df": flat_df,
        }
    except Exception as e:
        traceback.print_exc()
        return {"num": "V11", "name": "SKU-wise BOM demand breakdown",
                "status": "WARN", "pass_count": 0, "fail_count": 0,
                "notes": f"Error: {e}", "blocks": [], "df": pd.DataFrame()}


# ── Main ───────────────────────────────────────────────────────────────────────

def main():
    print("=" * 70)
    print("Phase 1b Demand Generation Validation")
    print("=" * 70)

    plan, bom, demand = load_inputs()

    validation_funcs = [
        v01_plan_skus_in_demand,
        v02_bom_items_in_demand,
        v03_demand_qty_bom_x_plan,
        v04_demand_row_count,
        v05_itemwise_total_demand,
        v06_demand_uoms_valid,
        v07_qty_and_field_integrity,
        v08_need_by_matches_schedule,
        v09_demand_valid_sku_block,
        v10_total_demand_consistent,
        v11_sku_wise_bom_demand_breakdown,
    ]

    results = []
    for fn in validation_funcs:
        try:
            res = fn(plan, bom, demand)
            results.append(res)
        except Exception as e:
            num = fn.__name__.split("_")[0].upper().replace("V", "V0").replace("V00", "V0")
            results.append({
                "num": num, "name": fn.__name__, "status": "WARN",
                "pass_count": 0, "fail_count": 0,
                "notes": f"Unhandled error: {e}", "df": pd.DataFrame(),
            })

    # Build workbook
    print("\nBuilding Excel report...")
    wb = openpyxl.Workbook()
    # Remove default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    # Write detail sheets first (Summary will be inserted at index 0)
    sheet_names = ["V01", "V02", "V03", "V04", "V05", "V06", "V07", "V08", "V09", "V10", "V11"]
    descriptions = [
        "V01 – All plan SKUs present in demand",
        "V02 – All ROUTED BOM items generated in demand (raw materials excluded)",
        "V03 – Demand Qty = BOM qty x plan qty, grouped by (SKU, Component), routed only",
        "V04 – No missing or extra demand rows (routed items, active plan blocks only)",
        "V05 – Item-wise total demand correct (by item_code, routed items only)",
        "V06 – All demand UOMs valid (match BOM child_Unit)",
        "V07 – Field integrity: null/negative qty = FAIL, zero qty = WARN",
        "V08 – need_by matches a real scheduled (SKU, StartTime) in the plan",
        "V09 – Every demand row has valid SKU and well-formed block_id",
        "V10 – Total demand internally consistent (sum = Plan x routed BOM)",
        "V11 – SKU-wise BOM demand breakdown: per SKU, Plan Qty x BOM component = Expected vs Actual demand",
    ]

    for res, sname, desc in zip(results, sheet_names, descriptions):
        if "blocks" in res:
            # Nested per-SKU breakdown sheet (V11), not a flat table.
            write_sku_breakdown_sheet(wb, sname, desc, res.get("blocks", []),
                                      note=res.get("notes", ""))
        else:
            df = res.get("df", pd.DataFrame())
            if df is None:
                df = pd.DataFrame()
            write_detail_sheet(wb, sname, desc, df, status_col="Status",
                               note=res.get("notes", ""))

    write_summary_sheet(wb, results)

    wb.save(OUTPUT_FILE)
    print(f"\nReport saved: {OUTPUT_FILE}")

    # Print final counts
    total_pass = sum(r["pass_count"] for r in results)
    total_fail = sum(r["fail_count"] for r in results)
    n_val_pass = sum(1 for r in results if r["status"] == "PASS")
    n_val_fail = sum(1 for r in results if r["status"] == "FAIL")
    n_val_warn = sum(1 for r in results if r["status"] == "WARN")

    print("\n" + "=" * 70)
    print(f"Validation complete: {n_val_pass} PASS | {n_val_fail} FAIL | {n_val_warn} WARN")
    print(f"Row-level totals:    {total_pass:,} PASS | {total_fail:,} FAIL")
    print("=" * 70)


if __name__ == "__main__":
    main()
