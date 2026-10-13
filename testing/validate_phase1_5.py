#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
validate_phase1_5.py — Phase 1.5 (Wave Builder) validation.

Phase 1.5 takes the production blocks (one per valid curing-plan row) and buckets
them into 3-day WAVES. This script re-derives everything independently from the raw
plan + the two phase-1.5 outputs and runs 10 checks:

  V01  Every production block assigned to exactly one wave (no missing/duplicate).
  V02  Total planned quantity == total quantity assigned across all waves.
  V03  Every wave_id in block-assignment exists in the wave master.
  V04  n_blocks per wave == actual assigned blocks.
  V05  n_SKUs per wave == actual unique SKUs assigned.
  V06  tyres per wave == sum of assigned block quantities.
  V07  first_need / last_need == min / max need_by of assigned blocks.
  V08  Every block's need_by falls within its wave period [wave_start, wave_end].
  V09  Waves are sequential and do not overlap.
  V10  Every wave has >=1 block and every wave_id is unique.

INPUTS (resolved under testing_phases/):
  output/PHASEWISE/PHASE1_5_WAVES/phase1_5_waves.csv          (wave master)
  output/PHASEWISE/PHASE1_5_WAVES/phase1_5_block_to_wave.csv  (block -> wave)
  output/PHASEWISE/00_INPUTS/plan.csv                         (curing plan)

OUTPUT:
  phase1_5_validation_report.xlsx   (Summary + V01..V10 sheets, PASS/FAIL/WARN colouring)
"""
from __future__ import annotations
import os
from pathlib import Path
from datetime import datetime

import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- paths
BASE = Path(__file__).parent
PW = BASE / "output" / "PHASEWISE"
WAVES_FILE = PW / "PHASE1_5_WAVES" / "phase1_5_waves.csv"
BTW_FILE   = PW / "PHASE1_5_WAVES" / "phase1_5_block_to_wave.csv"
PLAN_FILE  = PW / "00_INPUTS" / "plan.csv"
OUTPUT_FILE = BASE / "phase1_5_validation_report.xlsx"

NON_PRODUCT_PLAN_CODES = {"CHANGEOVER", "MOULD_CLEAN"}
TOL = 0.01
MAX_ROWS = 2000

# --------------------------------------------------------------------------- style
FILL_PASS   = PatternFill("solid", fgColor="C6EFCE")
FILL_FAIL   = PatternFill("solid", fgColor="FFC7CE")
FILL_WARN   = PatternFill("solid", fgColor="FFEB9C")
FILL_HEADER = PatternFill("solid", fgColor="1F3864")
FONT_PASS   = Font(color="276221")
FONT_FAIL   = Font(color="9C0006")
FONT_WARN   = Font(color="9C5700")
FONT_HEADER = Font(color="FFFFFF", bold=True)
FONT_TITLE  = Font(size=14, bold=True, color="1F3864")


def _status_style(v):
    v = str(v).upper()
    if v == "PASS": return FILL_PASS, FONT_PASS
    if v == "FAIL": return FILL_FAIL, FONT_FAIL
    return FILL_WARN, FONT_WARN


def write_detail_sheet(wb, name, description, df, status_col="Status"):
    ws = wb.create_sheet(name)
    ws.cell(row=1, column=1, value=f"{name} — {description}").font = FONT_TITLE
    if df is None or len(df) == 0:
        ws.cell(row=3, column=1, value="No rows.").font = FONT_PASS
        return
    show = df.head(MAX_ROWS)
    hdr = 3
    for j, c in enumerate(show.columns, 1):
        cell = ws.cell(row=hdr, column=j, value=str(c))
        cell.fill = FILL_HEADER; cell.font = FONT_HEADER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (_, row) in enumerate(show.iterrows(), start=hdr + 1):
        for j, c in enumerate(show.columns, 1):
            val = row[c]
            if isinstance(val, (np.integer,)): val = int(val)
            elif isinstance(val, (np.floating,)): val = None if pd.isna(val) else float(val)
            elif isinstance(val, pd.Timestamp): val = None if pd.isna(val) else val.to_pydatetime()
            elif pd.isna(val): val = None
            cell = ws.cell(row=i, column=j, value=val)
            if c == status_col and val is not None:
                fill, font = _status_style(val); cell.fill = fill; cell.font = font
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1).coordinate
    if len(show):
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(show.columns))}{hdr+len(show)}"
    for j, c in enumerate(show.columns, 1):
        w = max([len(str(c))] + [len(str(x)) for x in show[c].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(46, max(10, w + 2))
    if len(df) > MAX_ROWS:
        ws.cell(row=hdr + len(show) + 2, column=1,
                value=f"... {len(show):,} of {len(df):,} rows shown (summary counts reflect ALL rows).").font = FONT_WARN


def write_summary_sheet(wb, results):
    ws = wb.create_sheet("Summary")
    ws.cell(row=1, column=1, value="Phase 1.5 — Wave Builder Validation Report").font = FONT_TITLE
    ws.cell(row=2, column=1, value=f"Run timestamp: {datetime.now():%Y-%m-%d %H:%M:%S}").font = Font(italic=True)
    hdr = 4
    for j, h in enumerate(["#", "Validation", "Status", "Pass", "Fail", "Notes"], 1):
        cell = ws.cell(row=hdr, column=j, value=h); cell.fill = FILL_HEADER; cell.font = FONT_HEADER
    for i, r in enumerate(results, start=hdr + 1):
        ws.cell(row=i, column=1, value=r["id"])
        ws.cell(row=i, column=2, value=r["name"])
        sc = ws.cell(row=i, column=3, value=r["status"])
        fill, font = _status_style(r["status"]); sc.fill = fill; sc.font = font
        ws.cell(row=i, column=4, value=r["pass"])
        ws.cell(row=i, column=5, value=r["fail"])
        ws.cell(row=i, column=6, value=r["notes"])
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 62
    ws.column_dimensions["C"].width = 9
    ws.column_dimensions["D"].width = 8
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 90
    ws.sheet_view.showGridLines = False
    wb.move_sheet("Summary", -(len(wb.sheetnames) - 1))


# --------------------------------------------------------------------------- load
def load():
    waves = pd.read_csv(WAVES_FILE)
    btw = pd.read_csv(BTW_FILE)
    plan = pd.read_csv(PLAN_FILE, dtype=str)
    # normalise
    for c in ("wave_id",):
        waves[c] = waves[c].astype(str).str.strip()
        btw[c] = btw[c].astype(str).str.strip()
    btw["block_id"] = btw["block_id"].astype(str).str.strip()
    btw["sku"] = btw["sku"].astype(str).str.strip()
    btw["qty"] = pd.to_numeric(btw["qty"], errors="coerce")
    btw["need_by"] = pd.to_datetime(btw["need_by"], errors="coerce")
    for c in ("first_need", "last_need", "wave_start", "wave_end"):
        waves[c] = pd.to_datetime(waves[c], errors="coerce")
    for c in ("n_blocks", "n_SKUs", "tyres"):
        waves[c] = pd.to_numeric(waves[c], errors="coerce")
    # plan -> valid production blocks (matches phase1.5: qty>0, not changeover/mould_clean)
    plan["qty"] = pd.to_numeric(plan["qty"], errors="coerce")
    plan["skuCode"] = plan["skuCode"].astype(str).str.strip()
    valid = plan[(plan["qty"] > 0) & (~plan["skuCode"].str.upper().isin(NON_PRODUCT_PLAN_CODES))].copy()
    return waves, btw, valid


# --------------------------------------------------------------------------- checks
def run_checks(waves, btw, valid):
    results = []
    details = {}

    def add(vid, name, status, npass, nfail, notes, df=None):
        results.append({"id": vid, "name": name, "status": status,
                        "pass": npass, "fail": nfail, "notes": notes})
        if df is not None:
            details[vid] = df

    # ---- V01: every production block assigned exactly once ----
    dup = btw["block_id"].value_counts()
    dup_ids = dup[dup > 1]
    n_plan, n_btw = len(valid), len(btw)
    unassigned = n_plan - btw["block_id"].nunique()
    ok = len(dup_ids) == 0 and n_plan == n_btw and unassigned == 0
    d = pd.DataFrame({"block_id": dup_ids.index, "times_assigned": dup_ids.values,
                      "Status": "FAIL"}) if len(dup_ids) else pd.DataFrame(
        [{"check": "plan_blocks", "value": n_plan, "Status": "PASS"},
         {"check": "assigned_blocks", "value": n_btw, "Status": "PASS"},
         {"check": "duplicate_assignments", "value": int(len(dup_ids)),
          "Status": "PASS" if not len(dup_ids) else "FAIL"}])
    add("V01", "Every production block assigned to exactly one wave",
        "PASS" if ok else "FAIL", n_btw - len(dup_ids), len(dup_ids),
        f"plan blocks={n_plan:,}, assigned={n_btw:,}, duplicates={len(dup_ids)}, unassigned={unassigned}", d)

    # ---- V02: total planned qty == total assigned qty ----
    pq, aq = float(valid["qty"].sum()), float(btw["qty"].sum())
    ok = abs(pq - aq) <= TOL
    add("V02", "Total planned quantity == total assigned to waves",
        "PASS" if ok else "FAIL", 1 if ok else 0, 0 if ok else 1,
        f"plan qty={pq:,.2f}  assigned qty={aq:,.2f}  diff={pq-aq:,.4f}",
        pd.DataFrame([{"plan_total_qty": pq, "assigned_total_qty": aq, "difference": pq - aq,
                       "Status": "PASS" if ok else "FAIL"}]))

    # ---- V03: every wave_id in btw exists in the wave master ----
    master_ids = set(waves["wave_id"])
    btw_ids = set(btw["wave_id"])
    orphan = sorted(btw_ids - master_ids)
    ok = not orphan
    add("V03", "Every wave_id in block-assignment exists in wave master",
        "PASS" if ok else "FAIL", len(btw_ids) - len(orphan), len(orphan),
        f"{len(btw_ids)} distinct wave_ids in assignment; {len(orphan)} not in master",
        pd.DataFrame({"orphan_wave_id": orphan, "Status": "FAIL"}) if orphan else
        pd.DataFrame([{"wave_id": w, "Status": "PASS"} for w in sorted(master_ids)]))

    # ---- recompute per-wave actuals from btw ----
    act = btw.groupby("wave_id").agg(
        n_blocks_actual=("block_id", "nunique"),
        n_SKUs_actual=("sku", "nunique"),
        tyres_actual=("qty", "sum"),
        first_need_actual=("need_by", "min"),
        last_need_actual=("need_by", "max")).reset_index()
    m = waves.merge(act, on="wave_id", how="left")

    # ---- V04: n_blocks matches ----
    m["v04"] = m["n_blocks"].fillna(-1).astype(int) == m["n_blocks_actual"].fillna(-2).astype(int)
    d = m[["wave_id", "n_blocks", "n_blocks_actual", "v04"]].copy()
    d["Status"] = np.where(d["v04"], "PASS", "FAIL"); d = d.drop(columns="v04")
    add("V04", "n_blocks per wave matches actual assigned blocks",
        "PASS" if m["v04"].all() else "FAIL", int(m["v04"].sum()), int((~m["v04"]).sum()),
        f"{int((~m['v04']).sum())} wave(s) mismatch", d)

    # ---- V05: n_SKUs matches ----
    m["v05"] = m["n_SKUs"].fillna(-1).astype(int) == m["n_SKUs_actual"].fillna(-2).astype(int)
    d = m[["wave_id", "n_SKUs", "n_SKUs_actual", "v05"]].copy()
    d["Status"] = np.where(d["v05"], "PASS", "FAIL"); d = d.drop(columns="v05")
    add("V05", "n_SKUs per wave matches actual unique SKUs",
        "PASS" if m["v05"].all() else "FAIL", int(m["v05"].sum()), int((~m["v05"]).sum()),
        f"{int((~m['v05']).sum())} wave(s) mismatch", d)

    # ---- V06: tyres == sum of assigned qty ----
    m["v06"] = (m["tyres"].fillna(-1) - m["tyres_actual"].fillna(-2)).abs() <= TOL
    d = m[["wave_id", "tyres", "tyres_actual"]].copy()
    d["difference"] = m["tyres"].fillna(0) - m["tyres_actual"].fillna(0)
    d["Status"] = np.where(m["v06"], "PASS", "FAIL")
    add("V06", "Total tyres per wave == sum of assigned block quantities",
        "PASS" if m["v06"].all() else "FAIL", int(m["v06"].sum()), int((~m["v06"]).sum()),
        f"{int((~m['v06']).sum())} wave(s) mismatch", d)

    # ---- V07: first_need / last_need match min/max need_by ----
    m["v07"] = (m["first_need"] == m["first_need_actual"]) & (m["last_need"] == m["last_need_actual"])
    d = m[["wave_id", "first_need", "first_need_actual", "last_need", "last_need_actual"]].copy()
    d["Status"] = np.where(m["v07"], "PASS", "FAIL")
    add("V07", "first_need / last_need match min / max need_by of assigned blocks",
        "PASS" if m["v07"].all() else "FAIL", int(m["v07"].sum()), int((~m["v07"]).sum()),
        f"{int((~m['v07']).sum())} wave(s) mismatch", d)

    # ---- V08: every block's need_by within its wave period ----
    wb_bounds = waves.set_index("wave_id")[["wave_start", "wave_end"]]
    j = btw.merge(wb_bounds, left_on="wave_id", right_index=True, how="left")
    j["in_period"] = (j["need_by"] >= j["wave_start"]) & (j["need_by"] <= j["wave_end"])
    bad = j[~j["in_period"]]
    ok = len(bad) == 0
    d = bad[["block_id", "sku", "need_by", "wave_id", "wave_start", "wave_end"]].copy()
    if len(d): d["Status"] = "FAIL"
    add("V08", "Every block's need_by falls within its wave period",
        "PASS" if ok else "FAIL", len(j) - len(bad), len(bad),
        f"{len(bad):,} of {len(j):,} blocks outside their wave window",
        d if len(d) else pd.DataFrame([{"blocks_checked": len(j), "outside": 0, "Status": "PASS"}]))

    # ---- V09: waves sequential and non-overlapping ----
    ws = waves.sort_values("wave_start").reset_index(drop=True)
    rows = []
    overlaps = 0
    for i in range(len(ws) - 1):
        gap_ok = ws.loc[i, "wave_end"] <= ws.loc[i + 1, "wave_start"] + pd.Timedelta(seconds=1)
        if not gap_ok: overlaps += 1
        rows.append({"wave_id": ws.loc[i, "wave_id"], "wave_end": ws.loc[i, "wave_end"],
                     "next_wave_id": ws.loc[i + 1, "wave_id"],
                     "next_wave_start": ws.loc[i + 1, "wave_start"],
                     "Status": "PASS" if gap_ok else "FAIL"})
    ok = overlaps == 0
    add("V09", "Waves are sequential and do not overlap",
        "PASS" if ok else "FAIL", len(rows) - overlaps, overlaps,
        f"{overlaps} overlapping wave boundary(ies)", pd.DataFrame(rows))

    # ---- V10: every wave has >=1 block and wave_id unique ----
    dup_w = waves["wave_id"].value_counts()
    dup_w = dup_w[dup_w > 1]
    empty = m[m["n_blocks_actual"].fillna(0) < 1]
    ok = len(dup_w) == 0 and len(empty) == 0
    d = pd.DataFrame([{"wave_id": w, "n_blocks_actual": int(r), "issue": "empty wave", "Status": "FAIL"}
                      for w, r in zip(empty["wave_id"], empty["n_blocks_actual"].fillna(0))]
                     + [{"wave_id": w, "n_blocks_actual": "-", "issue": "duplicate wave_id", "Status": "FAIL"}
                        for w in dup_w.index])
    if not len(d):
        d = pd.DataFrame([{"wave_id": w, "n_blocks_actual": int(n), "issue": "-", "Status": "PASS"}
                          for w, n in zip(m["wave_id"], m["n_blocks_actual"].fillna(0))])
    add("V10", "Every wave has >=1 block and every wave_id is unique",
        "PASS" if ok else "FAIL", len(waves) - len(empty) - len(dup_w), len(empty) + len(dup_w),
        f"empty waves={len(empty)}, duplicate wave_ids={len(dup_w)}", d)

    return results, details


def main():
    print("=" * 66)
    print("PHASE 1.5 — WAVE BUILDER VALIDATION")
    print("=" * 66)
    waves, btw, valid = load()
    print(f"  waves: {len(waves)} | block->wave: {len(btw):,} | plan valid blocks: {len(valid):,}")
    results, details = run_checks(waves, btw, valid)

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for r in results:
        df = details.get(r["id"])
        write_detail_sheet(wb, r["id"], r["name"], df)
    write_summary_sheet(wb, results)
    wb.save(OUTPUT_FILE)

    print("\n  RESULTS")
    for r in results:
        print(f"    {r['id']}  {r['status']:4s}  {r['name']}")
    npass = sum(r["status"] == "PASS" for r in results)
    print(f"\n  {npass}/{len(results)} PASS")
    print(f"  Report -> {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
