"""Format the plant's TDA (TBR apexed bead) schedule (tda_schedule.csv) in the same layout as the plant mixing and 4RC
schedules. Data is only grouped, never changed. Run from the repository root."""
import re

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = 'latest_schedule/plant-schedule/'
SRC = P + 'tda_schedule.csv'
OUT = P + 'TDA_SCHEDULE_PLANT_formatted_2403-2405.xlsx'
MAP = {'TDA1': '2403', 'TDA2': '2404', 'TDA3': '2405'}   # plant machine -> scheduler machine (confirmed by the user)
I = 'latest_schedule/v72/run_output/PHASEWISE/00_INPUTS/'
LONG_BOOKING_MIN = 45      # a booking this long (normal ~12-16 min) holds idle / stop time

raw = pd.read_csv(SRC)
d = raw.copy()
d['item_s'] = d['item'].astype(str).str.strip()
d['ci'] = d['consumed_item'].astype(str).str.strip()
d['start_time'] = pd.to_datetime(d['start_time'])
d['end_time'] = pd.to_datetime(d['end_time'])
d['plant_day'] = (d['start_time'] - pd.Timedelta(hours=7)).dt.date
d['mcode'] = d['machine'].map(MAP)

# ---------------- masters
rt = pd.read_csv(I + 'routing.csv', dtype=str, low_memory=False).apply(lambda c: c.str.strip())
rt = rt.drop_duplicates('routed_product').set_index('routed_product')
bom = pd.read_csv(I + 'bom.csv', dtype=str, low_memory=False).apply(lambda c: c.str.strip())
bom['per'] = pd.to_numeric(bom['child_quantity'], errors='coerce') / pd.to_numeric(bom['Parent_qty'], errors='coerce')
bom = bom.drop_duplicates(['Parent', 'child'])
per = bom.groupby(['Parent', 'child'])['per'].first().to_dict()
kids = bom.groupby('Parent')['child'].agg(list).to_dict()
it = pd.read_csv(I + 'itemtype_master.csv', dtype=str).apply(lambda c: c.str.strip())
ity = dict(zip(it['ItemCode'], it['ItemType']))
ag = pd.read_csv(I + 'aging_master.csv')
ag['k'] = ag['ItemCode'].astype(str).str.strip()
mul = {'Days': 24, 'Hours': 1}
mxa = dict(zip(ag['k'], ag['MaxAging'] * ag['MaxAgingUnit'].map(mul)))
mna = dict(zip(ag['k'], ag['MinAging'] * ag['MinAgingUnit'].map(mul)))


def per_bead(apex, ci):
    """input quantity per apexed bead, via the bead bundle (TBWB...) for wire and S54"""
    if (apex, ci) in per:
        return per[(apex, ci)]
    for b in kids.get(apex, []):
        if b.startswith('TBWB') and (b, ci) in per:
            return per[(apex, b)] * per[(b, ci)]
    return np.nan


def machines(k):
    return rt.loc[k, 'machines'].replace('"', '') if k in rt.index else '(not in routing)'


def speed(k):
    if k not in rt.index:
        return np.nan
    return float(rt.loc[k, 'proc_time'])          # efficiency 1 (config efficiency_override), as the scheduler


GROUP = {'1040S28600': 'bead wire', '1035BW3200': 'nylon fabric', 'SNCT001': 'slitted chafer', 'NOT APPLICABLE': 'none'}
d['ckind'] = d['ci'].map(lambda c: GROUP.get(c, 'compound' if ity.get(c, '').endswith('COMPOUND') else ity.get(c, 'other')))

# ---------------- bookings (prod_id) and runs
bk = d.groupby('prod_id').agg(machine=('machine', 'first'), item=('item_s', 'first'), shift=('shift', 'first'),
                              s=('start_time', 'min'), e=('end_time', 'max')).sort_values(['machine', 's'])
bk['min'] = (bk['e'] - bk['s']).dt.total_seconds() / 60
q = d.pivot_table(index='prod_id', columns='ci', values='qty', aggfunc='sum')
for r in bk.itertuples():
    est = {c: q.loc[r.Index, c] / per_bead(r.item, c) for c in ('1040S28600', 'A5542', 'SNCT001')
           if c in q.columns and q.loc[r.Index, c] == q.loc[r.Index, c] and per_bead(r.item, c) == per_bead(r.item, c)}
    bk.loc[r.Index, 'beads_est'] = np.median(list(est.values())) if est else np.nan
bk['run'] = ((bk['item'] != bk['item'].shift()) | (bk['machine'] != bk['machine'].shift())).cumsum()
bk['long'] = bk['min'] > LONG_BOOKING_MIN
d['run'] = d['prod_id'].map(bk['run'])

rows = []
for (run, sh, pdy), x in d.groupby(['run', 'shift', 'plant_day'], sort=False):
    item, mc = x['item_s'].iloc[0], x['machine'].iloc[0]
    b = bk.loc[x['prod_id'].unique()]
    beads = b['beads_est'].sum()
    sp = speed(item)
    kg = x.groupby('ckind')['qty'].sum()
    rows.append({
        'line': 'TBR', 'plant_day': pdy, 'date': pd.to_datetime(x['date'].iloc[0]).date(), 'shift': sh,
        'process': 'TBR apexed bead (bead winding + apexing)', 'machine_code': MAP[mc], 'plant_machine': mc,
        'machine_codes_allowed (routing)': machines(item),
        'routing_allows_this_machine': 'yes' if MAP[mc] in machines(item).split(',') else 'NO', 'item': item, 'item_type': ity.get(item, '(not in item master)'),
        'bead_bundle (BOM)': ', '.join(c for c in kids.get(item, []) if c.startswith('TBWB')),
        'run_no': int(run), 'start_time': b['s'].min(), 'end_time': b['e'].max(),
        'booked_min': round(b['min'].sum(), 1), 'bookings (prod_id)': len(b),
        'long_bookings (> %d min)' % LONG_BOOKING_MIN: int(b['long'].sum()),
        'beads_est (from BOM)': round(beads), 'UOM': 'NOS',
        'routing_min for beads (3 NOS/min, eff 1)': round(beads / sp, 1) if sp == sp else None,
        'bead_wire_kg': round(kg.get('bead wire', 0), 2), 'compound_kg': round(kg.get('compound', 0), 2),
        'slitted_chafer (as received)': round(kg.get('slitted chafer', 0), 2),
        'nylon_fabric (as received)': round(kg.get('nylon fabric', 0), 2),
        'consumed_items (qty)': '; '.join(f'{a} ({v:g})' for a, v in x.groupby('ci', sort=False)['qty'].sum().items() if a != 'NOT APPLICABLE'),
        'prod_ids': ';'.join(dict.fromkeys(x['prod_id'].astype(str))),
        'consumption_id': ';'.join(dict.fromkeys(x['consumption_id'].dropna().astype(str))),
        'source_rows': len(x)})
m = pd.DataFrame(rows).sort_values(['machine_code', 'start_time']).reset_index(drop=True)


def busy(g):
    return (g['e'].max() - g['s'].min()).total_seconds() / 3600


summ = []
for mc, g in bk.groupby('machine'):
    mm = m[m['plant_machine'] == mc]
    summ.append({'machine_code': MAP[mc], 'plant_machine': mc, 'items': g['item'].nunique(), 'runs': g['run'].nunique(), 'bookings': len(g),
                 'beads_est': round(g['beads_est'].sum()), 'first_start': g['s'].min(), 'last_end': g['e'].max(),
                 'booked_h': round(g['min'].sum() / 60, 1),
                 'long_bookings': int(g['long'].sum()), 'h_in_long_bookings': round(g.loc[g['long'], 'min'].sum() / 60, 1),
                 'routing_h for beads': round(mm['routing_min for beads (3 NOS/min, eff 1)'].sum() / 60, 1),
                 'bead_wire_kg': round(mm['bead_wire_kg'].sum(), 1), 'compound_kg': round(mm['compound_kg'].sum(), 1)})
summ = pd.DataFrame(summ)
tot = summ.drop(columns=['machine_code', 'plant_machine', 'first_start', 'last_end']).sum(numeric_only=True)
summ = pd.concat([summ, pd.DataFrame([{'machine_code': 'ALL', 'plant_machine': 'ALL', **tot.round(1).to_dict(), 'items': bk['item'].nunique(),
                                       'first_start': bk['s'].min(), 'last_end': bk['e'].max()}])], ignore_index=True)

by_shift = m.pivot_table(index=['machine_code', 'item'], columns='shift', values='beads_est (from BOM)', aggfunc='sum', fill_value=0)
by_shift['day'] = by_shift.sum(axis=1)
by_item = m.groupby('item').agg(machines=('machine_code', lambda s: '/'.join(sorted(set(s)))), runs=('run_no', 'nunique'),
                                bookings=('bookings (prod_id)', 'sum'), beads_est=('beads_est (from BOM)', 'sum'),
                                booked_min=('booked_min', 'sum')).reset_index()
by_item['allowed (routing)'] = by_item['item'].map(machines)

# compound age at use (consumption id F01/F02yyyymmddhhmmss_n = mixing batch booking time)
ages = []
for r in d[d['ckind'] == 'compound'].itertuples():
    mt = re.match(r'^[A-Z]\d{2}(\d{14})_', str(r.consumption_id))
    if mt:
        t = pd.to_datetime(mt.group(1), format='%Y%m%d%H%M%S', errors='coerce')
        if pd.notna(t):
            ages.append({'machine_code': MAP[r.machine], 'plant_machine': r.machine, 'item': r.item_s, 'compound': r.ci, 'compound_batch': r.consumption_id,
                         'batch_time': t, 'used_at': r.start_time, 'age_h': round((r.start_time - t).total_seconds() / 3600, 1),
                         'min_aging_h': mna.get(r.ci), 'max_aging_h': mxa.get(r.ci)})
AG = pd.DataFrame(ages)

chk = []
for r in by_item.to_dict('records'):
    if r['allowed (routing)'].startswith('('):
        chk.append({'item': r['item'], 'issue': 'item not in routing'})
for c in sorted(set(d['ci'])):
    if c != 'NOT APPLICABLE' and not any(per_bead(i, c) == per_bead(i, c) for i in set(d.loc[d['ci'] == c, 'item_s'])):
        chk.append({'item': c, 'issue': 'consumed input not in the BOM of the apexed bead or its bead bundle'})
for (i, c), g in d[d['ci'].isin(['1040S28600', 'A5542', 'A1031', 'B107', 'SNCT001', 'S54'])].groupby(['item_s', 'ci']):
    pb = per_bead(i, c)
    tot_b = bk.loc[bk['item'] == i, 'beads_est'].sum()
    if pb == pb and tot_b:
        ratio = g['qty'].sum() / (tot_b * pb)
        if abs(ratio - 1) > 0.10:
            chk.append({'item': i, 'issue': f'{c}: plant uses {ratio * 100:.0f} % of the BOM quantity for the estimated beads'})
nl = int(bk['long'].sum())
if nl:
    chk.append({'item': '', 'issue': f'{nl} bookings longer than {LONG_BOOKING_MIN} min ({bk.loc[bk["long"], "min"].sum() / 60:.1f} h) -- '
                                     'bookings run back to back, so stops / idle are inside these bookings'})
na = int((d['ci'] == 'NOT APPLICABLE').sum())
chk.append({'item': '', 'issue': f'{na} rows with consumed_item NOT APPLICABLE and qty 0 (kept in the source sheet, not counted)'})
nb = int(((d['ckind'] == 'compound') & (d['consumption_id'].astype(str) == 'Not Available')).sum())
if nb:
    chk.append({'item': '', 'issue': f'{nb} compound rows with consumption_id "Not Available" (no batch, age unknown)'})
if len(AG):
    for c, g in AG.groupby('compound'):
        lo, hi, mn, mx = g['age_h'].min(), g['age_h'].max(), g['min_aging_h'].iloc[0], g['max_aging_h'].iloc[0]
        bad = int((g['age_h'] > mx).sum()) if mx == mx else 0
        young = int((g['age_h'] < mn).sum()) if mn == mn else 0
        chk.append({'item': c, 'issue': f'compound age at use {lo:.0f}-{hi:.0f} h (aging min {mn:g} / max {mx:g} h): '
                                        + (f'{bad} uses older than max; ' if bad else '') + (f'{young} uses younger than min' if young else '')
                                        + ('OK' if not bad and not young else '')})
chk.append({'item': 'TDA1/2/3', 'issue': 'plant TDA1 = 2403, TDA2 = 2404, TDA3 = 2405 (as given by the plant)'})
bad = m[m['routing_allows_this_machine'] == 'NO']
for r in bad.drop_duplicates(['machine_code', 'item']).itertuples():
    chk.append({'item': r.item, 'issue': f'made on {r.machine_code} but routing allows only {machines(r.item)}'})
if not len(bad):
    chk.append({'item': 'all', 'issue': 'every item was made on a machine its routing allows (2403 / 2404 / 2405)'})
chk = pd.DataFrame(chk)

note = pd.DataFrame({'note': [
    f'Source: {SRC} ({len(raw):,} rows) -- data NOT changed; the original rows are in sheet "Source data (as received)".',
    'TDA = TBR apexed bead machine. One booking (prod_id) uses bead wire 1040S28600 + S54 (bead bundle) AND the apex '
    'inputs A1031, A5542, B107, SNCT001 (and 1035BW3200 for TBAPBD005): bead winding and apexing in one step on the plant machine.',
    'TDA schedule = bookings grouped by production run (consecutive bookings of the same item on the same machine) and shift. '
    'start = first / end = last booking time; consumed quantities summed per input.',
    'beads_est: the file gives no output quantity. Beads per booking = median of (bead wire, A5542, SNCT001 consumed) / BOM '
    'quantity per apexed bead (bead wire and S54 via the bead bundle TBWB...). These three agree within a few % for most bookings.',
    'routing_min = beads_est / (3 NOS/min, efficiency 1 as in the scheduler) from BOM_Routing_9oct; booked_min is the plant booking time '
    f'(back to back, so bookings > {LONG_BOOKING_MIN} min contain stops).',
    'Item type, routing machines (2403/2404/2405), BOM and aging from the v72 inputs; compound age from the consumption id '
    '(F01/F02 + mixing batch time). Plant TDA1 / TDA2 / TDA3 = machines 2403 / 2404 / 2405.',
    f'Data window: {d["start_time"].min():%d-%b %H:%M} -> {d["end_time"].max():%d-%b %H:%M} (one plant day 08-Oct).']})

hdr, bold = PatternFill('solid', fgColor='FF1F3864'), Font(name='Arial', bold=True, color='FFFFFFFF')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    sheets = [('Summary', summ), ('Beads by item', by_item), ('Beads by shift', by_shift.reset_index()), ('TDA schedule', m)]
    sheets += [(f'{mc} runs', m[m['machine_code'] == mc]) for mc in sorted(m['machine_code'].unique())]
    sheets += [('Bookings (prod_id)', bk.reset_index().drop(columns='long')), ('Compound age', AG),
               ('Source data (as received)', raw), ('Notes', note), ('Data checks', chk)]
    for name, df in sheets:
        df.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical='center')
        for j, c in enumerate(df.columns, 1):
            wd = min(max(10, int(df[c].astype(str).str.len().head(300).max() if len(df) else 10) + 2, len(str(c)) + 2), 60)
            ws.column_dimensions[get_column_letter(j)].width = wd
            if str(c) in ('start_time', 'end_time', 'first_start', 'last_end', 's', 'e', 'batch_time', 'used_at'):
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=j).number_format = 'dd-mmm-yy hh:mm'
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'

pd.set_option('display.width', 250)
pd.set_option('display.max_colwidth', 140)
print('source rows', len(raw), '| bookings', raw['prod_id'].nunique(), '=', int(m['bookings (prod_id)'].sum()),
      '| qty', round(raw['qty'].sum(), 2), '=', round(m[['bead_wire_kg', 'compound_kg', 'slitted_chafer (as received)',
                                                          'nylon_fabric (as received)']].sum().sum(), 2))
print(summ.to_string(index=False))
print(by_item.to_string(index=False))
print(chk.to_string(index=False))
print(m[['machine_code', 'plant_machine', 'routing_allows_this_machine', 'shift', 'item', 'start_time', 'end_time', 'booked_min', 'bookings (prod_id)', 'beads_est (from BOM)',
         'routing_min for beads (3 NOS/min, eff 1)']].to_string(index=False))
print('written', OUT)
