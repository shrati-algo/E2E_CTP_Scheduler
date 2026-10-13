"""Plant 4-roll calender (four_roll_schedule (1).csv, 30-Sep -> 11-Oct, 11 plant days) vs v82 (10-20 Oct), day by day:
process pattern (what runs when, how much, how often), the 4RC rules, speeds, spools / dips, compound age.
Run from the repository root."""
import re

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = 'latest_schedule/plant-schedule/'
V = 'v82'
I = f'latest_schedule/{V}/run_output/PHASEWISE/00_INPUTS/'
OUT = P + f'PLANT_vs_{V}_4RC_PROCESS.xlsx'
ALIAS = {'CAP 66': 'CAP 66-MOTHERROLL'}
CREEL_MIN = 240.0

# ---------------- masters (v82 inputs)
fl = pd.read_csv(I + 'frc_lengths.csv')
fl['item'] = fl['item'].astype(str).str.strip()
cls = dict(zip(fl['item'], fl['class'].astype(str).str.lower()))
wire = dict(zip(fl['item'], fl['wire'].astype(str)))
dipl = dict(zip(fl['item'], fl['spool_dip_length_m']))
rt = pd.read_csv(I + 'routing.csv', dtype=str, low_memory=False)
rt = rt[rt['machines'].astype(str).str.contains('901')]
spd = {p.strip(): float(t) for p, t in zip(rt['routed_product'], rt['proc_time'])}
bom = pd.read_csv(I + 'bom.csv', dtype=str, low_memory=False)
bom = bom.apply(lambda c: c.str.strip() if c.dtype == object else c)
bom['kgpm'] = pd.to_numeric(bom['child_quantity'], errors='coerce') / pd.to_numeric(bom['Parent_qty'], errors='coerce')
kgpm = bom.drop_duplicates(['Parent', 'child']).groupby(['Parent', 'child'])['kgpm'].first().to_dict()
ag = pd.read_csv(I + 'aging_master.csv')
ag['k'] = ag['ItemCode'].astype(str).str.strip()
mul = {'Days': 24, 'Hours': 1}
mxa = dict(zip(ag['k'], ag['MaxAging'] * ag['MaxAgingUnit'].map(mul)))
mna = dict(zip(ag['k'], ag['MinAging'] * ag['MinAgingUnit'].map(mul)))


def kind(ci):
    ci = str(ci).strip()
    if re.match(r'^\d', ci):
        return 'cord'
    if ci.upper().startswith('CAP STRIP'):
        return 'cap strip'
    if ci.upper() == 'EFA':
        return 'other'
    return 'compound'


def pday(t):
    return (t - pd.Timedelta(hours=7)).normalize()


# ---------------- plant rolls and runs
d = pd.read_csv(P + 'four_roll_schedule (1).csv')
for c in ('item', 'consumed_item', 'consumption_id'):
    d[c] = d[c].astype(str).str.strip()
d['start_time'] = pd.to_datetime(d['start_time'])
d['end_time'] = pd.to_datetime(d['end_time'])
d['k'] = d['item'].map(lambda x: ALIAS.get(x, x))
d['kind'] = d['consumed_item'].map(kind)


def roll_metres(g):
    k = g['k'].iloc[0]
    for kd in ('compound', 'cord'):
        x = g[g['kind'] == kd]
        if len(x):
            m = [q / kgpm.get((k, c), np.nan) for c, q in x.groupby('consumed_item')['qty'].sum().items()]
            m = [v for v in m if v == v]
            if m:
                return max(m)
    return np.nan


rolls = d.groupby('prod_id').agg(s=('start_time', 'min'), e=('end_time', 'max'), k=('k', 'first')).sort_values('s')
rolls['m'] = d.groupby('prod_id').apply(roll_metres, include_groups=False)
rolls['min'] = (rolls['e'] - rolls['s']).dt.total_seconds() / 60
rolls['gap_before'] = (rolls['s'] - rolls['e'].shift()).dt.total_seconds() / 60
rolls['run'] = (rolls['k'] != rolls['k'].shift()).cumsum()
rolls['class'] = rolls['k'].map(cls).fillna('?')
rolls['pday'] = rolls['s'].map(pday)
runs = rolls.groupby('run').agg(k=('k', 'first'), s=('s', 'min'), e=('e', 'max'), m=('m', 'sum'), rolls=('m', 'size'),
                                run_min=('min', 'sum')).reset_index()
runs['class'] = runs['k'].map(cls).fillna('?')
runs['wire'] = runs['k'].map(wire)
runs['pday'] = runs['s'].map(pday)
runs['span_min'] = (runs['e'] - runs['s']).dt.total_seconds() / 60
runs['speed'] = runs['m'] / runs['run_min']
runs['routing'] = runs['k'].map(spd)
runs['src'] = 'PLANT'

# ---------------- v82 lots -> runs
L = pd.read_excel(f'latest_schedule/{V}/FOURRC_RULES_{V}.xlsx', '901 lots')
L['item'] = L['item'].astype(str).str.strip()
L = L.sort_values('scheduled_start').reset_index(drop=True)
L['run'] = (L['item'] != L['item'].shift()).cumsum()
vr = L.groupby('run').agg(k=('item', 'first'), s=('scheduled_start', 'min'), e=('scheduled_finish', 'max'), m=('metres', 'sum'),
                          rolls=('metres', 'size'), run_min=('run_min_sched', 'sum')).reset_index()
vr['class'] = vr['k'].map(cls).fillna('?')
vr['wire'] = vr['k'].map(wire)
vr['pday'] = vr['s'].map(pday)
vr['span_min'] = (vr['e'] - vr['s']).dt.total_seconds() / 60
vr['speed'] = vr['m'] / vr['run_min']
vr['routing'] = vr['k'].map(spd)
vr['src'] = V
W0, W1 = pd.Timestamp('2026-10-10'), pd.Timestamp('2026-10-19')
vr_w = vr[(vr['pday'] >= W0) & (vr['pday'] <= W1)]


# ---------------- per-day process measures + rule checks (same code for both)
def day_measures(R, src):
    out = []
    for dd, g in R.groupby('pday'):
        g = g.sort_values('s').reset_index(drop=True)
        st = g[g['class'] == 'steel'].reset_index(drop=True)
        cc = cg = c3 = 0
        for i in range(1, len(st)):
            if wire.get(st.loc[i - 1, 'k']) != wire.get(st.loc[i, 'k']):
                cc += 1
                cg += (st.loc[i, 's'] - st.loc[i - 1, 'e']).total_seconds() / 60 >= CREEL_MIN - 1
                c3 += ((g['s'] >= st.loc[i - 1, 'e']) & (g['s'] < st.loc[i, 's']) & (g['class'] == 'fabric')).any()
        wl = g['wire'].tolist()
        loads = [w for i, w in enumerate(wl) if i == 0 or w != wl[i - 1]]
        rep = pd.Series(loads).value_counts() if loads else pd.Series(dtype=int)
        one_go = g.groupby('k').size()
        gaps = [(g.loc[i, 's'] - g.loc[i - 1, 'e']).total_seconds() / 60 for i in range(1, len(g))]
        out.append({'src': src, 'plant_day': dd.date(), 'runs': len(g), 'items': g['k'].nunique(),
                    'steel km': round(g.loc[g['class'] == 'steel', 'm'].sum() / 1000, 1),
                    'fabric km': round(g.loc[g['class'] == 'fabric', 'm'].sum() / 1000, 1),
                    'total km': round(g['m'].sum() / 1000, 1),
                    'running h': round(g['run_min'].sum() / 60, 1),
                    'idle between runs h': round(sum(x for x in gaps if x > 0) / 60, 1),
                    'cord changes': cc, 'creel >=240 min': cg, 'fabric between cord changes': c3,
                    'items in >1 run': int((one_go > 1).sum()), 'wire loads repeated': int((rep > 1).sum()),
                    'median run km': round(g['m'].median() / 1000, 2)})
    return pd.DataFrame(out)


DP = day_measures(runs, 'PLANT')
DP = DP[(DP['plant_day'] >= pd.Timestamp('2026-10-01').date()) & (DP['plant_day'] <= pd.Timestamp('2026-10-10').date())]
DV = day_measures(vr_w, V)

# ---------------- item frequency and lot size
def item_profile(R, ndays, src):
    g = R.groupby('k').agg(runs=('k', 'size'), days=('pday', 'nunique'), km=('m', 'sum'), med_run_km=('m', 'median'))
    g['km'] = (g['km'] / 1000).round(1)
    g['med_run_km'] = (g['med_run_km'] / 1000).round(2)
    g['km per day'] = (g['km'] / ndays).round(1)
    g['run every n days'] = (ndays / g['runs']).round(1)
    g.columns = [f'{src} {c}' for c in g.columns]
    return g


pl_w = runs[(runs['pday'] >= pd.Timestamp('2026-10-01')) & (runs['pday'] <= pd.Timestamp('2026-10-10'))]
IP = item_profile(pl_w, len(DP), 'PLANT').join(item_profile(vr_w, len(DV), V), how='outer')
IP['class'] = IP.index.map(cls)

# ---------------- speeds: plant per item (rolls without a stop: roll minutes <= 1.5 x median roll minutes of the item)
rolls['rmed'] = rolls.groupby('k')['min'].transform('median')
clean = rolls[(rolls['min'] <= 1.5 * rolls['rmed']) & (rolls['m'] > 0)]
SP = clean.groupby('k').apply(lambda x: pd.Series({'rolls': len(x), 'plant m/min (running)': x['m'].sum() / x['min'].sum(),
                                                  'plant m/min incl. stops': rolls.loc[rolls['k'] == x.name, 'm'].sum()
                                                  / rolls.loc[rolls['k'] == x.name, 'min'].sum(),
                                                  'median roll m': x['m'].median()}), include_groups=False)
SP['routing m/min'] = SP.index.map(spd)
SP['plant / routing'] = (SP['plant m/min (running)'] / SP['routing m/min']).round(2)
SP['class'] = SP.index.map(cls)
SP = SP.round(1)

# ---------------- cord spools: how a spool is used across loads (consumption id = spool id _ creel _ position)
cord = d[d['kind'] == 'cord'].copy()
cord['spool'] = cord['consumption_id'].str.replace(r'_\d+_\d+$', '', regex=True)
cord['run'] = cord['prod_id'].map(rolls['run'])
SPL = cord.groupby('spool').agg(cord=('consumed_item', 'first'), items=('k', lambda s: '+'.join(dict.fromkeys(s))),
                                runs=('run', 'nunique'), first=('start_time', 'min'), last=('end_time', 'max'),
                                kg=('qty', 'sum'), positions=('consumption_id', 'nunique'))
SPL['days on creel'] = ((SPL['last'] - SPL['first']).dt.total_seconds() / 86400).round(2)

# ---------------- CAP 66 inputs and compound age
cap = d[d['k'] == 'CAP 66-MOTHERROLL']
CAPIN = cap.groupby('consumed_item')['qty'].sum().round(0)
ages = []
for r in d[d['kind'] == 'compound'].itertuples():
    m_ = re.match(r'^[A-Z]\d{2}(\d{14})_', r.consumption_id)
    if m_:
        t = pd.to_datetime(m_.group(1), format='%Y%m%d%H%M%S', errors='coerce')
        if pd.notna(t):
            ages.append({'compound': r.consumed_item, 'age_h': (r.start_time - t).total_seconds() / 3600,
                         'min_h': mna.get(r.consumed_item), 'max_h': mxa.get(r.consumed_item)})
AG = pd.DataFrame(ages)
AGS = AG.groupby('compound').agg(uses=('age_h', 'size'), min_age=('age_h', 'min'), median_age=('age_h', 'median'),
                                 max_age=('age_h', 'max'), master_min=('min_h', 'first'), master_max=('max_h', 'first'))
AGS['over max'] = AG.assign(o=AG['age_h'] > AG['max_h']).groupby('compound')['o'].sum()
AGS['under min'] = AG.assign(u=AG['age_h'] < AG['min_h']).groupby('compound')['u'].sum()
AGS = AGS.round(1)

# ---------------- shift pattern: share of steel / fabric running by shift
def shift_of(t):
    h = t.hour + t.minute / 60
    return 'A 07-15' if 7 <= h < 15 else ('B 15-23' if 15 <= h < 23 else 'C 23-07')


SH = pd.concat([pl_w.assign(shift=pl_w['s'].map(shift_of)), vr_w.assign(shift=vr_w['s'].map(shift_of))])
SHP = SH.pivot_table(index=['src', 'shift'], columns='class', values='m', aggfunc='sum').fillna(0).div(1000).round(1)

# ---------------- write
hdr, bold = PatternFill('solid', fgColor='FF1F3864'), Font(name='Arial', bold=True, color='FFFFFFFF')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    for name, df in [('Days plant', DP), (f'Days {V}', DV), ('Items', IP.reset_index()), ('Speeds plant', SP.reset_index()),
                     ('Cord spools plant', SPL.reset_index()), ('Compound age plant', AGS.reset_index()),
                     ('Shift pattern', SHP.reset_index()), ('Runs plant', runs), (f'Runs {V}', vr_w)]:
        df.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical='center')
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, len(str(c)) + 2), 40)
        ws.freeze_panes = 'A2'
pd.set_option('display.width', 260)
pd.set_option('display.max_rows', 200)
print(DP.to_string(index=False)); print(DV.to_string(index=False))
num = ['runs', 'items', 'steel km', 'fabric km', 'total km', 'running h', 'idle between runs h', 'cord changes', 'creel >=240 min',
       'fabric between cord changes', 'items in >1 run', 'wire loads repeated', 'median run km']
print(pd.DataFrame({'PLANT avg/day': DP[num].mean().round(2), f'{V} avg/day': DV[num].mean().round(2),
                    'PLANT sum': DP[num].sum(), f'{V} sum': DV[num].sum()}).to_string())
print(IP.to_string()); print(SP.to_string()); print(SHP.to_string())
print(SPL.sort_values('first').head(40).to_string()); print('spools', len(SPL), 'on >1 run', int((SPL['runs'] > 1).sum()))
print(CAPIN.to_string()); print(AGS.to_string())
print('written', OUT)
