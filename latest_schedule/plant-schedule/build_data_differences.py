"""Data-level differences only (masters vs what the plant schedules show): mixing, 4RC, TBR bead (TDA).
Process / practice differences are left out. Reads the comparison outputs already in this folder plus a re-run of the
mixing data checks against the v72 masters (BOM_Routing_9oct). Run from the repository root."""
import re
import sys

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = 'latest_schedule/plant-schedule/'
I = 'latest_schedule/v72/run_output/PHASEWISE/00_INPUTS/'
MIXCHK = sys.argv[1]                       # mixing data checks re-run on the v72 masters (compare_plant_v47 with v72 inputs)
OUT = P + 'PLANT_vs_MASTER_DATA_DIFFERENCES.xlsx'

rows = []


def add(area, master, field, item, machine, ours, plant, diff, note=''):
    rows.append({'area': area, 'master table': master, 'field': field, 'item': item, 'machine': machine,
                 'our master value': ours, 'plant value': plant, 'difference': diff, 'note': note})


# ---------------- mixing
mc = pd.read_excel(MIXCHK, 'Plant data checks')
for r in mc[mc['check'].str.startswith('D1')].itertuples():
    if r.status == 'NO DATA':
        add('Mixing', 'CTP_Master_Routing / BOM', 'routing + BOM row', r.compound, r.mixer, 'missing', f'{r.batches} batches mixed', 'item missing')
    else:
        add('Mixing', 'CTP_Master_Routing', 'machines', r.compound, r.mixer, r.detail.replace('routing allows ', ''), f'mixed on {r.mixer} ({r.batches} batches)', 'mixer not allowed')
cy = pd.read_excel(P + 'PLANT_vs_SCHEDULER_MIXING_CYCLE_TIMES.xlsx')
cy = cy[cy['our_cycle_s'].notna()]
for r in cy.itertuples():
    if r.plant_batches >= 5 and (r._8 < 90 or r._8 > 110):
        add('Mixing', 'CTP_Master_Routing', 'proc_time (cycle s / batch)', r.item, r.mixer, round(r.route_cycle_s, 1), round(r.plant_cycle_s, 1),
            f'{r._8:.0f} % of ours', f'{r.plant_batches} plant batches')
    if r.route_batch_kg and abs(r.plant_batch_kg - r.route_batch_kg) / r.route_batch_kg > 0.15:
        add('Mixing', 'CTP_Master_Routing', 'batch_size (kg)', r.item, r.mixer, r.route_batch_kg, r.plant_batch_kg,
            f'{(r.plant_batch_kg / r.route_batch_kg - 1) * 100:+.0f} %', f'{r.plant_batches} plant batches')
for r in mc[mc['check'].str.startswith('D3') & (mc['status'] == 'DEVIATION')].itertuples():
    devs = [int(x) for x in re.findall(r'\(([+-]\d+) %\)', r.detail)]
    if 'not in BOM' in r.detail or 'not consumed' in r.detail:
        kind = 'different ingredient(s)'
    elif len(devs) >= 2 and max(devs) - min(devs) <= 8 and min(abs(x) for x in devs) >= 15:
        kind = 'all ingredients off by the same % -> BOM Parent_qty vs sum of ingredients'
    else:
        kind = 'ingredient quantity'
    add('Mixing', 'BOM (BOM_Routing_9oct)', 'recipe (kg per batch)', r.compound, r.mixer, 'BOM', 'plant consumption', kind, r.detail[:400])

# ---------------- 4RC
fr = pd.read_excel(P + 'PLANT_vs_v72_4RC_COMPARISON.xlsx', 'Plant runs')
for r in fr.to_dict('records'):
    add('4RC', 'CTP_Master_Routing', 'line speed m/min (eff 1)', r['item_code'], '901', r['routing speed (m/min, eff 1)'], r['speed m/min (actual)'],
        f"run took {r['actual / routing time %']} % of routing time", f"{r['metres used']:,.0f} m in {r['minutes']:.0f} min (plant roll times; may include stops)")
    mc_, mf = r['metres (from compound kg / BOM kg per m)'], r['metres (from cord / fabric kg / BOM kg per m)']
    if mc_ and mf and abs(mc_ / mf - 1) > 0.05:
        add('4RC', 'BOM (BOM_Routing_9oct)', 'compound : cord / fabric kg per metre', r['item_code'], '901', 'BOM ratio', 'plant consumption',
            f'{(mc_ / mf - 1) * 100:+.0f} %', f'metres from compound {mc_:,.0f} vs from cord / fabric {mf:,.0f}')
    if isinstance(r['inputs not in BOM'], str) and r['inputs not in BOM']:
        add('4RC', 'BOM (BOM_Routing_9oct)', 'children', r['item_code'], '901', 'fabric + B616M', f"also {r['inputs not in BOM']}", 'input not in BOM', '4,415 units consumed')
add('4RC', 'CTP_Master_FRC', 'spool length (m)', 'CPJ1218 + CPJ370 (wire 1036S34500)', '901', 12600, 'creel changed after 29,441 m = 2.34 spools',
    'not whole spools', 'real spool length, or cord taken off before the end')
add('4RC', 'CTP_Master_FRC', 'dip length (m)', 'ECOPOLY350', '901', 2000, '11,074-11,830 m = 5.5-5.9 dips', 'not whole dips', '')
add('4RC', 'item master', 'item code', 'CAP 66-MOTHERROLL', '901', 'CAP 66-MOTHERROLL', 'CAP 66', 'name differs', '')

# ---------------- TBR bead (TDA)
ct = pd.read_excel(P + 'PLANT_vs_v72_TDA_COMPARISON.xlsx', 'Cycle time by item')
ct = ct[ct['plant beads / min (total)'].notna()]
for r in ct.to_dict('records'):
    add('TBR bead', 'CTP_Master_Routing', 'proc_time apex (+ winding) NOS/min', r['item'], r['machine'],
        f"apex {r['routing apex NOS/min (eff 1)']} + winding {r['routing winding NOS/min (eff 1)']} = {r['v72 min per bead (winding + apex)']} min/bead",
        f"{r['plant beads / min (total)']} beads/min = {r['plant min per bead (one booking)']} min/bead (winding included)",
        f"{r['plant as % of v72 winding + apex time']:.0f} % of our time", f"inch {r['inch']}")
co = pd.read_excel(P + 'PLANT_vs_v72_TDA_COMPARISON.xlsx', 'Plant changeovers')
for r in co.to_dict('records'):
    add('TBR bead', 'CTP_Master_ChangeoverMatrix (TBR Bead Apexing TDA n)', 'changeover_time_min', f"{r['from']} -> {r['to']}", r['machine'],
        r['matrix min'] if r['matrix min'] == r['matrix min'] else 'missing pair', f"{max(r['extra min (changeover / stop)'], 0):.0f} min extra in the first booking",
        'pair missing' if r['matrix min'] != r['matrix min'] else 'plant shorter' if r['extra min (changeover / stop)'] < r['matrix min'] - 5 else 'about same / longer', '')
add('TBR bead', 'machine master', 'bead inch per machine', '20" / 22.5" beads', '2403-2405', 'no inch data (any item on any routing machine)',
    '2404 = 22.5" only; 2403 / 2405 = 20" only', 'not in masters', 'inch derived from CTP_Master_BuildingAllowable.size: tbr_bead_inch.csv')
add('TBR bead', 'BOM (BOM_Routing_9oct)', 'children per bead', 'TBAPBD005', '2403', '1040S28600 via TBWB004', '89 % of BOM bead wire', '-11 %',
    'all other TDA inputs within +-10 % of the BOM (SNCT001 91-106 %)')

D = pd.DataFrame(rows)
S = D.groupby(['area', 'master table', 'field', 'difference' if False else 'area']).size() if False else \
    D.groupby(['area', 'master table', 'field']).size().rename('items / pairs').reset_index()

hdr, bold = PatternFill('solid', fgColor='FF1F3864'), Font(name='Arial', bold=True, color='FFFFFFFF')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    for name, df in [('Summary', S), ('All differences', D)] + [(a, D[D['area'] == a]) for a in D['area'].unique()]:
        df.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical='center')
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, int(df[c].astype(str).str.len().head(300).max() if len(df) else 10) + 2, len(str(c)) + 2), 70)
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'
pd.set_option('display.width', 250)
pd.set_option('display.max_colwidth', 90)
print(S.to_string(index=False))
d3 = D[D['field'] == 'recipe (kg per batch)']
print(d3['difference'].value_counts())
print(d3[d3['difference'].str.startswith('all')][['item', 'note']].to_string(index=False))
print(D[(D['area'] == 'Mixing') & (D['field'].str.startswith('proc'))]['difference'].str.extract(r'(\d+)')[0].astype(int).describe())
print('written', OUT)
