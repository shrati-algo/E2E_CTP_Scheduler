"""Compare the plant's 4-roll calender day (four_roll_schedule.csv, 08-Oct) with the v72 schedule on the same plant day,
rule by rule (4RC rules 1-10) plus data checks against the v72 masters (routing, BOM, CTP_Master_FRC, aging).
Run from the repository root."""
import re

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = 'latest_schedule/plant-schedule/'
V = 'v72'
I = f'latest_schedule/{V}/run_output/PHASEWISE/00_INPUTS/'
OUT = P + f'PLANT_vs_{V}_4RC_COMPARISON.xlsx'
ALIAS = {'CAP 66': 'CAP 66-MOTHERROLL'}
DAY0, DAY1 = pd.Timestamp('2026-10-08 07:00'), pd.Timestamp('2026-10-09 07:00')
CREEL_MIN = 240.0

# ---------------- masters
fl = pd.read_csv(I + 'frc_lengths.csv')
fl['item'] = fl['item'].astype(str).str.strip()
cls = dict(zip(fl['item'], fl['class'].astype(str).str.lower()))
wire = dict(zip(fl['item'], fl['wire'].astype(str)))
dipl = dict(zip(fl['item'], fl['spool_dip_length_m']))
minr = dict(zip(fl['item'], fl['min_roll_m']))
rt = pd.read_csv(I + 'routing.csv', dtype=str, low_memory=False)
rt = rt[rt['machines'].astype(str).str.contains('901')]
spd = {p.strip(): float(t) for p, t in zip(rt['routed_product'], rt['proc_time'])}   # efficiency 1 (config efficiency_override), as the scheduler
bom = pd.read_csv(I + 'bom.csv', dtype=str, low_memory=False)
bom = bom.apply(lambda c: c.str.strip() if c.dtype == object else c)
bom['kgpm'] = pd.to_numeric(bom['child_quantity'], errors='coerce') / pd.to_numeric(bom['Parent_qty'], errors='coerce')
kgpm = bom.drop_duplicates(['Parent', 'child']).groupby(['Parent', 'child'])['kgpm'].first().to_dict()
bom_children = bom.groupby('Parent')['child'].agg(lambda s: set(s)).to_dict()
ag = pd.read_csv(I + 'aging_master.csv')
ag['k'] = ag['ItemCode'].astype(str).str.strip()
mul = {'Days': 24, 'Hours': 1}
mxa = dict(zip(ag['k'], ag['MaxAging'] * ag['MaxAgingUnit'].map(mul)))
mna = dict(zip(ag['k'], ag['MinAging'] * ag['MinAgingUnit'].map(mul)))
mx = pd.read_csv(I + 'changeover_matrix.csv', low_memory=False)
mx = mx[mx['machine'].astype(str).str.strip() == '4 RC Calendar']
co = {(str(a).strip(), str(b).strip()): t for a, b, t in zip(mx['from_MaterialCode_O'], mx['to_MaterialCode_O'],
                                                            pd.to_numeric(mx['changeover_time_min'], errors='coerce'))}


def comp_kind(ci):
    ci = str(ci).strip()
    if re.match(r'^\d', ci):
        return 'cord / fabric'
    if ci.upper().startswith('CAP STRIP'):
        return 'cap strip'
    return 'compound'


# ---------------- plant runs
d = pd.read_csv(P + 'four_roll_schedule.csv')
d['item'] = d['item'].astype(str).str.strip()
d['consumed_item'] = d['consumed_item'].astype(str).str.strip()
d['start_time'] = pd.to_datetime(d['start_time'])
d['end_time'] = pd.to_datetime(d['end_time'])
d['k'] = d['item'].map(lambda x: ALIAS.get(x, x))
d['ck'] = d['consumed_item'].map(comp_kind)
rolls = d.groupby('prod_id').agg(s=('start_time', 'min'), e=('end_time', 'max'), k=('k', 'first')).sort_values('s')
rolls['run'] = (rolls['k'] != rolls['k'].shift()).cumsum()
d['run'] = d['prod_id'].map(rolls['run'])
pr = []
for run, x in d.groupby('run'):
    k = x['k'].iloc[0]
    rr = rolls[rolls['run'] == run]
    comp = x[x['ck'] == 'compound']
    cord = x[x['ck'] == 'cord / fabric']
    m_comp = sum(q / kgpm.get((k, c), np.nan) for c, q in comp.groupby('consumed_item')['qty'].sum().items())
    m_cord = sum(q / kgpm.get((k, c), np.nan) for c, q in cord.groupby('consumed_item')['qty'].sum().items())
    mins = (rr['e'].max() - rr['s'].min()).total_seconds() / 60
    metres = m_comp if m_comp == m_comp else m_cord
    rmin = metres / spd[k] if k in spd and metres == metres else np.nan
    extra = sorted(set(x['consumed_item']) - bom_children.get(k, set()))
    pr.append({'source': 'PLANT', 'run_no': int(run), 'item': x['item'].iloc[0], 'item_code': k, 'class': cls.get(k, '?'),
               'wire / dip': wire.get(k, ''), 'start': rr['s'].min(), 'end': rr['e'].max(), 'minutes': round(mins, 1),
               'rolls': len(rr), 'metres (from compound kg / BOM kg per m)': round(m_comp, 0) if m_comp == m_comp else None,
               'metres (from cord / fabric kg / BOM kg per m)': round(m_cord, 0) if m_cord == m_cord else None,
               'metres used': round(metres, 0), 'speed m/min (actual)': round(metres / mins, 1) if mins else None,
               'routing speed (m/min, eff 1)': round(spd.get(k, np.nan), 2),
               'routing minutes for these metres': round(rmin, 1), 'actual / routing time %': round(mins / rmin * 100) if rmin == rmin else None,
               'spool / dip length m': dipl.get(k), 'metres per roll': round(metres / len(rr), 0),
               'inputs not in BOM': ', '.join(extra)})
PR = pd.DataFrame(pr)
# compound age at use (consumption id F01/M02yyyymmddhhmmss_n = mixing batch booking time)
ages = []
for r in d[d['ck'] == 'compound'].itertuples():
    m_ = re.match(r'^[A-Z]\d{2}(\d{14})_', str(r.consumption_id))
    if m_:
        t = pd.to_datetime(m_.group(1), format='%Y%m%d%H%M%S', errors='coerce')
        if pd.notna(t):
            ages.append({'item': r.k, 'compound': r.consumed_item, 'compound_batch': r.consumption_id, 'batch_time': t,
                         'used_at': r.start_time, 'age_h': (r.start_time - t).total_seconds() / 3600,
                         'min_aging_h': mna.get(r.consumed_item), 'max_aging_h': mxa.get(r.consumed_item)})
AG = pd.DataFrame(ages)

# ---------------- v72 runs on the same plant day
L9 = pd.read_excel(f'latest_schedule/{V}/FOURRC_RULES_{V}.xlsx', '901 lots')
L9['item'] = L9['item'].astype(str).str.strip()
L9 = L9.sort_values('scheduled_start').reset_index(drop=True)
L9['run'] = (L9['item'] != L9['item'].shift()).cumsum()
vr = []
for run, x in L9.groupby('run'):
    k = x['item'].iloc[0]
    mins = (x['scheduled_finish'].max() - x['scheduled_start'].min()).total_seconds() / 60
    vr.append({'source': V, 'run_no': int(run), 'item': k, 'item_code': k, 'class': cls.get(k, '?'), 'wire / dip': wire.get(k, ''),
               'start': x['scheduled_start'].min(), 'end': x['scheduled_finish'].max(), 'minutes': round(mins, 1),
               'rolls': len(x), 'metres used': round(x['metres'].sum(), 0),
               'speed m/min (actual)': round(x['metres'].sum() / x['run_min_sched'].sum(), 1),
               'routing speed (m/min, eff 1)': round(spd.get(k, np.nan), 2),
               'routing minutes for these metres': round(x['run_min_routing'].sum(), 1),
               'actual / routing time %': round(x['run_min_sched'].sum() / x['run_min_routing'].sum() * 100),
               'spool / dip length m': dipl.get(k), 'metres per roll': round(x['metres'].mean(), 0)})
VR_all = pd.DataFrame(vr)
VR = VR_all[(VR_all['start'] >= DAY0) & (VR_all['start'] < DAY1)].reset_index(drop=True)
D9 = pd.Timestamp('2026-10-09 07:00')
VR9 = VR_all[(VR_all['start'] >= D9) & (VR_all['start'] < D9 + pd.Timedelta(days=1))].reset_index(drop=True)


# ---------------- rule checks (same functions for both)
def checks(R, src):
    R = R.sort_values('start').reset_index(drop=True)
    out = []

    def add(rule, what, occ, ok, note=''):
        out.append({'rule': rule, 'check': what, 'source': src, 'occurrences': int(occ), 'followed': int(ok),
                    'compliance %': round(100 * ok / occ, 1) if occ else None, 'note': note})
    # 4RC-7 one go per item per day
    g = R.groupby('item_code')['run_no'].nunique()
    add('4RC-7 one go per day', 'item made in ONE run in the plant day', len(g), int((g <= 1).sum()),
        ', '.join(f'{k} x{v}' for k, v in g[g > 1].items()))
    # 4RC-5 creel 240 min and rule 3 fabric between cord changes
    st = R[R['class'] == 'steel'].reset_index(drop=True)
    ncc = okc = ok3 = 0
    notes = []
    for i in range(1, len(st)):
        a, b = st.loc[i - 1], st.loc[i]
        if wire.get(a['item_code']) != wire.get(b['item_code']):
            ncc += 1
            gap = (b['start'] - a['end']).total_seconds() / 60
            fab = R[(R['start'] >= a['end']) & (R['start'] < b['start']) & (R['class'] == 'fabric')]
            okc += gap >= CREEL_MIN - 1
            ok3 += len(fab) > 0
            notes.append(f"{a['item_code']}->{b['item_code']}: {gap:.0f} min, fabric {', '.join(fab['item_code']) or 'none'}")
    add('4RC-5 one creel', f'cord change: next steel >= {CREEL_MIN:.0f} min after the old cord', ncc, okc, '; '.join(notes))
    add('4RC-9 rule 3', 'fabric run between steel cord changes', ncc, ok3)
    # wire campaigns: same wire not split by another wire within the day
    wl = R['wire / dip'].tolist()
    loads = [w for i, w in enumerate(wl) if i == 0 or w != wl[i - 1]]
    rep = pd.Series(loads).value_counts()
    add('4RC v42 wire campaign', 'each wire / dip loaded once in the day (same wire together)', len(rep), int((rep <= 1).sum()),
        ', '.join(f'{w} x{n}' for w, n in rep[rep > 1].items()))
    # 4RC-4 run time = metres / routing speed
    ok4 = R['actual / routing time %'].between(95, 105)
    add('4RC-4 routing line speed', 'run time = metres / routing speed (eff 1) +-5 %', len(R), int(ok4.sum()),
        '; '.join(f"{r['item_code']} {r['actual / routing time %']}%" for _, r in R[~ok4].iterrows()))
    # 4RC-6 roll <= one spool / dip
    ok6 = R['metres per roll'] <= R['spool / dip length m'] + 1
    add('4RC-6 roll lengths', 'average roll <= one spool / dip length', len(R), int(ok6.sum()))
    # 4RC-8 spool / dip consumed: a wire load ends on whole spools / dips
    R['load'] = (R['wire / dip'] != R['wire / dip'].shift()).cumsum()
    lo = R.groupby('load').agg(w=('wire / dip', 'first'), m=('metres used', 'sum'), D=('spool / dip length m', 'first'),
                               items=('item_code', lambda s: '+'.join(dict.fromkeys(s))))
    lo['spools'] = lo['m'] / lo['D']
    lo['whole'] = (lo['spools'] - lo['spools'].round()).abs() <= 0.05
    if src == 'PLANT':   # first / last load may continue from / into days outside the data
        lo = lo.iloc[1:-1]
    add('4RC-8 spool / dip used up', 'wire / dip load ends on whole spools / dips (+-5 %)' +
        (' -- first and last load of the day not judged' if src == 'PLANT' else ''), len(lo), int(lo['whole'].sum()),
        '; '.join(f"{r.items}: {r.spools:.2f} spools / dips" for r in lo.itertuples() if not r.whole))
    # changeover (matrix) visible between runs
    nco = okco = 0
    for i in range(1, len(R)):
        a, b = R.loc[i - 1], R.loc[i]
        t = co.get((ALIAS.get(a['item'], a['item']).replace('-MOTHERROLL', ''), b['item'].replace('-MOTHERROLL', '')))
        if t is None:
            t = co.get((a['item_code'], b['item_code']))
        if t is not None and t == t and t > 0:
            nco += 1
            okco += (b['start'] - a['end']).total_seconds() / 60 >= t - 0.5
    add('4RC-9 least changeover', 'gap between runs >= matrix changeover time', nco, okco,
        'plant bookings run back to back (no changeover time recorded)' if src == 'PLANT' else '')
    # 4RC-10 / idle
    gaps = [(R.loc[i, 'start'] - R.loc[i - 1, 'end']).total_seconds() / 60 for i in range(1, len(R))]
    add('4RC-10 no idle inside a load', 'no idle > 5 min between runs of the same wire / dip',
        sum(1 for i in range(1, len(R)) if R.loc[i, 'wire / dip'] == R.loc[i - 1, 'wire / dip']),
        sum(1 for i in range(1, len(R)) if R.loc[i, 'wire / dip'] == R.loc[i - 1, 'wire / dip'] and gaps[i - 1] <= 5))
    # capacity
    busy = R['minutes'].sum()
    add('4RC-3 capacity', 'run minutes in the plant day <= 1440', 1, int(busy <= 1440), f'{busy:.0f} min')
    return out


CHK = pd.DataFrame(checks(PR, 'PLANT') + checks(VR, V + ' 08-Oct') + checks(VR9, V + ' 09-Oct'))
# CAP 66 rule 2 (bypass slitted): plant consumes CAP STRIP 120 on the CAP 66 run
cap = PR[PR['item_code'] == 'CAP 66-MOTHERROLL']
CHK = pd.concat([CHK, pd.DataFrame([
    {'rule': '4RC-2 CAP 66 from fabric + B616M', 'check': 'CAP 66 run uses only its BOM inputs (fabric 000000001032281000 + B616M)',
     'source': 'PLANT', 'occurrences': len(cap), 'followed': int((cap['inputs not in BOM'] == '').sum()),
     'compliance %': None, 'note': 'plant also consumes ' + ', '.join(sorted(set(cap['inputs not in BOM']))) if len(cap) else ''},
    {'rule': '4RC-2 CAP 66 from fabric + B616M', 'check': 'no slitted-material lots', 'source': V, 'occurrences': 1,
     'followed': 1, 'compliance %': 100.0, 'note': 'FOURRC_RULES check 4RC-2'}])], ignore_index=True)
# day-slices of v72 cut runs at 07:00, so capacity / idle / changeover are taken from v72's own full-horizon checker
CHK = CHK[~(CHK['source'].str.startswith(V) & CHK['rule'].isin(['4RC-3 capacity', '4RC-10 no idle inside a load', '4RC-9 least changeover']))]
OFF = pd.read_excel(f'latest_schedule/{V}/FOURRC_RULES_{V}.xlsx', 'Summary')
CHK['compliance %'] = (100 * CHK['followed'] / CHK['occurrences'].replace(0, np.nan)).round(1)

# ---------------- summary measures
def meas(R):
    return {'runs': len(R), 'items': R['item_code'].nunique(), 'rolls': int(R['rolls'].sum()),
            'run minutes': round(R['minutes'].sum()), 'steel metres': int(R.loc[R['class'] == 'steel', 'metres used'].sum()),
            'fabric metres': int(R.loc[R['class'] == 'fabric', 'metres used'].sum()),
            'steel cord changes': int(sum(1 for i in range(1, len(R[R['class'] == 'steel']))
                                          if wire.get(R[R['class'] == 'steel'].iloc[i - 1]['item_code'])
                                          != wire.get(R[R['class'] == 'steel'].iloc[i]['item_code']))),
            'utilisation % of 24 h': round(R['minutes'].sum() / 14.4)}


VR_all['pday'] = (VR_all['start'] - pd.Timedelta(hours=7)).dt.date
steady = [meas(g) for dd, g in VR_all.groupby('pday') if pd.Timestamp('2026-10-09').date() <= dd <= pd.Timestamp('2026-10-15').date()]
SUM = pd.DataFrame([{'measure': k, 'PLANT 08-Oct': v, f'{V} 08-Oct': meas(VR)[k], f'{V} 09-Oct': meas(VR9)[k],
                     f'{V} avg 09-15 Oct': round(np.mean([m_[k] for m_ in steady]), 1)} for k, v in meas(PR).items()])
SUM.loc[SUM['measure'] == 'rolls', 'measure'] = 'rolls (plant: wound rolls; v72: lots = one spool / dip each)'
DAYS = pd.DataFrame([{'plant_day': dd, **meas(g)} for dd, g in VR_all.groupby('pday')])

DATA = []
for r in PR.to_dict('records'):
    if r['item'] != r['item_code']:
        DATA.append({'item': r['item'], 'check': 'item code', 'finding': f"plant '{r['item']}' = scheduler '{r['item_code']}'"})
    if r['inputs not in BOM']:
        DATA.append({'item': r['item_code'], 'check': 'BOM', 'finding': f"consumes inputs not in the BOM: {r['inputs not in BOM']}"})
    mc, mf = r['metres (from compound kg / BOM kg per m)'], r['metres (from cord / fabric kg / BOM kg per m)']
    if mc and mf and abs(mc / mf - 1) > 0.05:
        DATA.append({'item': r['item_code'], 'check': 'BOM ratio', 'finding': f'metres from compound {mc:,.0f} vs from cord / fabric {mf:,.0f} ({(mc / mf - 1) * 100:+.0f} %): compound : cord ratio differs from the BOM'})
    DATA.append({'item': r['item_code'], 'check': 'line speed', 'finding': f"actual {r['speed m/min (actual)']} m/min vs routing {r['routing speed (m/min, eff 1)']} m/min (run took {r['actual / routing time %']} % of routing time)"})
if len(AG):
    for (it_, cp), g in AG.groupby(['item', 'compound']):
        lo_ = g['age_h'].min()
        hi_ = g['age_h'].max()
        mn_, mx_ = g['min_aging_h'].iloc[0], g['max_aging_h'].iloc[0]
        f = []
        if mn_ == mn_ and lo_ < mn_:
            f.append(f'youngest {lo_:.1f} h < min aging {mn_:g} h')
        if mx_ == mx_ and hi_ > mx_:
            f.append(f'oldest {hi_:.1f} h > max aging {mx_:g} h')
        DATA.append({'item': it_, 'check': f'compound {cp} age at use',
                     'finding': (f'{lo_:.1f}-{hi_:.1f} h (min {mn_:g} h / max {mx_:g} h)' + (' -- ' + '; '.join(f) if f else ' -- OK'))})
DATA = pd.DataFrame(DATA)

hdr, bold = PatternFill('solid', fgColor='FF1F3864'), Font(name='Arial', bold=True, color='FFFFFFFF')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    for name, df in [('Summary', SUM), ('Rule checks', CHK), (f'{V} official 4RC checks', OFF), (f'{V} 4RC by day', DAYS), ('Plant runs', PR), (f'{V} runs (08-Oct)', VR),
                     ('Data checks (plant)', DATA), ('Compound age (plant)', AG)]:
        df.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical='center')
        for j, c in enumerate(df.columns, 1):
            wd = min(max(10, int(df[c].astype(str).str.len().head(300).max() if len(df) else 10) + 2, len(str(c)) + 2), 70)
            ws.column_dimensions[get_column_letter(j)].width = wd
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'
pd.set_option('display.width', 260)
pd.set_option('display.max_colwidth', 110)
print(SUM.to_string(index=False))
print(CHK[['rule', 'source', 'occurrences', 'followed', 'compliance %', 'note']].to_string(index=False))
print(PR[['item_code', 'class', 'minutes', 'rolls', 'metres (from compound kg / BOM kg per m)', 'metres (from cord / fabric kg / BOM kg per m)',
          'speed m/min (actual)', 'routing speed (m/min, eff 1)', 'actual / routing time %', 'inputs not in BOM']].to_string(index=False))
print(VR9[['item_code', 'class', 'start', 'minutes', 'rolls', 'metres used', 'speed m/min (actual)', 'actual / routing time %']].to_string(index=False))
print(DATA.to_string(index=False))
print('written', OUT)
