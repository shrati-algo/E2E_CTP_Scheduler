#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""compare_plant_v47.py -- plant mixing schedule (mixing_schedule.xlsx) vs the v47 schedule on the same day;
mixing-rule compliance of both; plant data checks against the masters v47 used (routing, BOM, aging).

    python latest_schedule/plant-schedule/compare_plant_v47.py
    -> latest_schedule/plant-schedule/PLANT_vs_v47_MIXING_COMPARISON.xlsx
"""
import os
import re

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
IN47 = os.path.join(ROOT, "latest_schedule", "v47", "run_output", "PHASEWISE", "00_INPUTS")
OUT = os.path.join(HERE, "PLANT_vs_v47_MIXING_COMPARISON.xlsx")

PMAP = {"Mixer270F": "F270", "Mixer310F": "K310", "Mixer440": "F440", "Mixer270": "K270M",
        "Mixer431": "K430-1", "Mixer432": "K430-2"}
ORDER = ["F270", "K310", "F440", "K270M", "K430-1", "K430-2"]
FINAL_BUTYL, FINAL_FRIENDLY = {"30G", "IL3", "IL5"}, {"B1", "B2", "B99"}
MASTER_BUTYL, MASTER_FRIENDLY = {"29MG", "ML405", "ML605"}, {"MB100", "MB200", "XMB98"}
IL = {"IL3", "IL5"}
LONG = {"B176", "B186", "B333", "B350", "B370", "B458", "B460", "B616M", "B96", "EFA", "S36", "IL3", "IL5"}
EXEMPT = {"B1", "B2", "B99", "MB100", "MB200", "XMB98"}
STICKY = {"JT444", "JT466"}
CLEAN_S, CO_S, TR100_S = 120, 60, 300
H7 = pd.Timedelta(hours=7)


def fam(it):                       # 30GWA etc. -> 30G family for the butyl rules
    it = str(it).strip()
    return "30G" if it.startswith("30G") else it


def n0(x):
    return str(x).strip().lstrip("0")


# ---------------------------------------------------------------- masters used by v47
it_m = pd.read_csv(os.path.join(IN47, "itemtype_master.csv"), dtype=str)
ITYPE = dict(zip(it_m.ItemCode.str.strip(), it_m.ItemType.str.strip().str.upper()))
rt = pd.read_csv(os.path.join(IN47, "routing.csv"), low_memory=False)
rt = rt[rt.operation_name.astype(str).str.upper().str.contains("MIXING")]
RT = {}
for r in rt.itertuples():
    ms = {m.strip().strip('"') for m in str(r.machines).split(",") if m.strip() and m.strip() != "nan"}
    ms = {"K310" if m == "310" else m for m in ms}
    p = str(r.routed_product).strip()
    e = float(r.efficiency) if r.efficiency == r.efficiency and float(r.efficiency) > 0 else 1.0
    sec = float(r.proc_time) * (60 if str(r.proc_time_UOM).upper().startswith("MIN") else 1)
    RT.setdefault(p, {"machines": set(), "cycle_s": sec / e, "batch_kg": float(r.batch_size or 0),
                      "op": r.operation_name})
    RT[p]["machines"] |= ms
ag = pd.read_csv(os.path.join(IN47, "aging_master.csv"))
UNIT = {"HOURS": 1, "DAYS": 24, "MONTH": 720, "MIN": 1 / 60}
AGE = {}
for r in ag.itertuples():
    try:
        AGE[str(r.ItemCode).strip()] = (float(r.MinAging) * UNIT.get(str(r.MinAgingUnit).upper(), 1),
                                        float(r.MaxAging) * UNIT.get(str(r.MaxAgingUnit).upper(), 1))
    except Exception:
        pass
bom = pd.read_csv(os.path.join(IN47, "bom.csv"), dtype=str, low_memory=False)
bom = bom[bom.Parent.notna()]
bom["pq"] = pd.to_numeric(bom.Parent_qty, errors="coerce")
bom["cq"] = pd.to_numeric(bom.child_quantity, errors="coerce")
bom = bom[(bom.pq > 0) & bom.cq.notna()]
bom["ratio"] = bom.cq / bom.pq
BOM = {}
for p, g in bom.groupby(bom.Parent.str.strip()):
    BOM[p] = g.groupby(g.child.map(n0))["ratio"].median().to_dict()

# ---------------------------------------------------------------- plant data
raw = pd.read_excel(os.path.join(HERE, "mixing_schedule.xlsx"), "in")
raw["mixer"] = raw.machine.map(PMAP)
raw["item"] = raw["item"].astype(str).str.strip()
first_lot = raw.groupby("prod_id").lot_id.first()
pl = raw[raw.lot_id == raw.prod_id.map(first_lot)].copy()          # each batch once (3 duplicated batches)
B = pl.groupby("prod_id").agg(mixer=("mixer", "first"), item=("item", "first"), lot=("lot_id", "first"),
                              shift=("shift", "first"), start=("start_time", "min"), lot_end=("end_time", "max"),
                              kg=("qty", "sum")).reset_index().sort_values(["mixer", "start"])
W0, W1 = B.start.min(), B.start.max()
B["plant_day"] = (B.start - H7).dt.date


def runs_plant(b):
    out = []
    for m, g in b.groupby("mixer"):
        g = g.sort_values("start")
        cur = None
        for r in g.itertuples():
            if cur and cur["item"] == r.item:
                cur["starts"].append(r.start)
                cur["n"] += 1
                cur["lots"].add(r.lot)
            else:
                if cur:
                    out.append(cur)
                cur = {"mixer": m, "item": r.item, "starts": [r.start], "n": 1, "lots": {r.lot}}
        if cur:
            out.append(cur)
    for c in out:
        iv = pd.Series(c["starts"]).diff().dt.total_seconds().dropna()
        cyc = float(iv.median()) if len(iv) else RT.get(c["item"], {}).get("cycle_s", 150.0)
        c["cycle_s"] = cyc
        c["start"], c["end"] = c["starts"][0], c["starts"][-1] + pd.Timedelta(seconds=cyc)
    return out


# ---------------------------------------------------------------- v47 on the same window
v = pd.read_excel(os.path.join(ROOT, "latest_schedule", "v47", "MIXING_SCHEDULE_v47.xlsx"), "Mixing schedule")
v["start_time"], v["end_time"] = pd.to_datetime(v.start_time), pd.to_datetime(v.end_time)
v["item"] = v["item"].astype(str).str.strip()


def runs_v47(df):
    out = []
    for m, g in df.groupby("machine"):
        g = g.sort_values("start_time")
        cur = None
        for r in g.itertuples():
            if cur and cur["item"] == r.item:
                cur["end"] = max(cur["end"], r.end_time)
                cur["n"] += r.batches
            else:
                if cur:
                    out.append(cur)
                cur = {"mixer": m, "item": r.item, "start": r.start_time, "end": r.end_time, "n": r.batches}
        if cur:
            out.append(cur)
    return out


RV_all = runs_v47(v)                                # whole v47 schedule (for gaps across the window edge)
RP = runs_plant(B)


# ---------------------------------------------------------------- mixing rules on runs
def rules(R, who, w0, w1):
    viol, occ = [], {k: 0 for k in ["M1", "M2", "M3", "M4", "M5", "M7", "M8", "CO"]}
    inw = lambda c: w0 <= c["start"] <= w1
    by_m = {}
    for c in R:
        by_m.setdefault(c["mixer"], []).append(c)
    for m, L in by_m.items():
        L.sort(key=lambda c: c["start"])
        for i, c in enumerate(L):
            if not inw(c):
                continue
            p = L[i - 1] if i else None
            n = L[i + 1] if i + 1 < len(L) else None
            f, pf, nf = fam(c["item"]), fam(p["item"]) if p else None, fam(n["item"]) if n else None
            gb = (c["start"] - p["end"]).total_seconds() if p else None
            ga = (n["start"] - c["end"]).total_seconds() if n else None
            base = {"schedule": who, "mixer": m, "compound": c["item"], "run_start": c["start"],
                    "previous": p["item"] if p else "", "next": n["item"] if n else "",
                    "gap_before_s": None if gb is None else round(gb), "gap_after_s": None if ga is None else round(ga)}
            if f in FINAL_BUTYL and p is not None:            # first run in the data: predecessor unknown
                occ["M1"] += 1
                pp = L[i - 2] if i > 1 else None
                ok = (pf in FINAL_FRIENDLY) or (pf in IL and f in IL and pf != f and pp is not None and fam(pp["item"]) in FINAL_FRIENDLY)
                if not ok:
                    viol.append({**base, "rule": "Final Mixer – Before Butyl", "detail": f"{c['item']} preceded by {pf or 'nothing'} (needs B1 / B2 / B99)"})
            if (f in FINAL_BUTYL or f in MASTER_BUTYL) and n is not None and not (f in IL and nf in IL and f != nf):
                occ["M2"] += 1
                if ga is not None and ga < CLEAN_S - 1:
                    viol.append({**base, "rule": "Final Mixer – After Butyl", "detail": f"next {n['item']} {ga:.0f} s after {c['item']} (needs >= 120 s clean-out)"})
            if f in MASTER_BUTYL and p is not None:
                occ["M3"] += 1
                bad = ([] if pf in MASTER_FRIENDLY else [f"before: {pf or 'nothing'}"]) + ([] if (n is None or nf in MASTER_FRIENDLY) else [f"after: {nf}"])
                if bad:
                    viol.append({**base, "rule": "Master Mixer – Butyl Sequence", "detail": "; ".join(bad) + " (needs MB100 / MB200 / XMB98)"})
            if m == "K430-2" and ITYPE.get(c["item"], "") == "FINAL COMPOUND":
                occ["M4"] += 1
                bad = ([] if (gb is None or gb >= CLEAN_S - 1) else [f"before {gb:.0f} s"]) + ([] if (ga is None or ga >= CLEAN_S - 1) else [f"after {ga:.0f} s"])
                if bad:
                    viol.append({**base, "rule": "430-2 Master Mixer – Final Compound", "detail": "; ".join(bad) + " (needs >= 120 s)"})
            if n is not None and nf != f:
                occ["CO"] += 1
                need = TR100_S if {ITYPE.get(c["item"], ""), ITYPE.get(n["item"], "")} == {"MASTER COMPOUND", "FINAL COMPOUND"} else CO_S
                if ga is not None and ga < need - 1:
                    viol.append({**base, "rule": "Changeover time (BRM-10 / TR100)", "detail": f"{c['item']} -> {n['item']} gap {ga:.0f} s (needs >= {need} s)" + (" [INDICATIVE: plant batch end estimated]" if who == "plant" else "")})
    # M5 IL3 + IL5 one go a day
    days = {}
    for c in R:
        if c["item"] in IL and w0 <= c["start"] <= w1:
            days.setdefault((c["start"] - H7).date(), []).append(c)
    for d, L in days.items():
        occ["M5"] += 1
        L.sort(key=lambda c: c["start"])
        camps, cur = [], [L[0]]
        for c in L[1:]:
            if c["mixer"] == cur[-1]["mixer"] and (c["start"] - cur[-1]["end"]).total_seconds() < 1800:
                cur.append(c)
            else:
                camps.append(cur)
                cur = [c]
        camps.append(cur)
        items = [x["item"] for x in camps[0]]
        if len(camps) > 1 or len(items) != len(set(items)):
            viol.append({"schedule": who, "mixer": ",".join(sorted({c["mixer"] for c in L})), "compound": "IL3 + IL5",
                         "run_start": L[0]["start"], "rule": "IL3 and IL5 together in one go a day",
                         "detail": f"{len(camps)} IL campaigns that plant day: " + " | ".join("->".join(x["item"] for x in cp) for cp in camps)})
    # M7 one mixer per day
    dm = {}
    for c in R:
        if w0 <= c["start"] <= w1:
            dm.setdefault((c["item"], (c["start"] - H7).date()), set()).add(c["mixer"])
    for (it, d), ms in dm.items():
        occ["M7"] += 1
        if len(ms) > 1:
            viol.append({"schedule": who, "mixer": ", ".join(sorted(ms)), "compound": it, "run_start": None,
                         "rule": "One compound in one mixer per day", "detail": f"plant day {d}: on {len(ms)} mixers"})
    # M8 repeat gap
    bi = {}
    for c in R:
        bi.setdefault(c["item"], []).append(c)
    for it, L in bi.items():
        if it in EXEMPT:
            continue
        L.sort(key=lambda c: c["start"])
        need = 24 if it in LONG else 15
        for a, b in zip(L, L[1:]):
            if not (w0 <= b["start"] <= w1):
                continue
            occ["M8"] += 1
            gh = (b["start"] - a["start"]).total_seconds() / 3600
            if gh < need - 1 / 60:
                viol.append({"schedule": who, "mixer": f"{a['mixer']} -> {b['mixer']}", "compound": it, "run_start": b["start"],
                             "rule": "Same compound not repeated within 15 h (24 h IL / 4RC)",
                             "detail": f"run {b['start']:%d-%b %H:%M} starts {gh:.1f} h after the run {a['start']:%d-%b %H:%M} (needs >= {need} h)"})
    return pd.DataFrame(viol), occ


VP, OP = rules(RP, "plant", W0, W1)
VV, OV = rules(RV_all, "v47", W0, W1)
RULES = [("M1", "Final Mixer – Before Butyl"), ("M2", "Final Mixer – After Butyl"), ("M3", "Master Mixer – Butyl Sequence"),
         ("M4", "430-2 Master Mixer – Final Compound"), ("M5", "IL3 and IL5 together in one go a day"),
         ("M7", "One compound in one mixer per day"), ("M8", "Same compound not repeated within 15 h (24 h IL / 4RC)"),
         ("CO", "Changeover time (BRM-10 / TR100)")]
comp = []
for k, nm in RULES:
    row = {"rule": nm}
    for who, V, O in (("plant", VP, OP), ("v47", VV, OV)):
        nv = int((V["rule"] == nm).sum()) if len(V) else 0
        row[f"{who}_checked"], row[f"{who}_violations"] = O[k], nv
        row[f"{who}_compliance_%"] = round(100 * (1 - nv / O[k]), 1) if O[k] else None
    comp.append(row)
comp.append({"rule": "Mixing output in batches", "plant_checked": len(B), "plant_violations": 0, "plant_compliance_%": 100.0,
             "v47_checked": int(((v.start_time >= W0) & (v.start_time <= W1)).sum()), "v47_violations": 0, "v47_compliance_%": 100.0})
COMP = pd.DataFrame(comp)

# ---------------------------------------------------------------- plant data checks vs v47 masters
chk = []
# D1 allowed mixer (routing)
for (m, it), g in B.groupby(["mixer", "item"]):
    al = RT.get(it, {}).get("machines")
    if al is None:
        chk.append({"check": "D1 mixer allowed by routing", "mixer": m, "compound": it, "batches": len(g), "status": "NO DATA",
                    "detail": "compound has no mixing row in the routing used by v47"})
    elif m not in al:
        chk.append({"check": "D1 mixer allowed by routing", "mixer": m, "compound": it, "batches": len(g), "status": "VIOLATION",
                    "detail": f"routing allows {sorted(al)}"})
# D2 batch weight vs routing batch size (+-5 %)
for (m, it), g in B.groupby(["mixer", "item"]):
    bk = RT.get(it, {}).get("batch_kg", 0)
    if bk:
        med = g.kg.median()
        dev = 100 * (med - bk) / bk
        if abs(dev) > 5:
            chk.append({"check": "D2 batch weight vs routing batch_size", "mixer": m, "compound": it, "batches": len(g),
                        "status": "DEVIATION", "detail": f"median batch {med:.1f} kg vs routing {bk:g} kg ({dev:+.1f} %)"})
# D3 recipe vs BOM
for it, g in pl.groupby("item"):
    rec = BOM.get(it)
    if not rec:
        chk.append({"check": "D3 recipe vs BOM", "mixer": ", ".join(sorted(g.mixer.unique())), "compound": it,
                    "batches": g.prod_id.nunique(), "status": "NO DATA", "detail": "no BOM rows for this compound in the v47 BOM"})
        continue
    used = g.groupby(g.consumed_item.map(n0)).qty.sum() / g.prod_id.nunique()
    bkg = g.groupby("prod_id").qty.sum().median()
    extra = sorted(set(used.index) - set(rec))
    miss = sorted(set(rec) - set(used.index))
    devs = []
    for c, r in rec.items():
        if c in used.index and r > 0:
            exp = r * bkg
            d = 100 * (used[c] - exp) / exp
            if abs(d) > 5:
                devs.append(f"{c} {used[c]:.1f} kg vs {exp:.1f} ({d:+.0f} %)")
    if extra or miss or devs:
        chk.append({"check": "D3 recipe vs BOM", "mixer": ", ".join(sorted(g.mixer.unique())), "compound": it,
                    "batches": g.prod_id.nunique(), "status": "DEVIATION",
                    "detail": "; ".join(([f"not in BOM: {extra}"] if extra else []) + ([f"BOM item not consumed: {miss}"] if miss else []) + devs)[:500]})
# D4 aging of consumed compounds (produced time from the consumption id)
ts_re = re.compile(r"(20\d{2})(\d{2})(\d{2})(\d{2})(\d{2})(\d{2})")
age_rows = []
for r in pl.dropna(subset=["consumption_id"]).itertuples():
    ci = n0(r.consumed_item)
    if ITYPE.get(str(r.consumed_item).strip(), ITYPE.get(ci, "")) not in ("MASTER COMPOUND", "FINAL COMPOUND"):
        continue
    mt = ts_re.search(str(r.consumption_id))
    if not mt:
        continue
    try:
        prod = pd.Timestamp(*[int(x) for x in mt.groups()])
    except Exception:
        continue
    lim = AGE.get(str(r.consumed_item).strip()) or AGE.get(ci)
    if not lim:
        continue
    age_h = (r.start_time - prod).total_seconds() / 3600
    st = "OK" if lim[0] - 1e-6 <= age_h <= lim[1] + 1e-6 else ("TOO FRESH (< min aging)" if age_h < lim[0] else "OVER-AGE (> max aging)")
    age_rows.append({"mixer": r.mixer, "compound": r.item, "batch": r.prod_id, "batch_start": r.start_time,
                     "consumed_compound": r.consumed_item, "consumption_id": r.consumption_id, "produced": prod,
                     "age_h": round(age_h, 1), "min_aging_h": lim[0], "max_aging_h": lim[1], "status": st})
AGEDF = pd.DataFrame(age_rows)
if len(AGEDF):
    for (m, it, cc, st), g in AGEDF[AGEDF.status != "OK"].groupby(["mixer", "compound", "consumed_compound", "status"]):
        chk.append({"check": "D4 aging of consumed compound", "mixer": m, "compound": it, "batches": g.batch.nunique(),
                    "status": "VIOLATION", "detail": f"{cc}: {st}; age {g.age_h.min():.1f}-{g.age_h.max():.1f} h vs min {g.min_aging_h.iloc[0]:g} / max {g.max_aging_h.iloc[0]:g} h"})
# D5 run >= 5 batches (BRM-01 MPQ; a run = consecutive batches of one compound on a mixer -- the plant lot_id is a
# sub-lot of ~4 batches, so the run is the comparable unit) / D7 sticky compounds in stacks of 3 (BRM-02)
for c in RP:
    if c["n"] < 5 and c["item"] not in STICKY and c["item"] not in EXEMPT:
        chk.append({"check": "D5 run >= 5 batches (BRM-01 MPQ)", "mixer": c["mixer"], "compound": c["item"], "batches": c["n"],
                    "status": "VIOLATION", "detail": f"run {c['start']:%d-%b %H:%M}: {c['n']} batches"})
    if c["item"] in STICKY and c["n"] % 3:
        chk.append({"check": "D7 sticky compound in stacks of 3 (BRM-02)", "mixer": c["mixer"], "compound": c["item"], "batches": c["n"],
                    "status": "VIOLATION", "detail": f"run {c['start']:%d-%b %H:%M}: {c['n']} batches"})
# D6 cycle time vs routing
for c in RP:
    rc = RT.get(c["item"], {}).get("cycle_s")
    if rc and c["n"] >= 3:
        if c["cycle_s"] < 0.9 * rc:
            chk.append({"check": "D6 batch cycle vs routing", "mixer": c["mixer"], "compound": c["item"], "batches": c["n"],
                        "status": "FASTER THAN ROUTING", "detail": f"run {c['start']:%d-%b %H:%M}: {c['cycle_s']:.0f} s/batch vs routing {rc:.0f} s"})
        elif c["cycle_s"] > 1.5 * rc:
            chk.append({"check": "D6 batch cycle vs routing", "mixer": c["mixer"], "compound": c["item"], "batches": c["n"],
                        "status": "SLOWER THAN ROUTING", "detail": f"run {c['start']:%d-%b %H:%M}: {c['cycle_s']:.0f} s/batch vs routing {rc:.0f} s"})
CHK = pd.DataFrame(chk)

# ---------------------------------------------------------------- comparison tables
vw = v[(v.start_time >= W0) & (v.start_time <= W1)].copy()
RV = [c for c in RV_all if W0 <= c["start"] <= W1]
def mixer_tab(who, b_by_m, runs, kg):
    rows = []
    for m in ORDER:
        rr = [c for c in runs if c["mixer"] == m]
        rows.append({"schedule": who, "mixer": m, "batches": int(b_by_m.get(m, 0)), "runs": len(rr),
                     "compounds": len({c["item"] for c in rr}), "compound_changes": max(0, len(rr) - 1), "kg": round(kg.get(m, 0))})
    return rows
cmpm = pd.DataFrame(mixer_tab("plant", B.groupby("mixer").size(), RP, B.groupby("mixer").kg.sum()) +
                    mixer_tab("v47", vw.groupby("machine").batches.sum(), RV,
                              (vw.batches * vw.item.map(lambda i: RT.get(i, {}).get("batch_kg", 0))).groupby(vw.machine).sum()))
MIX = cmpm.pivot_table(index="mixer", columns="schedule", values=["batches", "runs", "compounds", "compound_changes", "kg"]).reindex(ORDER)
MIX.columns = [f"{a}_{b}" for a, b in MIX.columns]
MIX = MIX.reset_index()
B["mixing"] = B["item"].map(lambda i: "Final" if ITYPE.get(i, "") == "FINAL COMPOUND" else "Master")
vw["mixing"] = vw["item"].map(lambda i: "Final" if ITYPE.get(i, "") == "FINAL COMPOUND" else "Master")
SHIFT = pd.concat([B.groupby(["mixing", "shift"]).size().rename("plant"), vw.groupby(["mixing", "shift"]).batches.sum().rename("v47")], axis=1).fillna(0).astype(int).reset_index()
pc = B.groupby("item").agg(plant_batches=("prod_id", "size"), plant_mixers=("mixer", lambda s: ", ".join(sorted(set(s)))))
vc = vw.groupby("item").agg(v47_batches=("batches", "sum"), v47_mixers=("machine", lambda s: ", ".join(sorted(set(s)))))
CPD = pc.join(vc, how="outer").fillna({"plant_batches": 0, "v47_batches": 0}).reset_index().rename(columns={"index": "compound"})
CPD["type"] = CPD.iloc[:, 0].map(lambda i: ITYPE.get(i, "?"))
CPD["difference (plant - v47)"] = CPD.plant_batches - CPD.v47_batches
CPD["in"] = np.where(CPD.plant_batches.eq(0), "v47 only", np.where(CPD.v47_batches.eq(0), "plant only", "both"))
CPD = CPD.sort_values("difference (plant - v47)")

OVR = pd.DataFrame([
    {"measure": "window (by batch / row start)", "plant": f"{W0:%d-%b %H:%M} -> {W1:%d-%b %H:%M}", "v47": "same window"},
    {"measure": "batches", "plant": len(B), "v47": int(vw.batches.sum())},
    {"measure": "  final / master", "plant": f"{int((B.mixing == 'Final').sum())} / {int((B.mixing == 'Master').sum())}",
     "v47": f"{int(vw[vw.mixing == 'Final'].batches.sum())} / {int(vw[vw.mixing == 'Master'].batches.sum())}"},
    {"measure": "compounds", "plant": B["item"].nunique(), "v47": vw["item"].nunique()},
    {"measure": "runs (consecutive same compound on a mixer)", "plant": len(RP), "v47": len(RV)},
    {"measure": "compound changes on the mixers", "plant": int(MIX.filter(like="compound_changes_plant").sum().iloc[0]),
     "v47": int(MIX.filter(like="compound_changes_v47").sum().iloc[0])},
    {"measure": "compounds in both / plant only / v47 only", "plant": f"{(CPD['in'] == 'both').sum()} / {(CPD['in'] == 'plant only').sum()} / {(CPD['in'] == 'v47 only').sum()}", "v47": ""},
    {"measure": "mixing-rule violations (8 rules)", "plant": int(len(VP)), "v47": int(len(VV))},
    {"measure": "plant data-check findings (routing / BOM / aging / MPQ / cycle)", "plant": int(len(CHK)), "v47": "n.a. (built from these masters)"},
])
NOTES = pd.DataFrame({"note": [
    "Plant: latest_schedule/plant-schedule/mixing_schedule.xlsx (data not changed); 3 MR5 batch ids recorded under two lot ids are counted once here.",
    "v47: latest_schedule/v47/MIXING_SCHEDULE_v47.xlsx, rows starting in the same window as the plant data.",
    "Masters = the files v47 used (v47 run_output/PHASEWISE/00_INPUTS): routing.csv (allowed mixers, batch size, cycle), bom.csv (recipe), aging_master.csv (min / max aging), itemtype_master.csv.",
    "Plant run end = last batch start + the run's median batch interval (the source end_time is the lot end, not the batch end); gaps between runs use that.",
    "Aging check: produced time read from the consumption id (e.g. M0220261007173857_1 = 07-Oct 17:38:57); only consumed MASTER / FINAL compounds with an aging row are checked.",
    "Rule names as in the plant rule list (01-Oct and 08-Oct). Repeat-gap check uses runs on any mixer; the plant file has no data before its first batch, so gaps to earlier runs cannot be checked.",
    "30GWA is treated as a 30G-family butyl compound for the butyl rules (not in the v47 item master)."]})

hdr, bold = PatternFill("solid", fgColor="FF1F3864"), Font(name="Arial", bold=True, color="FFFFFFFF")
red, grn = PatternFill("solid", fgColor="FFFFC7CE"), PatternFill("solid", fgColor="FFC6EFCE")
with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    sheets = [("Overview", OVR), ("Rule compliance", COMP), ("Plant rule violations", VP), ("v47 rule violations", VV),
              ("Plant data checks", CHK), ("Plant aging detail", AGEDF), ("By mixer", MIX), ("By shift", SHIFT),
              ("By compound", CPD), ("Notes", NOTES)]
    for name, df in sheets:
        df = df if len(df) else pd.DataFrame([{"result": "none"}])
        df.to_excel(xw, sheet_name=name, index=False)
        ws = xw.sheets[name]
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical="center")
        for j, c in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, int(df[c].astype(str).str.len().head(300).max()) + 2, len(str(c)) + 2), 80)
            if pd.api.types.is_datetime64_any_dtype(df[c]):
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=j).number_format = "dd-mmm hh:mm"
            if str(c) in ("status",):
                for r, val in enumerate(df[c].astype(str), 2):
                    ws.cell(row=r, column=j).fill = grn if val == "OK" else red
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
print(OVR.to_string(index=False))
print(COMP.to_string(index=False))
print(MIX.to_string(index=False))
print(CHK.groupby(["check", "status"]).size().to_string() if len(CHK) else "no data-check findings")
if len(AGEDF):
    print("aging rows checked", len(AGEDF), AGEDF.status.value_counts().to_dict())
print("wrote", OUT)
