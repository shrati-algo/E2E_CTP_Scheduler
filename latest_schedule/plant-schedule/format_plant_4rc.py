"""Format the plant's 4-roll calender schedule (four_roll_schedule.csv) in the same layout as the plant mixing schedule
(MIXING_SCHEDULE_PLANT_formatted.xlsx). Data is only grouped, never changed. Run from the repository root."""
import os
import re

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

P = 'latest_schedule/plant-schedule/'
SRC = P + 'four_roll_schedule.csv'
OUT = P + 'FOUR_ROLL_SCHEDULE_PLANT_formatted.xlsx'
I = 'latest_schedule/v72/run_output/PHASEWISE/00_INPUTS/'
ALIAS = {'CAP 66': 'CAP 66-MOTHERROLL'}          # plant code -> scheduler item

raw = pd.read_csv(SRC)
d = raw.copy()
d['item_s'] = d['item'].astype(str).str.strip()
d['start_time'] = pd.to_datetime(d['start_time'])
d['end_time'] = pd.to_datetime(d['end_time'])
d['plant_day'] = (d['start_time'] - pd.Timedelta(hours=7)).dt.date

fl = pd.read_csv(I + 'frc_lengths.csv')
fl['item'] = fl['item'].astype(str).str.strip()
cls = dict(zip(fl['item'], fl['class'].astype(str).str.lower()))
wire = dict(zip(fl['item'], fl['wire'].astype(str)))
dip = dict(zip(fl['item'], fl['spool_dip_length_m']))
itp = I + 'itemtype_master.csv' if os.path.exists(I + 'itemtype_master.csv') else 'inputs/itemtype_master.csv'
it = pd.read_csv(itp, dtype=str)
ity = dict(zip(it['ItemCode'].str.strip(), it['ItemType'].str.strip()))
up = pd.read_excel('latest_schedule/v72/CTP_06OCT-16OCT_SCHEDULE_v72.xlsx', 'Upstream schedule', usecols=['line', 'item'])
line = up.groupby(up['item'].astype(str).str.strip())['line'].agg(lambda s: '/'.join(sorted(set(s.dropna())))).to_dict()


def kind(ci):
    ci = str(ci).strip()
    if re.match(r'^\d', ci):
        return 'cord / fabric'
    if ci.upper().startswith('CAP STRIP'):
        return 'cap strip'
    return 'compound'


d['ckind'] = d['consumed_item'].map(kind)
rolls = d.groupby('prod_id').agg(s=('start_time', 'min'), e=('end_time', 'max'), item=('item_s', 'first')).sort_values('s')
rolls['run'] = (rolls['item'] != rolls['item'].shift()).cumsum()
d['run'] = d['prod_id'].map(rolls['run'])

rows = []
for (run, sh, pdy), x in d.groupby(['run', 'shift', 'plant_day'], sort=False):
    item = x['item_s'].iloc[0]
    k = ALIAS.get(item, item)
    cons = x.groupby('consumed_item', sort=False)['qty'].sum()
    rr = rolls.loc[x['prod_id'].unique()]
    rows.append({
        'line': line.get(k, ''), 'plant_day': pdy, 'date': pd.to_datetime(x['date'].iloc[0]).date(), 'shift': sh,
        'process': {'steel': 'Steel calendering', 'fabric': 'Fabric calendering'}.get(cls.get(k, ''), '(class not in CTP_Master_FRC)'),
        'machine': '4RC', 'machine_code_v72': '901', 'item': item, 'item_code_v72': k,
        'item_type': ity.get(k, '(not in item master)'),
        'FRC_wire / dip': wire.get(k, ''), 'spool_dip_length_m': dip.get(k, ''), 'run_no': int(run),
        'start_time': rr['s'].min(), 'end_time': rr['e'].max(),
        'run_min': round((rr['e'].max() - rr['s'].min()).total_seconds() / 60, 1),
        'rolls': x['prod_id'].nunique(), 'UOM': 'ROLL',
        'cord_fabric_consumed (as received)': round(x.loc[x['ckind'] == 'cord / fabric', 'qty'].sum(), 2),
        'compound_consumed_kg': round(x.loc[x['ckind'] == 'compound', 'qty'].sum(), 2),
        'cap_strip_consumed (as received)': round(x.loc[x['ckind'] == 'cap strip', 'qty'].sum(), 2),
        'roll_ids': ';'.join(dict.fromkeys(x['prod_id'].astype(str))),
        'consumed_items (qty)': '; '.join(f'{a} ({b:g})' for a, b in cons.items()),
        'consumption_id': ';'.join(dict.fromkeys(x['consumption_id'].dropna().astype(str))),
        'source_rows': len(x)})
m = pd.DataFrame(rows).sort_values('start_time').reset_index(drop=True)

# busy time = union of roll start -> end intervals (overlapping bookings counted once)
tot, cs, ce = pd.Timedelta(0), None, None
for r in rolls.itertuples():
    if cs is None:
        cs, ce = r.s, r.e
    elif r.s <= ce:
        ce = max(ce, r.e)
    else:
        tot += ce - cs
        cs, ce = r.s, r.e
tot += ce - cs
busy = tot.total_seconds() / 3600

summ = m.groupby(['process', 'machine', 'machine_code_v72']).agg(
    rows=('run_no', 'size'), runs=('run_no', 'nunique'), items=('item', 'nunique'), rolls=('rolls', 'sum'),
    run_min=('run_min', 'sum'), cord_fabric_consumed=('cord_fabric_consumed (as received)', 'sum'),
    compound_kg=('compound_consumed_kg', 'sum'), cap_strip=('cap_strip_consumed (as received)', 'sum')).reset_index()
summ = pd.concat([summ, pd.DataFrame([{
    'process': 'ALL', 'machine': '4RC', 'machine_code_v72': '901', 'rows': len(m), 'runs': m['run_no'].nunique(),
    'items': m['item'].nunique(), 'rolls': int(m['rolls'].sum()), 'run_min': round(m['run_min'].sum(), 1),
    'cord_fabric_consumed': round(m['cord_fabric_consumed (as received)'].sum(), 2),
    'compound_kg': round(m['compound_consumed_kg'].sum(), 2),
    'cap_strip': round(m['cap_strip_consumed (as received)'].sum(), 2)}])], ignore_index=True)
summ['busy_h (roll start->end, overlaps merged)'] = [None] * (len(summ) - 1) + [round(busy, 1)]

by_shift = m.pivot_table(index=['process', 'item'], columns='shift', values='rolls', aggfunc='sum', fill_value=0)
by_shift['day'] = by_shift.sum(axis=1)
by_day = m.pivot_table(index=['process', 'item'], columns='plant_day', values='rolls', aggfunc='sum', fill_value=0)
by_day['total'] = by_day.sum(axis=1)

chk = []
for r in m.itertuples():
    if r.process.startswith('('):
        chk.append({'run_no': r.run_no, 'item': r.item, 'issue': 'item not in CTP_Master_FRC (no steel / fabric class, no wire / dip length)'})
    if r.item_type.startswith('('):
        chk.append({'run_no': r.run_no, 'item': r.item, 'issue': 'item not in the item master'})
for a, b in ALIAS.items():
    if a in set(m['item']):
        chk.append({'run_no': '', 'item': a, 'issue': f'plant code "{a}" = scheduler item "{b}" (name differs)'})
st = rolls.sort_values('s')
ov = int((st['s'] < st['e'].shift()).sum())
if ov:
    chk.append({'run_no': '', 'item': '', 'issue': f'{ov} rolls start before the previous roll ends (overlapping booking times)'})
chk = pd.DataFrame(chk if chk else [{'run_no': '', 'item': '', 'issue': 'no issues found'}]).drop_duplicates()

note = pd.DataFrame({'note': [
    f'Source: {SRC} ({len(raw):,} rows) -- data NOT changed; the original rows are in sheet "Source data (as received)".',
    '4RC schedule = the source rows grouped by production run (consecutive rolls of the same item on the 4-roll calender) and '
    'shift, same layout as the plant mixing schedule: rolls = distinct prod_id, start = first / end = last time of those '
    'rolls, consumed quantities summed per input.',
    'Inputs split by type: cord / fabric = consumed codes starting with a digit (steel cord 1036S..., fabric 0000000010...), '
    'compound = B... codes (kg), cap strip = CAP STRIP codes; quantities as received (the file gives no unit for cord / fabric).',
    'process (steel / fabric), FRC wire / dip and spool / dip length from CTP_Master_FRC (v72 inputs frc_lengths.csv); item '
    'type from the item master; line (PCR / TBR) from the v72 schedule; plant "CAP 66" = scheduler "CAP 66-MOTHERROLL".',
    f'Data window: {d["start_time"].min():%d-%b %H:%M} -> {d["end_time"].max():%d-%b %H:%M}; busy h = union of roll '
    'start -> end intervals.']})

hdr, bold = PatternFill('solid', fgColor='FF1F3864'), Font(name='Arial', bold=True, color='FFFFFFFF')
with pd.ExcelWriter(OUT, engine='openpyxl') as xw:
    sheets = [('Summary', summ), ('Rolls by day', by_day.reset_index()), ('Rolls by shift', by_shift.reset_index()),
              ('4RC schedule', m), ('Steel runs', m[m['process'] == 'Steel calendering']),
              ('Fabric runs', m[m['process'] == 'Fabric calendering']),
              ('Source data (as received)', raw), ('Notes', note), ('Data checks', chk)]
    for name, df in sheets:
        df.to_excel(xw, sheet_name=name[:31], index=False)
        ws = xw.sheets[name[:31]]
        for c in ws[1]:
            c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical='center')
        for j, c in enumerate(df.columns, 1):
            wd = min(max(10, int(df[c].astype(str).str.len().head(300).max() if len(df) else 10) + 2, len(str(c)) + 2), 60)
            ws.column_dimensions[get_column_letter(j)].width = wd
            if str(c) in ('start_time', 'end_time'):
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=j).number_format = 'dd-mmm-yy hh:mm:ss'
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'

print('source rows', len(raw), '| rolls', raw['prod_id'].nunique(), '=', int(m['rolls'].sum()),
      '| qty', round(raw['qty'].sum(), 2), '=', round(m['cord_fabric_consumed (as received)'].sum()
                                                    + m['compound_consumed_kg'].sum() + m['cap_strip_consumed (as received)'].sum(), 2))
print(summ.to_string(index=False))
print(chk.to_string(index=False))
print(by_shift.to_string())
print(m[['shift', 'process', 'item', 'start_time', 'end_time', 'run_min', 'rolls', 'compound_consumed_kg']].to_string(index=False))
print('written', OUT)
