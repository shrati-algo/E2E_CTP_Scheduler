#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""data_issues_4rc.py -- data problems in the two 4-roll-calender (901) masters and how the scheduler uses them.

    python data_issues_4rc.py v41   -> latest_schedule/v41/DATA_ISSUES_4RC_v41.xlsx

Reads the live downloads of the run (db_pipeline/raw): JK_CTP_PPDS.CTP_Master_ChangeoverMatrix (machine '4 RC Calendar'),
JK_CTP_PPDS.CTP_Master_FRC (is_current = 1), JK_CTP_PPDS.CTP_Master_Routing (machine 901), and the v41 901 lots.
"""
import os
import re
import sys

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
RAW = os.path.join(HERE, "db_pipeline", "raw")
n = lambda x: re.sub(r"[\s\-]", "", str(x).upper()).replace("MOTHERROLL", "")
ALIAS = {"MONOCHAFER": "MONOCHAFFER"}
k = lambda x: ALIAS.get(n(x), n(x))


def main(v):
    vdir = os.path.join(HERE, "latest_schedule", v)
    m = pd.read_csv(os.path.join(RAW, "JK_CTP_PPDS__CTP_Master_ChangeoverMatrix.csv"), low_memory=False)
    m = m[m["machine"].astype(str).str.strip() == "4 RC Calendar"].copy()
    f = pd.read_csv(os.path.join(RAW, "JK_CTP_PPDS__CTP_Master_FRC.csv"))
    f = f[pd.to_numeric(f.get("is_current", 1), errors="coerce").fillna(1) == 1].copy()
    r = pd.read_csv(os.path.join(RAW, "JK_CTP_PPDS__CTP_Master_Routing.csv"), dtype=str, low_memory=False)
    r = r[r["machines"].astype(str).str.contains(r'"901"|^901$', regex=True)]
    L = pd.read_excel(os.path.join(vdir, f"FOURRC_RULES_{v}.xlsx"), "901 lots")
    L["scheduled_start"] = pd.to_datetime(L["scheduled_start"], errors="coerce")
    L["scheduled_finish"] = pd.to_datetime(L["scheduled_finish"], errors="coerce")
    used = sorted(set(L["item"].astype(str)))

    m["a"], m["b"] = m["from_MaterialCode_O"].map(k), m["to_MaterialCode_O"].map(k)
    m["type"] = m["changeover_type"].astype(str).str.strip()
    f["key"] = f["item_code"].map(k)
    fw = dict(zip(f["key"], f["wire_type"].astype(str).str.strip()))
    ft = dict(zip(f["key"], f["item_type"].astype(str).str.strip().str.lower()))
    rk = {k(x): x for x in r["routed_product"].astype(str).str.strip()}
    mitems = set(m["a"]) | set(m["b"])
    issues = []

    def add(table, col, issue, count, example, effect, action):
        issues.append({"table": table, "column(s)": col, "issue": issue, "rows / items": count,
                       "example": example, "effect on the schedule": effect, "action needed": action})

    # ---- 1. matrix vs FRC: what is a steel cord change (creel reload)
    st = m[(m["a"] != m["b"]) & m["a"].isin(fw) & m["b"].isin(fw) & (m["a"].map(ft) == "steel") & (m["b"].map(ft) == "steel")].copy()
    st["wire_from"], st["wire_to"] = st["a"].map(fw), st["b"].map(fw)
    st["same_wire_FRC"] = st["wire_from"] == st["wire_to"]
    tl = st["type"].str.lower()
    st["matrix_says"] = tl.map(lambda t: "same cord: compound change, no creel reload" if "compound" in t else
                               ("new cord: creel reload" if "steel to steel changeover" in t else t))
    st["FRC_says"] = st["same_wire_FRC"].map({True: "same wire: same spool, no reload", False: "different wire: new cord, reload"})
    st["disagree"] = (tl.str.contains("compound") & ~st["same_wire_FRC"]) | (tl.str.contains("steel to steel changeover") & st["same_wire_FRC"])
    D = st[st["disagree"]][["from_MaterialCode_O", "to_MaterialCode_O", "changeover_letter", "changeover_time_min", "type",
                            "wire_from", "wire_to", "matrix_says", "FRC_says"]].rename(columns={"type": "changeover_type"})
    n1 = int((tl.str.contains("compound") & ~st["same_wire_FRC"]).sum())
    n2 = int((tl.str.contains("steel to steel changeover") & st["same_wire_FRC"]).sum())
    add("CTP_Master_ChangeoverMatrix vs CTP_Master_FRC", "changeover_type vs wire_type",
        "The two tables disagree on whether a steel -> steel change is a new cord (creel reload)", f"{len(D)} of {len(st)} steel pairs",
        f"{n1} pairs: matrix 'Steel to steel compound change' (25 min, no reload) but FRC wires differ (e.g. CPJ 1218 1036S34500 -> "
        f"CP56542 1036S56500); {n2} pairs: matrix 'Steel to Steel Changeover' (64 min + reload) but FRC same wire (e.g. CPJ 1218 -> CPJ 120, both 1036S34500)",
        "Creel reload (240 min) applied inconsistently; part of the 901 idle (773 min in v41) may not be real",
        "Plant to confirm what wire_type means (cord on the creel?) and correct one of the two tables")
    # pairs that occurred in v41 with their gaps
    s = L[L["class"].astype(str).str.lower() == "steel"].sort_values("scheduled_start").reset_index(drop=True)
    occ = []
    for a, b in zip(s.itertuples(), s.iloc[1:].itertuples()):
        if a.item == b.item:
            continue
        row = st[(st["a"] == k(a.item)) & (st["b"] == k(b.item))]
        occ.append({"from": a.item, "to": b.item, "old_end": a.scheduled_finish, "new_start": b.scheduled_start,
                    "gap_min": round((b.scheduled_start - a.scheduled_finish).total_seconds() / 60),
                    "matrix_changeover_type": row["type"].iloc[0] if len(row) else "not in matrix",
                    "matrix_min": row["changeover_time_min"].iloc[0] if len(row) else None,
                    "wire_from (FRC)": fw.get(k(a.item)), "wire_to (FRC)": fw.get(k(b.item)),
                    "tables_disagree": bool(row["disagree"].iloc[0]) if len(row) else None})
    O = pd.DataFrame(occ)

    # ---- 2. FRC problems
    nl = f[pd.to_numeric(f["Spool/Dip Length"], errors="coerce").isna()]
    add("CTP_Master_FRC", "Spool/Dip Length", "Spool/Dip Length blank", f"{len(nl)} items", ", ".join(nl["item_code"].astype(str)),
        "Rule 2 (consume the whole spool) cannot be applied to these items if they are needed", "Fill the spool / dip length")
    ns = f[pd.to_numeric(f["line_speed_m_min"], errors="coerce").isna()]
    add("CTP_Master_FRC", "line_speed_m_min", "line speed blank", f"{len(ns)} items", ", ".join(ns["item_code"].astype(str)[:12]),
        "None today: the scheduler takes the 901 speed from CTP_Master_Routing.proc_time", "Fill, or drop the column to avoid two sources")
    sp = []
    for it in f["key"]:
        if it in rk:
            rr = r[r["routed_product"].astype(str).str.strip() == rk[it]]
            rs = pd.to_numeric(rr["proc_time"], errors="coerce").min()
            fs = pd.to_numeric(f.loc[f["key"] == it, "line_speed_m_min"], errors="coerce").iloc[0]
            if pd.notna(fs) and pd.notna(rs) and abs(fs - rs) > 0.01:
                sp.append({"item": rk[it], "FRC line_speed_m_min": fs, "Routing proc_time (M/MIN, used)": rs})
    SP = pd.DataFrame(sp)
    add("CTP_Master_FRC vs CTP_Master_Routing", "line_speed_m_min vs proc_time", "Line speed differs between FRC and routing",
        f"{len(SP)} items", "; ".join(f"{x['item']}: FRC {x['FRC line_speed_m_min']:g} vs routing {x['Routing proc_time (M/MIN, used)']:g}"
                                      for x in sp[:4]),
        "Run time follows routing (rule 4); FRC speed ignored", "Confirm which speed is right; keep one source")
    naming = [(fc, rk[k(fc)]) for fc in f["item_code"].astype(str) if k(fc) in rk and fc.strip() != rk[k(fc)]]
    add("CTP_Master_FRC (and ChangeoverMatrix)", "item_code / from_MaterialCode_O / to_MaterialCode_O",
        "Item codes written differently from routing / BOM (spaces, '-MOTHERROLL', spelling)", f"{len(naming)} items",
        ", ".join(f"'{a}' vs '{b}'" for a, b in naming[:6]) + "; MONOCHAFFER (FRC/matrix) vs MONOCHAFER (routing)",
        "Matched only through a normalising rule + alias in the scheduler; a new code could silently not match",
        "Use the same item_code in every master")
    add("CTP_Master_FRC", "tyre_type, pattern", "Columns empty", f"{int(f['tyre_type'].isna().sum())} / {len(f)} rows",
        "tyre_type and pattern blank for every row", "PCR / TBR of a calender item comes from the BOM instead", "Fill or drop")
    add("CTP_Master_FRC", "spool_length", "Same value (300) for every row; meaning unclear (used as the minimum roll)",
        f"{len(f)} rows", "spool_length = 300 everywhere", "Used as the minimum roll length (4RC-6)", "Confirm the meaning / rename")
    miss_f = [x for x in used if k(x) not in set(f["key"])]
    add("CTP_Master_FRC", "item_code", "901 items scheduled but not in FRC", f"{len(miss_f)} items", ", ".join(miss_f) or "none",
        "Rules 2 / 3 cannot use spool / wire", "Add rows" if miss_f else "-")

    # ---- 3. matrix problems
    cw = m[m["a"] == "CWBP60"]["type"].value_counts().to_dict()
    add("CTP_Master_ChangeoverMatrix", "changeover_type", "CWBP60 appears only as 'Dual groove roll set up CWBP60' -- no steel / fabric type",
        "1 item", str(cw), "Its class (steel) had to be set in config (class_override: CWBP60 steel)", "Give CWBP60 its steel / fabric pair types")
    sp_types = m["type"].unique()
    odd = [t for t in sp_types if "  " in t or t != t.strip()]
    add("CTP_Master_ChangeoverMatrix", "changeover_type", "Inconsistent text (double spaces, case)", f"{len(odd)} type values",
        "; ".join(repr(t) for t in odd), "Matched only after lower-casing / trimming", "Standardise the type names")
    pairs = {(a, b) for a, b in zip(m["a"], m["b"])}
    cal = [k(x) for x in used]
    missing = [(a, b) for a in cal for b in cal if a != b and (a, b) not in pairs]
    add("CTP_Master_ChangeoverMatrix", "from_MaterialCode_O / to_MaterialCode_O", "Pairs of scheduled 901 items missing from the matrix",
        f"{len(missing)} pairs", ", ".join(f"{a}->{b}" for a, b in missing[:6]) or "none",
        "Missing pair -> default minutes / treated as a cord change", "Add the pairs" if missing else "-")
    asym = []
    for (a, b), t in zip(zip(m["a"], m["b"]), m["type"]):
        bt = m[(m["a"] == b) & (m["b"] == a)]["type"]
        if len(bt) and bt.iloc[0].lower().replace("fabric to steel", "x").replace("steel to fabric", "x") != t.lower().replace("fabric to steel", "x").replace("steel to fabric", "x"):
            asym.append((a, b, t, bt.iloc[0]))
    add("CTP_Master_ChangeoverMatrix", "changeover_type", "Pair type differs by direction (A->B vs B->A), other than steel/fabric direction",
        f"{len(asym)} pairs", "; ".join(f"{a}->{b} '{t}' / back '{u}'" for a, b, t, u in asym[:3]) or "none", "Different creel decision by direction", "Check")
    cls_conf = []
    for it in set(m["a"]):
        g = m[m["a"] == it]["type"].str.lower()
        fr = "steel" if g.str.startswith("steel to").any() else ("fabric" if g.str.startswith("fabric to").any() else "")
        if fr and it in ft and ft[it] != fr:
            cls_conf.append((it, fr, ft[it]))
    add("CTP_Master_ChangeoverMatrix vs CTP_Master_FRC", "changeover_type vs item_type", "Steel / fabric class differs between the tables",
        f"{len(cls_conf)} items", ", ".join(f"{a}: matrix {b}, FRC {c}" for a, b, c in cls_conf) or "none", "Wrong changeover class", "Align")

    # ---- 4. which table the scheduler uses for what
    U = pd.DataFrame([
        ("Changeover minutes between any two 901 items", "CTP_Master_ChangeoverMatrix", "changeover_time_min (by from/to pair)", "phase5 changeover lookup"),
        ("Steel / fabric class of an item", "CTP_Master_ChangeoverMatrix", "changeover_type text ('Steel to ...', 'Fabric to ...') + config class_override CWBP60", "4RC creel / day plan"),
        ("Creel reload needed? (one-creel rule, placement + final repair)", "CTP_Master_ChangeoverMatrix", "changeover_type: 'compound change' = no reload, else reload", "4RC rule 5 creel pass"),
        ("Creel reload needed? (rule 3 fabric bridge + rule 1 join, after the last repair)", "CTP_Master_FRC", "wire_type: different wire = reload", "4RC rule 3 / 4RC-7 passes  <-- INCONSISTENT with the line above"),
        ("Creel reload length 240 min", "config.yaml (not a DB table)", "fourrc_rules.creel.setup_min", "plant value to confirm"),
        ("Spool / dip length (rule 2 whole spools)", "CTP_Master_FRC", "Spool/Dip Length, wire_type (shared spool)", "4RC day plan"),
        ("Minimum roll length", "CTP_Master_FRC", "spool_length (300)", "4RC-6"),
        ("901 line speed (rule 4)", "CTP_Master_Routing", "proc_time (M/MIN), efficiency (overridden 1.0)", "phase2 run time"),
        ("Which items run on 901", "CTP_Master_Routing", "routed_product, machines", "routing"),
    ], columns=["what", "table", "column(s)", "used by"])
    I = pd.DataFrame(issues)
    out = os.path.join(vdir, f"DATA_ISSUES_4RC_{v}.xlsx")
    sheets = [("Issues", I), ("Which table is used", U), ("Matrix vs FRC (all pairs)", D),
              ("Steel changes in schedule", O), ("Line speed FRC vs routing", SP if len(SP) else pd.DataFrame([{"result": "no difference"}])),
              ("FRC rows", f.drop(columns=["key"])), ("Matrix rows 4RC", m.drop(columns=["a", "b", "type"]))]
    hdr, bold = PatternFill("solid", fgColor="FF1F3864"), Font(bold=True, color="FFFFFFFF")
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for c in ws[1]:
                c.fill, c.font, c.alignment = hdr, bold, Alignment(wrap_text=True, vertical="center")
            for j, c in enumerate(df.columns, 1):
                w = min(max(10, int(df[c].astype(str).str.len().head(300).max() if len(df) else 10) + 2, len(str(c)) + 2), 80)
                ws.column_dimensions[get_column_letter(j)].width = w
                if name in ("Issues", "Which table is used"):
                    for rr_ in range(2, len(df) + 2):
                        ws.cell(row=rr_, column=j).alignment = Alignment(wrap_text=True, vertical="top")
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
    print(I[["table", "issue", "rows / items"]].to_string(index=False))
    print(f"steel changes in {v}: {len(O)}, on disagreeing pairs: {int(O['tables_disagree'].fillna(False).sum()) if len(O) else 0}")
    print(f"wrote {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "v41")
