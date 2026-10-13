#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_schedule_v18.py -- write one run in the CTP_e2e_v18_SCHEDULE.xlsx layout (2026-10-01).

    python make_schedule_v18.py <run_dir> <plan_dir> <template_v18.xlsx> <out.xlsx>

Sheets (same order, columns, header style, number formats and widths as the v18 template):
  Summary, Daily PCR TBR, Curing schedule, Building schedule, Sheet1 (line x process x machine x UOM
  by date), Upstream schedule, Assumptions, Rule compliance -- then the run's extra sheets (mixing
  batches, mixer batch plan, mixer clean-outs) in the same style.

Traceability (strict, strict_links.py): every Upstream / Building / Curing row carries lot_id and
consumption_id, rebuilt with a time-aware material ledger -- a lot consumes only supply that is
usable (finish + min aging) and not past max aging at its start, FEFO, each kg / m / tyre drawn once:
  <lot_id>(q)  scheduled lot | OPENING_WIP:<item>(q)  opening stock (production ids in 'Lot genealogy')
  RAW_MATERIAL:<item>  bought-in input | SHORT:<item>(q)  not covered by ready supply at the start
  curing lot_id = CUR<block>; its consumption_id = the build lots / opening green tyres it cures.
"""
import os
import re
import sys
from collections import defaultdict, deque

import openpyxl
import pandas as pd
import yaml
from openpyxl.styles import Alignment, Font, PatternFill
from schedule_changeovers import LIGHT_BLUE
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = PatternFill("solid", fgColor="FF1F3864")
HFONT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
HALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DT = "yyyy\\-mm\\-dd\\ hh:mm:ss"
D = "yyyy\\-mm\\-dd"
H = pd.Timedelta(hours=1)
# plan window / WIP of a run_v15.py run (env set by the runner); None = the 28-Sep..08-Oct runs
WIN = None
if os.environ.get("CTP_PLAN_START") and os.environ.get("CTP_PLAN_END"):
    _wc = os.environ.get("CTP_WIP_CUTOFF") or os.environ["CTP_PLAN_START"]
    WIN = dict(start=pd.Timestamp(os.environ["CTP_PLAN_START"]), end=pd.Timestamp(os.environ["CTP_PLAN_END"]),
               wip=f"{pd.Timestamp(_wc):%d-%b %H:%M}",
               wip_src=("jkplanning_CTP.opening_wip of that day ONLY (all items; rows with no item_type dropped; "
                        "compounds 1 batch x routing batch weight; blank units = BOM unit)"
                        if os.environ.get("CTP_WIP_SOURCE", "").lower() == "old" else
                        "jkplanning_CTP.opening_wip of that day (compounds, green tyres) + components from the latest "
                        "JK_CTP_PPDS.CTP_Opening_WIP snapshot at or before that time (WIP MIX)"
                        if os.environ.get("CTP_WIP_MIX", "") in ("1", "true", "yes")
                        else "JK_CTP_PPDS.CTP_Opening_WIP snapshot at or before that time (live DB)"))


def pday(s):
    return (pd.to_datetime(s) - pd.Timedelta(hours=7)).dt.normalize()


# ───────────────────────────────────────────────────────────────── data
def load(run, plans):
    cfg = yaml.safe_load(open(os.path.join(HERE, "config.yaml"), encoding="utf-8"))

    def inp(name):                                    # the run's own input snapshot when it has one
        p = os.path.join(run, "PHASEWISE", "00_INPUTS", name)
        return p if os.path.exists(p) else os.path.join(HERE, "inputs", name)
    rep = os.path.join(run, "DAILY_REPORT.xlsx")
    _src = [os.path.join(run, f) for f in ("phase5_schedule_updated.csv", "phase6_curing_revision_updated.csv")
            if os.path.exists(os.path.join(run, f))]
    if not os.path.exists(rep) or any(os.path.getmtime(f) > os.path.getmtime(rep) for f in _src):
        # build the daily report first -- and REBUILD it when the run is newer (a rerun into the same folder
        # otherwise kept the previous attempt's built / cured numbers)
        import subprocess
        subprocess.run([sys.executable, os.path.join(HERE, "daily_report.py"), run, plans], cwd=HERE,
                       stdout=subprocess.DEVNULL, check=True)
    summ = pd.read_excel(rep, sheet_name="Coverage by plan end")
    daily = pd.read_excel(rep, sheet_name="Daily PCR TBR")
    for d in (summ, daily):
        for a, b in (("built", "build_plan"), ("cured", "cure_plan")):
            col = a if a in d.columns else f"{a}_by_plan_end"
            if f"{a}_%" not in d.columns:
                d[f"{a}_%"] = (100 * d[col] / d[b].where(d[b] > 0)).round(1)
    fs = pd.read_csv(os.path.join(run, "phase5_floor_schedule_updated.csv"), low_memory=False,
                     parse_dates=["start_time", "end_time"])
    fs["machine"] = fs["machine"].astype(str)
    fs["item"] = fs["item"].astype(str).str.strip()
    cu = pd.read_csv(os.path.join(run, "phase6_curing_revision_updated.csv"),
                     parse_dates=["startTime", "endTime", "kit_ready"])
    plan = pd.read_csv(inp("plan.csv"), parse_dates=["startTime", "endTime"])
    bplan = pd.read_csv(inp("building_plan.csv"), parse_dates=["startTime"])
    bom = pd.read_csv(inp("bom.csv"), dtype=str, low_memory=False,
                      usecols=["Super_parent", "Parent", "Parent_qty", "Parent_unit", "child", "child_quantity",
                               "child_Unit", "child_description"])
    inv = pd.read_csv(inp("inventory.csv"), low_memory=False)
    ag = pd.read_csv(inp("aging_master.csv"))
    tt = pd.read_csv(os.path.join(HERE, "ctp_inputfiles", "item_tyre_type_live.csv"), dtype=str)
    ed = os.path.join(run, "PHASEWISE", "PHASE3_DAG", "phase3_dag_edges_updated.csv")
    edges = pd.read_csv(ed, low_memory=False, usecols=["consumer_lot", "via_sku"]) if os.path.exists(ed) else None
    # v76: the scheduler's own input quantity per unit of each consumer lot (phase-3 DAG edge qty / phase-2 lot qty),
    # so the workbook genealogy draws exactly what the scheduler planned (BOM ratio only for inputs with no edge)
    lot_inputs = {}
    lp = os.path.join(run, "PHASEWISE", "PHASE2_LOTS", "phase2_lots_updated.csv")
    if os.path.exists(ed) and os.path.exists(lp):
        _e = pd.read_csv(ed, low_memory=False, usecols=["producer_item", "consumer_lot", "consume_qty"])
        _e["consume_qty"] = pd.to_numeric(_e["consume_qty"], errors="coerce").fillna(0)
        _q = pd.read_csv(lp, low_memory=False, usecols=["lot_id", "lot_qty"]).set_index("lot_id")["lot_qty"]
        _s = _e[_e["consume_qty"] > 0].groupby(["consumer_lot", "producer_item"])["consume_qty"].sum().reset_index()
        _s["lq"] = _s["consumer_lot"].map(_q)
        _s = _s[_s["lq"] > 0]
        for c, p, v in zip(_s["consumer_lot"], _s["producer_item"].astype(str).str.strip(), _s["consume_qty"] / _s["lq"]):
            lot_inputs.setdefault(c, {})[p] = float(v)
    tbr_skus = set()
    for f in os.listdir(plans):
        if re.match(r"ctp_curing_schedule_.*tbr.*\.xlsx$", f, re.I):
            tbr_skus |= set(pd.read_excel(os.path.join(plans, f), "Shift Schedule")["SKUCode"].astype(str).str.strip())
    return dict(cfg=cfg, summ=summ, daily=daily, fs=fs, cu=cu, plan=plan, bplan=bplan, bom=bom, inv=inv,
                ag=ag, tt=tt, tbr_skus=tbr_skus, run=run, edges=edges, lot_inputs=lot_inputs)


def plan_end(D_):
    return D_["plan"]["endTime"].max()


def sku2gt_map(bom):
    g = bom[bom.child_description.astype(str).str.strip().str.upper().isin(["GREEN TYRES", "GREEN_TYRE"])]
    return dict(zip(g.Super_parent.astype(str).str.strip(), g.child.astype(str).str.strip()))


def build_lot_sku(D_):
    """sku per build lot: match the lot to its plan row by (TBM, qty) in time order"""
    fs, bplan = D_["fs"], D_["bplan"]
    b = fs[fs.department.str.upper() == "BUILDING"]
    lots = b.groupby("lot_id").agg(m=("machine", "first"), s=("start_time", "min"),
                                   q=("produce_qty", "sum")).sort_values("s")
    free = defaultdict(list)
    for _, r in bplan.sort_values("startTime").iterrows():
        free[(str(r.buildMachine), int(round(r.qty)))].append(str(r.skuCode).strip())
    sku = {}
    for lid, r in lots.iterrows():
        k = (r.m, int(round(r.q)))
        sku[lid] = free[k].pop(0) if free.get(k) else ""
    return sku


def strict(D_):
    """strict genealogy (strict_links.py), computed once per workbook"""
    if "strict" in D_:
        return D_["strict"]
    sys.path.insert(0, HERE)
    import strict_links
    fs = D_["fs"][D_["fs"].department.str.upper() != "CURING"]
    lot_sku = {}
    if D_.get("edges") is not None:
        e = D_["edges"].dropna()
        lot_sku = e.groupby("consumer_lot")["via_sku"].agg(lambda x: x.mode().iat[0]).astype(str).to_dict()
    lot_sku.update({k: v for k, v in build_lot_sku(D_).items() if v})
    cu = D_["cu"].copy()
    cu["skuCode"] = cu["skuCode"].astype(str).str.strip()
    tt = D_["cfg"].get("gt_transfer_time_min") or 90.0          # routing GT BUILDING transfer (phase 6 S6)
    cmap, cure_map, gen, st = strict_links.build(fs, D_["bom"], D_["inv"], D_["ag"], cu, sku2gt_map(D_["bom"]),
                                                 lot_sku=lot_sku, gt_transfer_min=tt, lot_inputs=D_.get("lot_inputs"))
    lots = fs.groupby("lot_id").agg(item=("item", "first"), dep=("department", "first"), e=("end_time", "max"))
    cur_id = {i: f"CUR{int(b):05d}" for i, b in cu.block_id.items()}
    gen["consumer_lot"] = [cur_id[k] if kd else k for k, kd in zip(gen.consumer, gen.kind)]
    gen["consumer_stage"] = [("CURING" if kd else lots.at[k, "dep"]) for k, kd in zip(gen.consumer, gen.kind)]
    gen["consumer_item"] = [(cu.at[k, "skuCode"] if kd else lots.at[k, "item"]) for k, kd in zip(gen.consumer, gen.kind)]
    gen["qty"] = gen["qty"].round(3)
    gen = gen.rename(columns={"consumer_start": "first_draw"})[
        ["consumer_lot", "consumer_stage", "consumer_item", "first_draw", "last_draw", "input_item", "supply_type",
         "supply_id", "qty", "supply_usable", "ready_at_draw"]].sort_values(["first_draw", "consumer_lot"])
    D_["strict"] = (cmap, cure_map, gen, st)
    return D_["strict"]


def curing_sheet(D_):
    cu, plan, bom, inv, fs = D_["cu"], D_["plan"], D_["bom"], D_["inv"], D_["fs"]
    f = plan[(plan.skuCode.astype(str).str.strip() != "CHANGEOVER")
             & (pd.to_numeric(plan.qty, errors="coerce").fillna(0) > 0)].reset_index(drop=True)
    cu = cu.join(f[["startTime"]].rename(columns={"startTime": "plan_start"}), on="block_id")
    cu["skuCode"] = cu["skuCode"].astype(str).str.strip()
    _, cure_map, _, _ = strict(D_)                  # strict FEFO: only green tyres ready at the cure start
    lot = {i: f"CUR{int(b):05d}" for i, b in cu.block_id.items()}
    cons = cure_map
    pe = plan_end(D_)
    out = pd.DataFrame({
        "line": cu.skuCode.map(lambda s: "TBR" if s in D_["tbr_skus"] else "PCR"),
        "skuCode": cu.skuCode, "pressNo": cu.pressNo, "shift": cu["shift"],
        "startTime": cu.startTime, "endTime": cu.endTime, "qty": cu.qty, "cycleTime": cu.cycleTime,
        "lot_id": pd.Series(lot), "consumption_id": pd.Series(cons),
        "gt_short": pd.Series(cons).astype(str).str.contains("SHORT:"), "kit_status": cu.kit_status,
        "kit_ready": cu.kit_ready, "green_age_h": pd.to_numeric(cu.green_age_h, errors="coerce").round(2),
        "within_horizon": cu.endTime <= pe,
        "push_h": ((cu.startTime - cu.plan_start) / H).round(2)})
    return out.sort_values(["startTime", "pressNo"])


def building_sheet(D_):
    fs, bplan = D_["fs"], D_["bplan"]
    b = fs[fs.department.str.upper() == "BUILDING"].copy()
    sku = build_lot_sku(D_)
    # strict kit: consumption_id = the supply actually drawn at the build start (strict_links);
    # kit_ready = when the last drawn component became usable (finish + min aging)
    cmap, _, gen, _ = strict(D_)
    ready = gen[gen.consumer_stage.astype(str).str.upper() == "BUILDING"].groupby("consumer_lot")["supply_usable"].max()
    pe = plan_end(D_)
    kr = b.lot_id.map(ready)
    out = pd.DataFrame({
        "line": b.machine.map(lambda m: "TBR" if m.startswith("TBM-") else "PCR"),
        "date": pd.to_datetime(b["date"]), "shift": b["shift"], "machine": b.machine, "item": b["item"],
        "sku": b.lot_id.map(sku), "start_time": b.start_time, "end_time": b.end_time,
        "qty": pd.to_numeric(b.produce_qty, errors="coerce").round(0), "lot_id": b.lot_id,
        "prod_id": b.prod_id, "consumption_id": b.lot_id.map(cmap), "kit_ready": kr,
        "kit_incomplete": b.lot_id.map(cmap).astype(str).str.contains("SHORT:"),
        "within_horizon": b.end_time <= pe})
    return out.sort_values(["start_time", "machine"])


def upstream_sheet(D_):
    fs, bom, inv, tt = D_["fs"], D_["bom"], D_["inv"], D_["tt"]
    u = fs[~fs.department.str.upper().isin(["BUILDING", "CURING"])].copy()
    line = dict(zip(tt["item"].astype(str).str.strip(), tt["tyre_type"].str.upper()))
    kids = bom.assign(P=bom.Parent.astype(str).str.strip(), C=bom.child.astype(str).str.strip()) \
        .groupby("P")["C"].apply(lambda s: list(dict.fromkeys(s))).to_dict()
    produced = set(fs["item"])
    stock = set(inv.loc[pd.to_numeric(inv["inventory"], errors="coerce").fillna(0) > 0, "itemcode"]
                .astype(str).str.strip())
    lot_item = fs.groupby("lot_id")["item"].first().to_dict()

    def expand(item, c):
        toks = [t for t in str(c).split(";") if t and t != "nan"]
        linked = {lot_item.get(t) for t in toks if t in lot_item}
        out = []
        for t in toks:
            if t in ("RAW_MATERIAL", "WIP"):
                # every BOM input of this item that no linked lot supplies, labelled by where it comes
                # from: opening stock, a bought-in raw material, or produced in the plan but unlinked
                for k in kids.get(item, []):
                    if k in linked:
                        continue
                    out.append(f"OPENING_WIP:{k}" if k in stock else
                               f"RAW_MATERIAL:{k}" if k not in produced else f"UNLINKED_SUPPLY:{k}")
                if not kids.get(item):                   # the item has no inputs in the BOM (e.g. BP01)
                    out.append(("RAW_MATERIAL" if t == "RAW_MATERIAL" else "OPENING_WIP") + ":NO_BOM_INPUTS")
            else:
                out.append(t)
        return ";".join(dict.fromkeys(out))

    cmap = strict(D_)[0]                             # strict ledger links (expand() kept for reference)
    qty = pd.to_numeric(u.produce_qty, errors="coerce")
    uom = u.UOM.astype(str)
    if "batches" in u.columns:                        # BR-V1: mixing planned in batches
        isb = u["batches"].notna()
        qty = qty.where(~isb, pd.to_numeric(u["batches"], errors="coerce"))
        uom = uom.where(~isb, "BATCH")
    out = pd.DataFrame({
        "line": u["item"].map(line).fillna("").replace({"BOTH": "PCR+TBR"}),
        "date": u["date"].astype(str), "shift": u["shift"], "machine": u.machine, "item": u["item"],
        "item_type": u.item_type, "process": u.process, "start_time": u.start_time, "end_time": u.end_time,
        "produce_qty": qty.round(3), "UOM": uom, "lot_id": u.lot_id, "consumption_id": u.lot_id.map(cmap),
        "input_short": u.lot_id.map(cmap).astype(str).str.contains("SHORT:"), "FG_SKU_CODE": u.FG_SKU_CODE})
    # user 2026-10-07: the line filter holds PCR or TBR only -- an item used by both lines (or with no tyre type) is
    # given the line its LOT feeds: the majority of the lot's SKUs (phase2 skus_set) by the TBR plan SKU list
    _sh = ~out["line"].isin(["PCR", "TBR"])
    if _sh.any():
        try:
            _L = pd.read_csv(os.path.join(D_["run"], "phase2_lots_updated.csv"), low_memory=False,
                             usecols=["lot_id", "skus_set", "equipment"]).drop_duplicates("lot_id").set_index("lot_id")
            _tb = D_["tbr_skus"]

            def _ln(lid):
                ss = [x.strip() for x in str(_L["skus_set"].get(lid, "") or "").split(",") if x.strip() and x.strip() != "nan"]
                if ss:
                    t = sum(1 for x in ss if x in _tb)
                    return "TBR" if t * 2 >= len(ss) and t else "PCR"
                eq = str(_L["equipment"].get(lid, "")).upper()
                return "TBR" if "TBR" in eq else "PCR"
            out.loc[_sh, "line"] = out.loc[_sh, "lot_id"].map(_ln)
        except Exception:
            out.loc[_sh, "line"] = "PCR"
    return out.sort_values(["start_time", "machine"])


def sheet1_table(up):
    p = up.pivot_table(index=["line", "process", "machine", "UOM"], columns="date", values="produce_qty",
                       aggfunc="sum", fill_value=0)
    p["Grand Total"] = p.sum(axis=1)
    p = p.reset_index().sort_values(["line", "process", "machine", "UOM"])
    tot = {c: (p[c].sum() if c not in ("line", "process", "machine", "UOM") else None) for c in p.columns}
    tot["line"] = "Grand Total"
    return pd.concat([p, pd.DataFrame([tot])], ignore_index=True)


def _co_txt(br, grp):
    b = br.get("butyl_sequence_rule") or {}
    g = b.get(grp) or {}
    basis = str(g.get("cleanout_basis", b.get("cleanout_basis", "fixed"))).lower()
    return ("none - 60 s compound changeover only" if basis == "none" else
            "1 batch of the compound cleaned out" if basis == "one_batch_cycle" else
            f"{g.get('cleanout_sec', b.get('cleanout_sec'))} s")


def assumptions_sheet(D_):
    cfg, rows = D_["cfg"], []
    gl = pd.read_csv(os.path.join(HERE, "inputs", "routing_gapfill_log.csv"))
    for _, r in gl.iterrows():
        rows.append(("", "phase 0 - masters", "", r.routed_product, str(r.operation), str(r.operation),
                     "proc_time" if pd.notna(r.proc_time) else "routing",
                     f"{r.proc_time} {r.uom}" if pd.notna(r.proc_time) else "not filled", r.source))
    ca = pd.read_csv(os.path.join(HERE, "inputs", "ctb_approximations_log.csv"))
    for _, r in ca.iterrows():
        rows.append(("", "phase 0 - masters", "", r["item"], "", "", r["kind"], r.value_used, r.basis))
    br = cfg.get("business_rules") or {}
    rows += [
        ("PCR+TBR", "mixing", "", "XMT180, GL7", "MIXING", "MIXING", "batch size",
         "300 kg, 215 kg", "blank in JK_CTP_PPDS.CTP_Master_Routing -> gap-filled value used (BR-V1)"),
        ("PCR+TBR", "mixing", "", "all compounds", "MIXING", "MIXING", "BRM-13 buffer batches",
         str(br.get("buffer_batches")), "baseline v5: 0 (plant rule BRM-13 says +5) -- user decision 2026-10-01"),
        ("PCR+TBR", "mixing", "", "all mixers", "MIXING", "MIXING", "changeover",
         f"{cfg.get('mixer_changeover_min')} min", "BR-74 / user rule: 60 s (Upstream Business Rules BRM-10 says 2 min)"),
        ("PCR+TBR", "mixing", "", "30G/IL3/IL5 (final)", "MIXING", "MIXING", "butyl clean-out",
         _co_txt(br, "final"), "BR-V2/V13; B1, B2 not in routing/BOM -> only B99 usable as prep"),
        ("PCR+TBR", "mixing", "", "29MG/ML405/ML605 (master)", "MIXING", "MIXING", "butyl clean-out",
         _co_txt(br, "master"), "BR-V2/V13; XMB98 Repass not in routing/BOM"),
        ("PCR", "cutting", "", "machine 2001", "CUTTING", "BIAS CUTTER", "rate",
         "x1.66", "provisional calibration (FLAW C-2)"),
        ("PCR+TBR", "aging", "", "green tyres", "BUILDING", "BUILDING", "max aging", "72 h",
         "green_tyre_shelf_life_all_h (user)"),
    ]
    if WIN:   # run_v15.py: window / WIP taken from the run's own settings
        rows += [
            ("PCR+TBR", "calendar", "", "all machines", "", "", "holiday / PM", "from downtime_reason (live DB)",
             "jkplanning_CTP.downtime_reason read at run time; windows inside the plan: see Rule compliance BRG-01"),
            ("PCR+TBR", "horizon", "", "all machines", "", "", "machine open", f"plan start {WIN['start']:%d-%b %H:%M}",
             f"BR-V3 open_machines_at_plan_start; plan rows {WIN['start']:%d-%b %H:%M} -> {WIN['end']:%d-%b %H:%M} kept"),
            ("PCR+TBR", "inventory", "", "opening WIP", "", "", "snapshot", WIN["wip"], WIN["wip_src"]),
        ]
    else:
        rows += [
            ("PCR+TBR", "calendar", "", "all machines", "", "", "holiday", "02-Oct 07:00 -> 03-Oct 07:00",
             "plant closed; HOLIDAY rows added to downtime_reason working copy"),
            ("PCR+TBR", "horizon", "", "all machines", "", "", "machine open", "plan start 28-Sep 07:00",
             "BR-V3 open_machines_at_plan_start"),
            ("PCR+TBR", "inventory", "", "opening WIP", "", "", "snapshot", "28-Sep 06:30",
             "JK_CTP_PPDS.CTP_Opening_WIP (saved 30-Sep export)"),
        ]
    df = pd.DataFrame(rows, columns=["line", "stage", "code", "item", "department", "operation", "gap",
                                     "value", "source"])
    df["code"] = [f"A{i + 1:02d}" for i in range(len(df))]
    return df


IMPL = {  # rule id prefix -> (implemented_in, how_it_is_enforced)
    "BRG-01": ("phase5 BR-10 downtime + final_repair_pass", "PM/MTC and holiday windows are busy time on every machine"),
    "BRG-02": ("phase5 aging gates (BR-12)", "min/max aging from the aging master; builds released only with an aging-valid kit"),
    "BRG-03": ("build_sep_inputs.py + adapt_inputs.py", "DB MPQ (TBR) + plant-corrected local MPQ workbook (PCR)"),
    "BRM-01": ("phase2_lot_sizing.py", "mixing lot >= mixing_mpq_min_batches x routing batch size"),
    "BRM-02": ("phase2 business-rule block", "JT444/JT466 lots rounded up to 3-batch stacks"),
    "BRM-03": ("butyl_rule.py + phase5 (BR-V2)", "prep compound before/after protected compounds; broken final sequence -> 120 s clean-out, master -> 60 s changeover only"),
    "BRM-05": ("phase5 rule_repair_pass", "ML405/ML406 followed by MK3 when an MK3 lot is ready"),
    "BRM-06": ("phase5 rule_repair_pass", "TR100 clean-out time charged on a master<->final switch"),
    "BRM-09": ("phase5 rule_repair_pass", "same compound pulled together (<= pull_early_h)"),
    "BRM-10": ("phase5 compute_changeover_min", "fixed mixer changeover (config mixer_changeover_min)"),
    "BRM-11": ("routing machine eligibility", "lots only on routed machines"),
    "BRM-12": ("routing proc_time x efficiency", "MIN/BATCH rate per mixer"),
    "BRM-13": ("phase2 business-rule block", "buffer_batches added to each compound's last lot"),
    "BRE-01": ("phase5 placement (JIT)", "components timed to building need"),
    "BRE-02": ("phase5 rule_repair_pass", "same item pulled into one run per shift"),
    "BRE-04": ("build_sep_inputs.py", "plant MPQ master"),
    "BRE-05": ("phase5 changeover matrix", "machine-specific matrix (config machine_to_changeover_line)"),
    "BRE-06": ("routing machine eligibility", "lots only on routed machines"),
    "BRIL-01": ("phase5 placement (JIT)", "inner liner timed to building need"),
    "BRIL-02": ("phase5 rule_repair_pass", "same item pulled into one run per shift"),
    "BRIL-04": ("build_sep_inputs.py", "plant MPQ master"),
    "BRIL-05": ("phase5 changeover matrix", "machine-specific matrix"),
    "BRIL-06": ("routing machine eligibility", "lots only on routed machines"),
    "BRC-01": ("phase5 rule_repair_pass", "Belt-2 preferred right after Belt-1 of the same code"),
    "BRC-02": ("phase5 rule_repair_pass", "same item pulled into one run per shift"),
    "BRC-03": ("phase5 changeover matrix", "machine-specific matrix"),
    "BRC-04": ("routing machine eligibility", "lots only on routed machines"),
    "BR4RC-05": ("phase5 rule overrides", ">= 15 min clean-out before/after CAP66"),
    "BR4RC-07": ("phase5 changeover matrix", "4 RC Calendar line of the plant matrix"),
    "BR4RC-08": ("routing line speeds", "rates from routing (not audited)"),
}


def rules_sheet(D_):
    x = pd.read_excel(os.path.join(HERE, "businees_rules", "Upstream Business Rules.xlsx"), header=None)
    hi = next(i for i in range(10) if "Rule ID" in [str(v).strip() for v in x.iloc[i]])
    x.columns = [str(c).strip() for c in x.iloc[hi]]
    x = x.iloc[hi + 1:]
    x = x[x["Rule ID"].notna()]
    a = pd.read_csv(os.path.join(D_["run"], "BUSINESS_RULE_AUDIT.csv"))
    amap = []
    for _, r in a.iterrows():
        ids = set()
        for part in str(r.rule_id).split("/"):
            tok = part.strip().split(" ")[0]
            ids.add(tok)
            mm = re.match(r"([A-Za-z0-9]+)-(\d+(?:,\d+)*)$", tok)
            if mm:
                ids |= {f"{mm.group(1)}-{n}" for n in mm.group(2).split(",")}
        amap.append((ids, str(r.status), str(r.measured), str(r.detail) if pd.notna(r.detail) else ""))
    rows = []
    for _, r in x.iterrows():
        rid = str(r["Rule ID"]).strip()
        st, meas, det = next(((s, m, d) for ids, s, m, d in amap if rid in ids), ("NOT AUDITED", "", ""))
        verdict = ("APPLIED" if st.startswith("FOLLOWED") else "PARTIAL" if st.startswith("PARTLY") else
                   "VIOLATED" if st == "VIOLATED" else "NOT APPLICABLE" if st.startswith("N/A") else
                   "NO DATA" if st.startswith("NO DATA") else "NOT MODELLED" if st == "NOT MODELLED" else st)
        impl, how = IMPL.get(rid, ("-", "-"))
        why = "" if verdict == "APPLIED" else (det or st)
        close = {"VIOLATED": "tighten the rule_repair_pass / add a hard constraint; plant to confirm priority",
                 "PARTIAL": "extend the rule to the remaining cases",
                 "NOT MODELLED": "needs sequencing logic and plant family/die data",
                 "NO DATA": "plant to supply the master data"}.get(verdict, "")
        rows.append((rid, r.get("Section", ""), r.get("Category", ""), r.get("Sub-Category", ""),
                     str(r.get("Business Rule", "")).replace("\n", " "), verdict, impl, how, meas, why, close))
    for rid, rule, impl, how in [
        ("BR-V1", "Mixing output planned in batches (kg -> batches from CTP_Master_Routing)",
         "mixing_batches.py, phase2, phase5", "lots rounded up to whole batches; shift rows in whole batches"),
        ("BR-V2", "Butyl-friendly sequence (final: B1/B2/B99 before 30G/IL3/IL5; master: MB100/MB200/XMB98 around 29MG/ML405/ML605); broken final -> 120 s clean-out, broken master -> 60 s changeover only (BR-V13)",
         "butyl_rule.py, phase5", "prep preferred in re-sequencing; clean-out reserved as busy time"),
        ("BR-V3", "Machines open at the plan start", "phase5 schedule-open floor", "open_machines_at_plan_start"),
        ("BR-V4", "Changeover against the real time-neighbours", "phase5 first_fit_nb", "neighbour_aware_changeover")]:
        rows.append((rid, "CTP versions", "Scheduler", "", rule, "APPLIED", impl, how,
                     "see latest_schedule/BUSINESS_RULES_LOG.md", "", ""))
    return pd.DataFrame(rows, columns=["rule_id", "section", "category", "sub_category", "business_rule",
                                       "verdict", "implemented_in", "how_it_is_enforced",
                                       "measured_from_the_schedule", "why_missed", "how_to_close_it"])


def extra_sheets(D_, up):
    fs = D_["fs"]
    out = {}
    if "batches" in fs.columns and fs["batches"].notna().any():
        mix = fs[fs["batches"].notna()].copy()
        mix["batches"] = pd.to_numeric(mix["batches"], errors="coerce").astype(int)
        a = (mix.groupby(["machine", "item", "item_type", "batch_size_kg"], as_index=False)
             .agg(batches=("batches", "sum"), kg=("produce_qty", "sum"), runs=("lot_id", "nunique"),
                  first_start=("start_time", "min"), last_end=("end_time", "max")))
        out["Mixing batches"] = a.sort_values(["machine", "first_start"])
        p = mix.pivot_table(index="machine", columns="date", values="batches", aggfunc="sum", fill_value=0)
        p["Total"] = p.sum(axis=1)
        p.loc["ALL MIXERS"] = p.sum()
        out["Mixer batch plan"] = p.reset_index()
    try:
        sys.path.insert(0, HERE)
        from make_schedule_from_template import cleanouts
        co = cleanouts(D_["run"])
        if co is not None:
            out["Mixer clean-outs"] = co
    except Exception:
        pass
    il = il_daily_plan(D_)
    if il is not None:
        out["IL3-IL5 daily plan"] = il
    cmap, cure_map, gen, st = strict(D_)
    out["Lot genealogy"] = gen
    sh = gen[gen.supply_type == "SHORT"]
    out["Genealogy check"] = pd.DataFrame([
        ("links drawn (scheduled lots + opening-stock lots)", int((gen.supply_type != "SHORT").sum())),
        ("links where a unit was drawn before it was usable", int((~gen.ready_at_draw).sum())),
        ("links to a supply id not in the schedule / inventory", int(gen[gen.supply_type != "SHORT"].supply_usable.isna().sum())),
        ("consumer lots with a SHORT input", int(sh.consumer_lot.nunique())),
        ("  of which curing blocks", int(sh[sh.consumer_stage == "CURING"].consumer_lot.nunique())),
        ("  of which building lots", int(sh[sh.consumer_stage.astype(str).str.upper() == "BUILDING"].consumer_lot.nunique())),
        ("  of which upstream lots", int(sh[~sh.consumer_stage.astype(str).str.upper().isin(["CURING", "BUILDING"])].consumer_lot.nunique())),
        ("bought-in raw-material inputs (RAW_MATERIAL:)", st.get("raw", 0)),
        ("items with stock used / left past max aging", st.get("expired_items", 0)),
        ("links drawing stock past max aging (EXPIRED_STOCK; aging is a report)", int((gen.supply_type == "EXPIRED_STOCK").sum())),
        ("  consumer lots using expired stock", int(gen[gen.supply_type == "EXPIRED_STOCK"].consumer_lot.nunique())),
        ("  qty drawn past max aging", round(float(gen.loc[gen.supply_type == "EXPIRED_STOCK", "qty"].sum()), 1)),
    ], columns=["check", "value"])
    return out


def il_daily_plan(D_):
    """BR-V10 validation: every IL3/IL5 campaign (consecutive IL lots on a mixer) per plant day."""
    cfg = D_["cfg"]
    ilc = cfg.get("il_daily_campaign") or {}
    if not ilc.get("enabled"):
        return None
    items = [str(x) for x in (ilc.get("items") or ["IL3", "IL5"])]
    prep = {"B1", "B2", str(ilc.get("prep", "B99"))}
    s = pd.read_csv(os.path.join(D_["run"], "phase5_schedule_updated.csv"), low_memory=False,
                    parse_dates=["scheduled_start", "scheduled_finish"]).drop_duplicates("lot_id")
    s = s[s["department"].astype(str).str.upper().str.contains("MIX")].copy()
    s["machine"], s["item"] = s["machine"].astype(str), s["item"].astype(str).str.strip()
    L = pd.read_csv(os.path.join(D_["run"], "phase2_lots_updated.csv"), low_memory=False,
                    usecols=["lot_id", "batches"]).set_index("lot_id")["batches"]
    rows = []
    for m, g in s.groupby("machine"):
        r = list(g.sort_values("scheduled_start").itertuples())
        i = 0
        while i < len(r):
            if r[i].item in items:
                j = i
                while j + 1 < len(r) and r[j + 1].item in items:
                    j += 1
                blk = r[i:j + 1]
                prev = r[i - 1] if i > 0 else None
                gaps = [round((b.scheduled_start - a.scheduled_finish).total_seconds()) for a, b in zip(blk, blk[1:])]
                rows.append({
                    "plant_day": (blk[0].scheduled_start - pd.Timedelta(hours=7)).normalize(), "machine": m,
                    "sequence": " -> ".join(x.item for x in blk),
                    "IL3_batches": sum(float(L.get(x.lot_id, 0) or 0) for x in blk if x.item == "IL3"),
                    "IL5_batches": sum(float(L.get(x.lot_id, 0) or 0) for x in blk if x.item == "IL5"),
                    "campaign_start": blk[0].scheduled_start, "campaign_end": blk[-1].scheduled_finish,
                    "preceded_by": prev.item if prev is not None else "",
                    "prep_ok": bool(prev is not None and prev.item in prep),
                    "gap_before_sec": round((blk[0].scheduled_start - prev.scheduled_finish).total_seconds()) if prev is not None else None,
                    "gaps_inside_sec": ", ".join(str(x) for x in gaps), "lots": ";".join(x.lot_id for x in blk)})
                i = j + 1
            else:
                i += 1
    df = pd.DataFrame(rows).sort_values(["plant_day", "campaign_start"])
    df["campaigns_that_day"] = df.groupby("plant_day")["plant_day"].transform("size")
    return df


def split_line_sheets(cur, bld, up):
    """User 2026-10-06: separate PCR and TBR sheets (own filter each); building / upstream rows sorted by machine
    and start with a light-blue CHANGEOVER / CLEANOUT row (duration) between lots of different items.
    Upstream items used by both lines go to 'Upstream PCR+TBR'."""
    from schedule_changeovers import Changeovers, with_changeovers
    co = Changeovers()
    bld2 = with_changeovers(
        bld, "machine", "item", "start_time", "end_time",
        lambda p, r: co.building(r["line"], p["item"], r["item"]),
        lambda p, r, mins, kind, basis: {"line": r["line"], "date": r["date"], "shift": r["shift"], "item": kind,
                                         "sku": f"{p['item']} -> {r['item']} | {basis} | {round(mins, 1):g} min"})
    up2 = with_changeovers(
        up, "machine", "item", "start_time", "end_time",
        lambda p, r: co.upstream(r["machine"], str(p["item"]), str(r["item"]), p.get("item_type", ""), r.get("item_type", "")),
        lambda p, r, mins, kind, basis: {"line": r["line"], "date": r["date"], "shift": r["shift"], "item": kind,
                                         "item_type": f"{p['item']} -> {r['item']}",
                                         "process": f"{basis} | {round(mins, 1):g} min", "UOM": "min"})
    cf = {"startTime": DT, "endTime": DT, "kit_ready": DT}
    bf = {"date": D, "start_time": DT, "end_time": DT, "kit_ready": DT}
    uf = {"start_time": DT, "end_time": DT}
    shared = ~up2["line"].isin(["PCR", "TBR"])
    return [("Curing PCR", cur[cur["line"] == "PCR"], cf, "A2"),
            ("Curing TBR", cur[cur["line"] == "TBR"], cf, "A2"),
            ("Building PCR", bld2[bld2["line"] == "PCR"], bf, "F2"),
            ("Building TBR", bld2[bld2["line"] == "TBR"], bf, "F2"),
            ("Upstream PCR", up2[up2["line"] == "PCR"], uf, "B2"),
            ("Upstream TBR", up2[up2["line"] == "TBR"], uf, "B2"),
            ("Upstream PCR+TBR", up2[shared], uf, "B2")]


# ───────────────────────────────────────────────────────────────── writing
def write_df(ws, df, start_row=1, widths=None, fmts=None):
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start_row, column=j, value=str(c))
        cell.font, cell.fill, cell.alignment = HFONT, NAVY, HALIGN
    for i, row in enumerate(df.itertuples(index=False, name=None), start_row + 1):
        for j, v in enumerate(row, 1):
            if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
                continue
            if isinstance(v, pd.Timestamp):
                v = v.floor("s").to_pydatetime()
            elif hasattr(v, "item"):
                v = v.item()
            elif isinstance(v, str):
                v = ILLEGAL_CHARACTERS_RE.sub("", v)       # control chars in some DB production ids
            ws.cell(row=i, column=j, value=v)
    n = len(df) + start_row
    for j, c in enumerate(cols, 1):
        L = get_column_letter(j)
        f = (fmts or {}).get(c)
        if f:
            for r in range(start_row + 1, n + 1):
                ws[f"{L}{r}"].number_format = f
        w = (widths or {}).get(c)
        if w is None:
            s = df[c].head(300).astype(str).str.len().max() if len(df) else 8
            w = min(max(10, int(max(s, len(str(c))) + 2)), 60)
        ws.column_dimensions[L].width = w
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{n}"
    return n


def template_widths(tpl, name):
    if name not in tpl.sheetnames:
        return {}
    ws = tpl[name]
    hr = 4 if name == "Sheet1" else 1
    out = {}
    for c in ws[hr]:
        if c.value is not None:
            d = ws.column_dimensions.get(c.column_letter)
            if d is not None and d.width:
                out[str(c.value)] = round(d.width, 2)
    return out


def main():
    run, plans, template, out = sys.argv[1:5]
    D_ = load(run, plans)
    tpl = openpyxl.load_workbook(template)
    cur, bld, up = curing_sheet(D_), building_sheet(D_), upstream_sheet(D_)
    _WB = D_["cfg"].get("schedule_workbook") or {}          # user 2026-10-06: original layout unless switched on
    s1 = sheet1_table(up)
    sheets = [
        ("Summary", D_["summ"][["line", "build_plan", "built_by_plan_end", "cure_plan", "cured_by_plan_end",
                                "built_%", "cured_%"]], {}, "A2"),
        ("Daily PCR TBR", D_["daily"].loc[D_["daily"]["line"].isin(["PCR", "TBR"]),
                                          ["line", "day", "build_plan", "built", "cure_plan", "cured",
                                           "built_%", "cured_%"]], {}, "A2"),
    ] + (split_line_sheets(cur, bld, up) if _WB.get("split_by_line", False) else [
        ("Curing schedule", cur, {"startTime": DT, "endTime": DT, "kit_ready": DT}, "A2"),
        ("Building schedule", bld, {"date": D, "start_time": DT, "end_time": DT, "kit_ready": DT}, "D2"),
    ]) + [
        ("Sheet1", s1, None, "E5"),
    ] + ([] if _WB.get("split_by_line", False) else [
        ("Upstream schedule", up, {"start_time": DT, "end_time": DT}, "A2"),
    ]) + [
        ("Assumptions", assumptions_sheet(D_), {}, "A2"),
        ("Rule compliance", rules_sheet(D_), {}, "E2"),
    ]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    blue = PatternFill("solid", fgColor=LIGHT_BLUE)
    for name, df, fmts, frz in sheets:
        ws = wb.create_sheet(name)
        if name == "Sheet1":
            ws["A3"], ws["E3"] = "Sum of produce_qty", "date"
            ws["A3"].font = ws["E3"].font = Font(bold=True)
            dcols = [c for c in df.columns if c not in ("line", "process", "machine", "UOM")]
            n = write_df(ws, df, start_row=4, widths=template_widths(tpl, name),
                         fmts={c: "#,##0.0" for c in dcols})
            ws.auto_filter.ref = f"A4:{get_column_letter(len(df.columns))}{n - 1}"
            for j in range(1, len(df.columns) + 1):
                ws.cell(row=n, column=j).font = Font(bold=True)
        else:
            write_df(ws, df, widths=template_widths(tpl, name), fmts=fmts)
            if "row_type" in df.columns:                  # changeover / clean-out rows in light blue
                for i, rt in enumerate(df["row_type"].tolist(), 2):
                    if rt != "LOT":
                        for j in range(1, len(df.columns) + 1):
                            ws.cell(row=i, column=j).fill = blue
        ws.freeze_panes = frz
        print(f"  {name:<20} {len(df):>6,} rows")
    ws = wb["Summary"]
    _w = (f"{WIN['start']:%d-%b %H:%M} -> {WIN['end']:%d-%b %H:%M}" if WIN
          else "28-Sep 07:00 -> 08-Oct 07:00, 02-Oct holiday")
    _e = f"{WIN['end']:%d-%b}" if WIN else "08-Oct"
    ws.cell(row=6, column=1, value=(
        f"build_plan / cure_plan are the plant plan ({_w}); "
        "built / cured are what this schedule achieves inside that window. A shortfall is work that lands "
        f"after the window ({_e} onward) -- the plan itself is not rewritten."))
    if _WB.get("split_by_line", False):
        ws.cell(row=7, column=1, value=(
        "Sheets per line: Curing / Building / Upstream PCR and TBR (Upstream PCR+TBR = items used by both lines), "
        "each with its own filter. Light-blue rows = CHANGEOVER / CLEANOUT between two lots of different items on a "
        "machine (from -> to, basis, duration_min); filter row_type = LOT to see production only."))
    _xs = extra_sheets(D_, up)
    # v40: 'Mixer clean-outs' = every clean-out scheduled on a mixer (after-butyl 120 s, TR100, butyl-sequence), taken
    # from the light-blue CLEANOUT rows of the Upstream sheets (the old list held only broken-sequence clean-outs)
    _mx = {str(x) for x in ((D_["cfg"].get("business_rules") or {}).get("mixers") or [])}
    _ups_ = [df for nm_, df, _, _ in sheets if str(nm_).startswith("Upstream ") and "row_type" in df.columns]
    _co = pd.concat(_ups_, ignore_index=True) if _ups_ else pd.DataFrame(columns=["row_type", "machine"])
    if "row_type" not in _co.columns:
        _co = _co.iloc[0:0].assign(row_type=pd.Series(dtype=str))
    _co = _co[(_co.get("row_type", pd.Series(dtype=str)) == "CLEANOUT") & _co["machine"].astype(str).isin(_mx)]
    if len(_co) and _WB.get("split_by_line", False):
        _xs["Mixer clean-outs"] = pd.DataFrame({
            "machine": _co["machine"], "from -> to": _co["item_type"], "clean-out": _co["process"].astype(str).str.split("|").str[0].str.strip(),
            "start": _co["start_time"], "end": _co["end_time"], "duration_sec": (_co["duration_min"] * 60).round(0),
            "line": _co["line"]}).sort_values(["machine", "start"])
    for name, df in _xs.items():
        ws = wb.create_sheet(name)
        write_df(ws, df, fmts={c: DT for c in df.columns if str(c) in ("first_start", "last_end", "cleanout_start", "start", "end",
                                                                        "next_start", "campaign_start", "campaign_end",
                                                                        "first_draw", "last_draw", "supply_usable")}
                 | ({"plant_day": D} if "plant_day" in df.columns else {}))
        ws.freeze_panes = "A2"
        print(f"  {name:<20} {len(df):>6,} rows  (extra sheet)")
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
