#!/usr/bin/env python3
"""make_validation_result_simple.py - a run's plan validation in the layout of the
original outputs/validation_result.xlsx (Summary, By day, Blocked SKUs, Blocking
items, Blocking detail, All combinations) + one short "Schedule checks" sheet.

    python make_validation_result_simple.py --run e2e_v45b [--out <file.xlsx>]

Follows the run's own data-blocker policy (phase0_sku_status.csv):
  BLOCKED                  missing data with no safe stand-in (e.g. no BOM) - not scheduled,
                           left out of the final figures
  SCHEDULABLE (ASSUMED)    a gap filled by a stated assumption (A01 machine pool, A02
                           department-median process time, ...) - scheduled and counted
  SCHEDULABLE              all data present
A combination = SKU x machine group of the building plan. Days are plant days (07:00).
"""
from __future__ import annotations

import argparse
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ctpio as io

YEL = PatternFill("solid", fgColor="FFFF00")
ORA = PatternFill("solid", fgColor="FCE4D6")
GRN = PatternFill("solid", fgColor="E2EFDA")
HDR = PatternFill("solid", fgColor="D9E1F2")
REASON = {"no_bom": "no BOM for this SKU+group (CTP_Master_BOM)",
          "no_building_routing": "no GT BUILDING routing (CTP_Master_Routing)",
          "no_routing": "component has no routing (CTP_Master_Routing)",
          "no_machines": "no machines in CTP_Master_Routing -> A01 department pool (assumed)",
          "no_proc_time": "no proc_time in CTP_Master_Routing -> A02 department median (assumed)"}


def write(ws, df, fill=None):
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HDR
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else
                   (v.to_pydatetime() if isinstance(v, pd.Timestamp) else v) for v in r])
        if fill:
            f = fill(dict(zip(df.columns, r)))
            if f:
                for c in ws[ws.max_row]:
                    c.fill = f
    for j, col in enumerate(df.columns, start=1):
        w = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, w + 2), 70)
    ws.freeze_panes = "A2"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    D = io.OUT_ROOT / a.run
    out = a.out or (D / f"validation_result_{a.run}.xlsx")

    combos, detail, days = [], [], []
    for ln in ("pcr", "tbr"):
        O = D / ln
        plan = pd.read_csv(O / "phase0_plan_full.csv", parse_dates=["build_start"])
        plan = plan[plan["qty"] > 0]
        st = pd.read_csv(O / "phase0_sku_status.csv", dtype=str).fillna("")
        st = st[st["source"] == "building_plan"].drop_duplicates("sku").set_index("sku")
        ep = pd.read_csv(O / "phase0_effective_plan.csv").set_index("sku")
        fnd = pd.read_csv(O / "phase0_findings.csv") if (O / "phase0_findings.csv").exists() else pd.DataFrame(
            columns=["sku", "item", "issue", "detail", "tyres"])
        itypes = io.load_itemtypes(io.load_bom(ln))
        plan["mg"] = plan.get("plan_mg", "").fillna("") if "plan_mg" in plan else ""
        for (sku, mg), g in plan.groupby(["sku", "mg"], dropna=False):
            state = st["state"].get(sku, "CLEAN")
            f = fnd[fnd["sku"] == sku]
            if state == "DATA_BLOCKED":
                status = "BLOCKED"
            elif state == "ASSUMED":
                status = "SCHEDULABLE (ASSUMED " + (st["assumption_codes"].get(sku, "") or "") + ")"
            else:
                status = "SCHEDULABLE"
            items = [i for i in f["item"].dropna().astype(str).unique() if i]
            ex = "; ".join((f"{r.item} ({REASON.get(r.issue, r.issue)})"
                            if isinstance(r.item, str) and r.item else REASON.get(r.issue, r.issue))
                           for r in f.drop_duplicates(["item", "issue"]).head(3).itertuples())
            combos.append({"line": ln.upper(), "sku": sku, "machine_group": mg,
                           "plan_qty": int(g["qty"].sum()), "plan_rows": len(g), "status": status,
                           "components": (int(ep["components"].get(sku)) if "components" in ep
                                          and pd.notna(ep["components"].get(sku)) else None),
                           "blocking_items": (max(1, len(items)) if status == "BLOCKED" else len(items)),
                           "examples": ex})
            for r in f.drop_duplicates(["item", "issue"]).itertuples():
                detail.append({"line": ln.upper(), "sku": sku, "group": mg, "qty": int(g["qty"].sum()),
                               "item": r.item if isinstance(r.item, str) else "",
                               "reason": REASON.get(r.issue, r.issue),
                               "item_type": itypes.get(r.item, "") if isinstance(r.item, str) else "",
                               "effect": "BLOCKS the SKU" if status == "BLOCKED" else "assumed - SKU scheduled"})
            g = g.assign(day=g["build_start"].map(lambda t: pd.Timestamp(io.plant_day(t)).normalize()),
                         kept=g["qty"] if status != "BLOCKED" else 0)
            for dday, gd in g.groupby("day"):
                days.append({"line": ln.upper(), "plant_day": dday, "planned": gd["qty"].sum(),
                             "effective": gd["kept"].sum()})

    C = pd.DataFrame(combos).sort_values(["status", "plan_qty"], ascending=[True, False])
    DT = pd.DataFrame(detail)
    BD = pd.DataFrame(days).groupby(["line", "plant_day"]).sum().reset_index()
    both = BD.groupby("plant_day")[["planned", "effective"]].sum().reset_index().assign(line="BOTH")
    BD = pd.concat([both[BD.columns], BD], ignore_index=True)
    BD["dropped"] = BD["planned"] - BD["effective"]
    BD["kept_pct"] = (100 * BD["effective"] / BD["planned"]).round(1)
    BD = BD[["line", "plant_day", "planned", "effective", "dropped", "kept_pct"]]
    blk = C[C["status"] == "BLOCKED"]
    BI = (DT[DT["effect"] == "BLOCKS the SKU"].assign(item=lambda x: x["item"].replace("", "(whole SKU)"))
          .groupby(["item", "reason"]).agg(skus_blocked=("sku", "nunique"), tyres=("qty", "sum"))
          .reset_index().sort_values("tyres", ascending=False)) if len(DT) else pd.DataFrame()

    # curing-side figures and the schedule checks from VALIDATION_REPORT
    try:
        vs = pd.read_csv(D / "VALIDATION_SUMMARY.csv", encoding="utf-8-sig")
        tot = vs[vs["line"] == "TOTAL"].iloc[0]
    except Exception:
        tot = None
    try:
        chk = pd.read_excel(D / "VALIDATION_REPORT.xlsx", sheet_name="Schedule checks")
        chk = chk[["gate", "line", "check", "value", "verdict"]]
    except Exception:
        chk = pd.DataFrame()

    asm = C[C["status"].str.startswith("SCHEDULABLE (ASSUMED")]
    S = [("run", a.run),
         ("planned combinations", len(C)),
         ("planned tyres", int(C["plan_qty"].sum())),
         ("schedulable combinations", int((C["status"] != "BLOCKED").sum())),
         ("  of which on a stated assumption (A01/A02)", len(asm)),
         ("EFFECTIVE PLANNED QUANTITY", int(C.loc[C["status"] != "BLOCKED", "plan_qty"].sum())),
         ("blocked combinations", len(blk)),
         ("blocked tyres", int(blk["plan_qty"].sum())),
         ("kept %", round(100 * C.loc[C["status"] != "BLOCKED", "plan_qty"].sum() / C["plan_qty"].sum(), 1)),
         ("plan rows kept", int(C.loc[C["status"] != "BLOCKED", "plan_rows"].sum())),
         ("plan rows total", int(C["plan_rows"].sum()))]
    for ln in ("PCR", "TBR"):
        c = C[C["line"] == ln]
        S.append((f"{ln}: planned / effective tyres",
                  f"{int(c['plan_qty'].sum()):,} / {int(c.loc[c['status'] != 'BLOCKED', 'plan_qty'].sum()):,}"))
    if tot is not None:
        S += [("curing plan tyres (raw)", int(tot["plan_tyres_raw"])),
              ("curing plan tyres excluded (data-blocked)", int(tot["data_blocked_tyres"])),
              ("curing plan tyres eligible", int(tot["eligible_tyres"])),
              ("cured in horizon (eligible)", int(tot["cured_eligible"])),
              ("fulfilment % of eligible", float(tot["fulfilment_pct_eligible"]))]
    if len(chk):
        S.append(("schedule checks", f"{int((chk['verdict'] == 'PASS').sum())} PASS, "
                                     f"{int((chk['verdict'] == 'FAIL').sum())} FAIL, "
                                     f"{int((~chk['verdict'].isin(['PASS', 'FAIL'])).sum())} not measured"))
    S += [("bom table", "JK_CTP_PPDS.CTP_Master_BOM"), ("routing table", "JK_CTP_PPDS.CTP_Master_Routing"),
          ("plan", "October 2026 building + curing plan (plan-latest)"),
          ("WIP snapshot", "JK_CTP_PPDS.CTP_Opening_WIP 22 Oct shift A (06:30)")]
    S = pd.DataFrame(S, columns=["metric", "value"])

    wb = Workbook()
    write(wb.active, S)
    wb.active.title = "Summary"
    write(wb.create_sheet("By day"), BD, fill=lambda r: ORA if r["kept_pct"] < 95 else None)
    write(wb.create_sheet("Blocked SKUs"), blk.drop(columns=[]), fill=lambda r: YEL)
    write(wb.create_sheet("Blocking items"), BI if len(BI) else pd.DataFrame({"(none)": []}), fill=lambda r: YEL)
    write(wb.create_sheet("Blocking detail"), DT if len(DT) else pd.DataFrame({"(none)": []}),
          fill=lambda r: YEL if r.get("effect") == "BLOCKS the SKU" else ORA)
    write(wb.create_sheet("All combinations"), C,
          fill=lambda r: YEL if r["status"] == "BLOCKED" else (ORA if "ASSUMED" in r["status"] else GRN))
    if len(chk):
        write(wb.create_sheet("Schedule checks"), chk,
              fill=lambda r: GRN if r["verdict"] == "PASS" else (YEL if r["verdict"] == "FAIL" else None))
    wb.save(out)
    print(S.to_string(index=False))
    print(f"[OK] {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
