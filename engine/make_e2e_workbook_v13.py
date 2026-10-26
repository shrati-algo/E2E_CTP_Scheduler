#!/usr/bin/env python3
"""make_e2e_workbook_v13.py — the e2e schedule and its sums check, in one book.

    python make_e2e_workbook_v13.py --run e2e_v13

ONE WORKBOOK PER LINE, TWO SHEETS:

  1  E2E Schedule   every row of phase5_floor_schedule_{line}.csv in the plant's
                    own 16 columns - the v10 format, plus `incomplete_components`
                    which names the kit parts that were not ready together.

  2  Sums Check     the same quantities pivoted department / machine x plant day,
                    exactly as `phase5_floor_schedule_tbr.xlsx` lays it out, with
                    the routing validation and the machine register above it so
                    the totals can be trusted rather than merely read.

The two sheets are reconciled before writing: the pivot's grand total must equal
the schedule's `produce_qty` sum, or the book is not written. A sums sheet that
does not add up to its own schedule is worse than no sums sheet.
"""
from __future__ import annotations

import argparse
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import ctpio as io

NAVY = PatternFill("solid", fgColor="1F3864")
SUB = PatternFill("solid", fgColor="D9E2F3")
TOT = PatternFill("solid", fgColor="FFF2CC")
BAD = PatternFill("solid", fgColor="FFC7CE")
GOOD = PatternFill("solid", fgColor="C6EFCE")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)
WHITE_B = Font(bold=True, color="FFFFFF")


def _pivot(d: pd.DataFrame) -> pd.DataFrame:
    d = d.copy()
    d["day"] = pd.to_datetime(d["start_time"], errors="coerce").map(
        lambda t: io.plant_day(t).date() if pd.notna(t) else None)
    d["produce_qty"] = pd.to_numeric(d["produce_qty"], errors="coerce").fillna(0.0)
    return pd.pivot_table(d, index=["department", "machine"], columns="day",
                          values="produce_qty", aggfunc="sum",
                          fill_value=0.0).sort_index()


def sheet_schedule(wb: Workbook, d: pd.DataFrame) -> None:
    ws = wb.create_sheet("E2E Schedule")
    cols = list(d.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, str(c))
        cell.font = WHITE_B
        cell.fill = NAVY
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    dt = {"start_time", "end_time"}
    for i, (_, row) in enumerate(d.iterrows(), 2):
        for j, c in enumerate(cols, 1):
            v = row[c]
            if pd.isna(v):
                cell = ws.cell(i, j, None)
            elif c in dt:
                cell = ws.cell(i, j, pd.to_datetime(v).to_pydatetime())
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif c == "produce_qty":
                cell = ws.cell(i, j, float(v))
                cell.number_format = "#,##0.00"
            else:
                cell = ws.cell(i, j, str(v))
            cell.border = THIN
        k = str(row.get("incomplete_components", ""))
        if k.startswith(("LATE", "INCOMPLETE")):
            ws.cell(i, cols.index("incomplete_components") + 1).fill = BAD
        elif k == "COMPLETE":
            ws.cell(i, cols.index("incomplete_components") + 1).fill = GOOD
    widths = {"date": 12, "machine": 12, "shift": 7, "item": 22, "item_type": 20,
              "process": 20, "department": 20, "start_time": 20, "end_time": 20,
              "produce_qty": 13, "UOM": 7, "lot_id": 22, "prod_id": 11,
              "consumption_id": 40, "FG_SKU_CODE": 20,
              "incomplete_components": 34}
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(d)+1}"


TARGET = {"pcr": 13500, "tbr": 3500}


def _daywise(run: str, line: str) -> pd.DataFrame:
    """GT planned / built and curing planned / cured, per plant day.

    Bucketed on FEASIBLE_END, never on build_start: bucketing the achieved
    figure on the planned day just restates the plan and reads 100% every day.
    """
    D = io.OUT_ROOT / run / line
    b = pd.read_csv(D / "phase5_building_plan.csv",
                    parse_dates=["build_start", "feasible_end"])
    c = pd.read_csv(D / "phase7_curing_plan.csv",
                    parse_dates=["planned_start", "feasible_end"])
    bw, cw = b[b["within_horizon"]], c[c["within_horizon"]]
    t = pd.concat([
        b.groupby(b["build_start"].dt.date)["qty"].sum().rename("GT planned"),
        bw.groupby(bw["feasible_end"].dt.date)["qty"].sum().rename("GT built"),
        c.groupby(c["planned_start"].dt.date)["qty_planned"].sum().rename("cure planned"),
        cw.groupby(cw["feasible_end"].dt.date)["qty"].sum().rename("cured"),
    ], axis=1).fillna(0.0).sort_index()
    t["target"] = TARGET[line]
    t["gap vs target"] = t["cured"] - TARGET[line]
    # `.replace(0, pd.NA)` on a float column returns OBJECT dtype in pandas 2.x,
    # and Series.round() then refuses it ("Expected numeric dtype, got object").
    # `.where()` keeps the column float and puts NaN in the divide-by-zero slots,
    # which is what was meant.
    _cp = t["cure planned"].astype(float)
    t["cured %"] = (100 * t["cured"].astype(float) /
                    _cp.where(_cp > 0)).round(1)
    return t.reset_index().rename(columns={"index": "day"})


def _blockers(run: str, line: str) -> pd.DataFrame:
    """What stopped a trolley, what it cost, and what stock it had."""
    D = io.OUT_ROOT / run / line
    b = pd.read_csv(D / "phase5_building_plan.csv")
    s = pd.read_csv(D / "phase5_schedule.csv")
    o = pd.read_csv(io.DATA / "opening_wip.csv", encoding="utf-8-sig",
                    low_memory=False)
    st = o.groupby(o["itemcode"].astype(str).str.strip())["inventory"].sum()
    ty = dict(zip(s["item_code"], s["item_type"]))
    lost = b[~b["within_horizon"]]
    if not len(lost):
        return pd.DataFrame(columns=["blocking_item", "item_type", "lots",
                                     "tyres lost", "opening stock"])
    g = (lost.groupby("blocking_item")
         .agg(lots=("qty", "size"), tyres=("qty", "sum"))
         .sort_values("tyres", ascending=False).reset_index())
    g["item_type"] = [ty.get(i, "") for i in g["blocking_item"]]
    g["opening stock"] = [round(float(st.get(str(i), 0.0)), 1)
                          for i in g["blocking_item"]]
    return g.rename(columns={"tyres": "tyres lost"})[
        ["blocking_item", "item_type", "lots", "tyres lost", "opening stock"]]


def sheet_sums(wb: Workbook, d: pd.DataFrame, val: pd.DataFrame,
               cov: pd.DataFrame, line: str, run: str) -> float:
    ws = wb.create_sheet("Sums Check")
    r = 1
    ws.cell(r, 1, f"{line.upper()} — SUMS CHECK").font = Font(bold=True, size=13)
    r += 2

    # ---- day-wise, against the plant target -------------------------------
    dw = _daywise(run, line)
    ws.cell(r, 1, f"DAY-WISE  (target {TARGET[line]:,} cured/day)").font = Font(bold=True)
    r += 1
    cols = list(dw.columns)
    for j, h in enumerate(cols, 1):
        c = ws.cell(r, j, str(h))
        c.font = WHITE_B
        c.fill = NAVY
        c.alignment = Alignment(horizontal="center")
    r += 1
    for _, row in dw.iterrows():
        for j, h in enumerate(cols, 1):
            v = row[h]
            cell = ws.cell(r, j, None if pd.isna(v) else
                           (str(v) if h == "day" else float(v)))
            cell.border = THIN
            if h != "day":
                cell.number_format = "#,##0" if h != "cured %" else "0.0"
        # colour the gap: red short of target, green at or above
        gj = cols.index("gap vs target") + 1
        ws.cell(r, gj).fill = GOOD if row["gap vs target"] >= 0 else BAD
        r += 1
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    for h in ("GT planned", "GT built", "cure planned", "cured"):
        c = ws.cell(r, cols.index(h) + 1, float(dw[h].sum()))
        c.font = Font(bold=True)
        c.fill = TOT
        c.number_format = "#,##0"
    r += 2

    # ---- major blockers ---------------------------------------------------
    bl = _blockers(run, line)
    ws.cell(r, 1, "MAJOR BLOCKERS — what stopped a trolley").font = Font(bold=True)
    r += 1
    bcols = list(bl.columns)
    for j, h in enumerate(bcols, 1):
        c = ws.cell(r, j, str(h))
        c.font = WHITE_B
        c.fill = NAVY
    r += 1
    for _, row in bl.head(12).iterrows():
        for j, h in enumerate(bcols, 1):
            v = row[h]
            cell = ws.cell(r, j, str(v) if isinstance(v, str) else float(v))
            cell.border = THIN
            if h in ("tyres lost", "opening stock", "lots"):
                cell.number_format = "#,##0.0" if h == "opening stock" else "#,##0"
        # near-zero opening stock is the story: mark it
        if float(row["opening stock"]) < 10:
            ws.cell(r, bcols.index("opening stock") + 1).fill = BAD
        r += 1
    ws.cell(r, 1, f"TOTAL lost at building").font = Font(bold=True)
    c = ws.cell(r, bcols.index("tyres lost") + 1, float(bl["tyres lost"].sum()))
    c.font = Font(bold=True)
    c.fill = TOT
    c.number_format = "#,##0"
    r += 2

    # ---- validation block -------------------------------------------------
    ws.cell(r, 1, "ROUTING VALIDATION").font = Font(bold=True)
    r += 1
    for j, h in enumerate(["test", "verdict", "failures", "detail"], 1):
        c = ws.cell(r, j, h)
        c.font = WHITE_B
        c.fill = NAVY
    r += 1
    for _, v in val.iterrows():
        ws.cell(r, 1, str(v["test"]))
        c = ws.cell(r, 2, str(v["verdict"]))
        c.fill = BAD if v["verdict"] == "FAIL" else GOOD
        c.font = Font(bold=True)
        ws.cell(r, 3, int(v["failures"]))
        ws.cell(r, 4, str(v["detail"]))
        r += 1
    r += 1

    # ---- idle machines ----------------------------------------------------
    idle = cov[~cov["planned"].astype(str).str.lower().isin(["true"])]
    ws.cell(r, 1, "MACHINES WITH NO WORK").font = Font(bold=True)
    r += 1
    if len(idle):
        for _, m in idle.iterrows():
            ws.cell(r, 1, str(m["machine"]))
            ws.cell(r, 2, str(m["name"]))
            ws.cell(r, 4, str(m.get("issue", "")))
            ws.cell(r, 1).fill = BAD
            r += 1
    else:
        ws.cell(r, 1, "none — every register machine carries work").fill = GOOD
        r += 1
    r += 1

    # ---- the pivot --------------------------------------------------------
    p = _pivot(d)
    days = list(p.columns)
    ws.cell(r, 1, "Sum of produce_qty").font = Font(bold=True, size=12)
    r += 1
    head = r
    ws.cell(r, 1, "Row Labels").font = WHITE_B
    ws.cell(r, 1).fill = NAVY
    for j, day in enumerate(days, 2):
        c = ws.cell(r, j, str(day))
        c.font = WHITE_B
        c.fill = NAVY
        c.alignment = Alignment(horizontal="center")
    c = ws.cell(r, len(days) + 2, "Grand Total")
    c.font = WHITE_B
    c.fill = NAVY
    r += 1

    for dept, grp in p.groupby(level=0):
        tot = grp.sum()
        ws.cell(r, 1, dept).font = Font(bold=True)
        ws.cell(r, 1).fill = SUB
        for j, day in enumerate(days, 2):
            c = ws.cell(r, j, float(tot[day]) or None)
            c.fill = SUB
            c.number_format = "#,##0.00"
        c = ws.cell(r, len(days) + 2, float(tot.sum()))
        c.fill = SUB
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
        r += 1
        for (_, mach), row in grp.iterrows():
            ws.cell(r, 1, f"    {mach}")
            for j, day in enumerate(days, 2):
                c = ws.cell(r, j, float(row[day]) or None)
                c.number_format = "#,##0.00"
            c = ws.cell(r, len(days) + 2, float(row.sum()))
            c.number_format = "#,##0.00"
            r += 1
    ws.cell(r, 1, "Grand Total").font = Font(bold=True)
    ws.cell(r, 1).fill = TOT
    for j, day in enumerate(days, 2):
        c = ws.cell(r, j, float(p[day].sum()))
        c.fill = TOT
        c.font = Font(bold=True)
        c.number_format = "#,##0.00"
    c = ws.cell(r, len(days) + 2, float(p.values.sum()))
    c.fill = TOT
    c.font = Font(bold=True)
    c.number_format = "#,##0.00"

    ws.column_dimensions["A"].width = 34
    for j in range(2, max(len(days) + 3, 9)):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.column_dimensions["D"].width = 60
    ws.freeze_panes = ws.cell(head + 1, 2).coordinate
    return float(p.values.sum())


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", default="e2e_v13")
    a = ap.parse_args()
    D = io.OUT_ROOT / a.run
    val = pd.read_csv(D / "VALIDATION.csv", encoding="utf-8-sig")
    cov = pd.read_csv(D / "MACHINE_COVERAGE.csv", encoding="utf-8-sig")

    print("\n" + "=" * 80)
    print(f"  E2E WORKBOOK — schedule + sums check, one per line   [{a.run}]")
    print("=" * 80)
    rc = 0
    for line in ("pcr", "tbr"):
        f = D / f"phase5_floor_schedule_{line}.csv"
        if not f.exists():
            print(f"  {line}: no floor schedule")
            continue
        d = pd.read_csv(f, encoding="utf-8-sig")
        wb = Workbook()
        wb.remove(wb.active)
        sheet_schedule(wb, d)
        total = sheet_sums(wb, d,
                           val[val["line"] == line.upper()],
                           cov[cov["line"].isin([line.upper(), "BOTH"])],
                           line, a.run)

        # RECONCILE before writing: the two sheets must agree.
        raw = float(pd.to_numeric(d["produce_qty"], errors="coerce").fillna(0).sum())
        if abs(raw - total) > 0.5:
            print(f"  [FAIL] {line.upper()}: pivot {total:,.2f} != schedule "
                  f"{raw:,.2f} — not written")
            rc = 1
            continue
        out = D / f"CTP_E2E_SCHEDULE_{line}.xlsx"
        wb.save(out)
        nf = int((val[val["line"] == line.upper()]["verdict"] == "FAIL").sum())
        print(f"  {line.upper()}: {out.name}")
        print(f"     sheet 1  E2E Schedule : {len(d):,} rows x {len(d.columns)} cols")
        print(f"     sheet 2  Sums Check   : totals {total:,.2f} == schedule "
              f"{raw:,.2f}  [reconciled]")
        print(f"              validation   : {len(val[val['line']==line.upper()])} "
              f"checks, {nf} FAIL")
    return rc


if __name__ == "__main__":
    sys.exit(main())
