#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""rule_adherence_report.py -- one rule-adherence test report for a schedule version (2026-10-06).

    python rule_adherence_report.py v34          -> latest_schedule/v34/RULE_ADHERENCE_REPORT_v34.xlsx

Collects every check workbook of the version (building / curing, mixing, MPQ, 4RC, extruders, general + IL rules)
into one table: section, rule, test, occurrences, followed, adherence %, status, source tables. Status:
PASS >= 99.5 %, PARTIAL 80-99.5 %, FAIL < 80 %; report-only rows and rules without plant data are listed but not
scored. The Summary sheet scores each section with formulas over the Checks sheet (recalculates in Excel).
"""
import glob
import re
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
T = {"ROUT": "JK_CTP_PPDS.CTP_Master_Routing", "MX": "JK_CTP_PPDS.CTP_Master_ChangeoverMatrix",
     "BOM": "JK_CTP_PPDS.CTP_Master_BOM", "MPQ": "JK_CTP_PPDS.CTP_Master_MPQ", "FRC": "JK_CTP_PPDS.CTP_Master_FRC",
     "AGE": "JK_CTP_PPDS.CTP_Master_Aging", "WIP": "JK_CTP_PPDS.CTP_Opening_WIP", "BA": "JK_CTP_PPDS.CTP_Master_BuildingAllowable",
     "DT": "jkplanning_CTP.downtime_reason", "PLAN": "plan files (latest_plan_file)"}


def _src(*k):
    return "; ".join(T[x] for x in k)


FILES = [  # (file prefix, section, default source tables)
    ("BUILD_CURE_VALIDATION", "Building & curing", _src("ROUT", "MX", "PLAN")),
    ("MIXING_RULE_COMPLIANCE", "Mixing", _src("ROUT", "BOM")),
    ("MIXING_RULE_AUDIT", "Mixing (rules 01-Oct + v47)", _src("ROUT", "BOM", "FRC")),
    ("MPQ_COMPLIANCE", "MPQ", _src("MPQ", "BOM", "ROUT", "BA")),
    ("FOURRC_RULES", "4-roll calender (901)", _src("ROUT", "MX", "FRC")),
    ("QUINTOPLEX_RULES", "Extrusion - Quintoplex (1002)", _src("ROUT", "MX", "BOM")),
    ("DUPLEX_RULES", "Extrusion - Duplex (1801)", _src("ROUT", "MX", "BOM")),
    ("TRIPLEX_RULES", "Extrusion - Triplex (1001)", _src("ROUT", "MX", "BOM")),
    ("QUADRAPLEX_RULES", "Extrusion - Quadraplex (1803)", _src("ROUT", "MX", "BOM")),
    ("GENERAL_RULES", "General rules", ""),
]
GEN_SRC = {"BR-1": _src("PLAN", "ROUT"), "BR-2": _src("WIP", "BOM"), "BR-3": _src("PLAN", "WIP"), "BR-4": _src("ROUT"),
           "BR-5": _src("ROUT"), "BR-7": _src("MX"), "BR-8": _src("BOM", "MX"), "BR-9": _src("ROUT"),
           "BR-10": _src("DT"), "BR-12": _src("AGE", "ROUT"), "BR-14": _src("MX"), "BR-15": _src("WIP", "DT", "PLAN"),
           "BR-16": _src("MX", "BOM"), "BR-17": "all JK_CTP_PPDS masters (live)", "BR-37": _src("BOM", "WIP", "PLAN"),
           "BR-38": _src("ROUT"), "BR-39": _src("ROUT"), "BR-41": _src("BOM", "MX"), "CD-3": _src("BOM", "MX", "PLAN")}


def _status(p, note):
    t = str(note or "").upper()
    if "NO DATA" in t:
        return "NO DATA"
    if p is None or p != p:
        return "REPORT"
    return "PASS" if p >= 99.5 else ("PARTIAL" if p >= 80 else "FAIL")


def collect(v):
    vdir = os.path.join(HERE, "latest_schedule", v)
    rows = []
    for pre, sec, src in FILES:
        fs = [f for f in glob.glob(os.path.join(vdir, f"{pre}_*.xlsx"))]
        if not fs:
            continue
        f = fs[0]
        if pre == "MIXING_RULE_COMPLIANCE":
            d = pd.read_excel(f, "By rule")
            for r in d.itertuples():
                rows.append((sec, r.rule, r.occurrence, r.occurrences, r.followed, r[7], str(r.note) if r.note == r.note else "", src))
            try:
                for r in pd.read_excel(f, "Not in the %").itertuples():
                    rows.append((sec, r.rule, r.description, 0, 0, None, f"NOT SCORED: {r[3]}", src))
            except Exception:
                pass
            continue
        if pre == "MIXING_RULE_AUDIT":
            for r in pd.read_excel(f, "Summary").to_dict("records"):
                n = r.get("occurrences checked")
                vv = r.get("violations")
                p = r.get("compliance_%")
                tgt = str(r.get("status", "")).startswith("TARGET")
                rows.append((sec, r["rule"], r.get("rule text (as evaluated)"), n,
                             (None if tgt or vv != vv or vv is None else (n or 0) - vv),
                             (None if tgt or not n else p),
                             (str(r.get("status")) + f" -- {p}% of target" if tgt else ""), src))
            continue
        if pre == "MPQ_COMPLIANCE":
            d = pd.read_excel(f, "Summary").set_index("check")["value"]
            n = int(d.get("upstream lots with an MPQ (DB master)", 0))
            b = int(d.get("lots below MPQ (must be 0)", 0))
            rows.append((sec, "MPQ-DB", "upstream lot >= its CTP_Master_MPQ minimum", n, n - b, 100 * (n - b) / n if n else None,
                         f"{int(d.get('lots with no MPQ row in CTP_Master_MPQ', 0))} lots without an MPQ row; "
                         f"{int(d.get('items on the smallest gauge row (gauge not in the DB)', 0))} edge items on the smallest gauge", src))
            continue
        d = pd.read_excel(f, "Summary")
        rc = "check" if "check" in d.columns else "occurrence"
        for r in d.to_dict("records"):
            rule = str(r.get("rule", ""))
            note = str(r.get("note", "")) if r.get("note") == r.get("note") else ""
            p = r.get("compliance_%")
            if "report" in str(r.get(rc, "")).lower():
                p = None
            if "switch is same die" in str(r.get(rc, "")):     # strict view; the rule itself allows "where coverage permits"
                p, note = None, ("strict view (scored row: '... or coverage did not permit'). " + note).strip()
            s, sc = sec, src
            if pre == "GENERAL_RULES":
                key = rule.split()[0]
                s = ("IL / calendering" if key in ("BR-37", "BR-38", "BR-39", "BR-41")
                     else "Dispatch (cover first, CD)" if key.startswith("CD") else "General rules")
                sc = GEN_SRC.get(key, "")
                if rule in ("BR-7",) and "without a matrix" in str(r.get(rc, "")):
                    note = "NO DATA: " + note
                if rule == "BR-10" and "shear" in str(r.get(rc, "")):
                    note = "NO DATA: " + note
            rows.append((s, rule, r.get(rc), r.get("occurrences"), r.get("followed"), p, note, sc))
    D = pd.DataFrame(rows, columns=["section", "rule", "test", "occurrences", "followed", "adherence_%", "note", "source_tables"])
    m = re.match(r"v(\d+)", v)                       # also variants such as v75_relaxed
    if m and int(m.group(1)) >= 36:                  # v36+: a rule switched off in config is measured, not scored
        import make_version_readme as _mvr
        off = _mvr._off_rules() - {"BR-3", "BR-16"}

        def _is_off(rule, test):
            k = str(rule).split()[0]
            if k not in off:
                return False
            if k == "BR-1":
                return "plan shift" in str(test)
            if k == "BR-12":
                return "fresh" in str(test)
            return True
        _o = [_is_off(r, t) for r, t in zip(D["rule"], D["test"])]
        D.loc[_o, "note"] = "NOT SCORED: switched off in this version (config) -- measured only. " + D.loc[_o, "note"].astype(str)
    D["status"] = [("NOT SCORED" if str(n).startswith("NOT SCORED") else _status(p, n)) for p, n in zip(D["adherence_%"], D["note"])]
    D.loc[D["status"].isin(["REPORT", "NO DATA", "NOT SCORED"]), "adherence_%"] = None
    return D


def write(v, D, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ck = wb.create_sheet("Checks")
    f_h = Font(name="Arial", bold=True, color="FFFFFF")
    fill_h = PatternFill("solid", fgColor="1F3864")
    f_n = Font(name="Arial", size=10)
    thin = Side(style="thin", color="BFBFBF")
    fills = {"PASS": "C6EFCE", "PARTIAL": "FFEB9C", "FAIL": "FFC7CE", "NO DATA": "D9D9D9", "REPORT": "DDEBF7", "NOT SCORED": "EDEDED"}
    cols = ["section", "rule", "test", "occurrences", "followed", "adherence_%", "status", "note", "source_tables"]
    for j, c in enumerate(cols, 1):
        x = ck.cell(1, j, c.replace("_", " "))
        x.font, x.fill, x.alignment = f_h, fill_h, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(D[cols].itertuples(index=False), 2):
        for j, (c, val) in enumerate(zip(cols, r), 1):
            if val is None or (isinstance(val, float) and val != val):
                val = None
            x = ck.cell(i, j, val)
            x.font = f_n
            x.border = Border(bottom=thin)
            if c == "adherence_%" and val is not None:
                x.number_format = "0.00"
            if c == "status":
                x.fill = PatternFill("solid", fgColor=fills.get(val, "FFFFFF"))
                x.alignment = Alignment(horizontal="center")
            if c in ("test", "note", "source_tables"):
                x.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate([24, 22, 52, 12, 11, 12, 12, 48, 46], 1):
        ck.column_dimensions[get_column_letter(j)].width = w
    ck.freeze_panes = "A2"
    ck.auto_filter.ref = ck.dimensions
    n = len(D) + 1
    # Summary with formulas over Checks
    ws["A1"] = f"Rule adherence test report -- schedule {v}"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = ("Adherence % per check = followed / occurrences. PASS >= 99.5 %, PARTIAL 80-99.5 %, FAIL < 80 %. "
                "Report-only rows, rules without plant data and not-applicable rules are listed on 'Checks' but not scored.")
    ws["A2"].font = Font(name="Arial", italic=True, size=9)
    hdr = ["Section", "Checks scored", "PASS", "PARTIAL", "FAIL", "No data / report / not scored",
           "Average adherence % (per check)", "Occurrence-weighted adherence %"]
    for j, h in enumerate(hdr, 1):
        x = ws.cell(4, j, h)
        x.font, x.fill, x.alignment = f_h, fill_h, Alignment(horizontal="center", vertical="center", wrap_text=True)
    secs = list(dict.fromkeys(D["section"]))
    rng = lambda col: f"Checks!${col}$2:${col}${n}"
    for i, s in enumerate(secs, 5):
        ws.cell(i, 1, s).font = f_n
        ws.cell(i, 2, f'=COUNTIFS({rng("A")},$A{i},{rng("F")},"<>")')
        for j, st in ((3, "PASS"), (4, "PARTIAL"), (5, "FAIL")):
            ws.cell(i, j, f'=COUNTIFS({rng("A")},$A{i},{rng("G")},"{st}")')
        ws.cell(i, 6, f'=COUNTIFS({rng("A")},$A{i})-B{i}')
        ws.cell(i, 7, f'=IFERROR(AVERAGEIFS({rng("F")},{rng("A")},$A{i}),"")')
        ws.cell(i, 8, f'=IFERROR(SUMPRODUCT(({rng("A")}=$A{i})*({rng("F")}<>"")*{rng("E")})/'
                      f'SUMPRODUCT(({rng("A")}=$A{i})*({rng("F")}<>"")*{rng("D")})*100,"")')
        for j in range(1, 9):
            ws.cell(i, j).border = Border(bottom=thin)
            if j >= 7:
                ws.cell(i, j).number_format = "0.00"
    t = 5 + len(secs)
    ws.cell(t, 1, "ALL RULES").font = Font(name="Arial", bold=True)
    for j, col in ((2, "B"), (3, "C"), (4, "D"), (5, "E"), (6, "F")):
        ws.cell(t, j, f"=SUM({col}5:{col}{t - 1})").font = Font(name="Arial", bold=True)
    ws.cell(t, 7, f'=AVERAGE({rng("F")})').number_format = "0.00"
    ws.cell(t, 8, f'=SUMPRODUCT(({rng("F")}<>"")*{rng("E")})/SUMPRODUCT(({rng("F")}<>"")*{rng("D")})*100').number_format = "0.00"
    for j in (7, 8):
        ws.cell(t, j).font = Font(name="Arial", bold=True)
    ws.cell(t + 2, 1, "Source data: live JK_CTP_PPDS / jkplanning_CTP tables of the run (see README_" + v + ".md); "
                      "check workbooks *_" + v + ".xlsx in the same folder.").font = Font(name="Arial", italic=True, size=9)
    for j, w in enumerate([34, 12, 9, 10, 9, 18, 18, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 45
    wb.save(out)


if __name__ == "__main__":
    for v in sys.argv[1:]:
        D = collect(v)
        out = os.path.join(HERE, "latest_schedule", v, f"RULE_ADHERENCE_REPORT_{v}.xlsx")
        write(v, D, out)
        sc = D[D["adherence_%"].notna()]
        print(f"{v}: {len(D)} rows, {len(sc)} scored checks | PASS {int((D.status == 'PASS').sum())} PARTIAL "
              f"{int((D.status == 'PARTIAL').sum())} FAIL {int((D.status == 'FAIL').sum())} | avg {sc['adherence_%'].mean():.2f}% -> {out}")
