"""Back-calculation of the no-cutter-slitter scheduler (v1): TBM kit -> calenders / extruders -> mixing, from the DB
BOM and routing AFTER the cutter / slitter bypass, every value compared with the SMART BOM.

    python backcalc/back_calculation.py [vNN]      -> latest_schedule/vNN/BACK_CALCULATION_vNN.xlsx

Reads (read-only): ctp_inputfiles/bom_live.csv and routing_live.csv (written by build_sep_inputs with cutter_bypass),
ctp_inputfiles/cutter_bypass_links.csv / _transfer.csv / _stock_ratios.csv, inputs/building_plan.csv (TBM plan),
inputs/cutter_bypass_stock_converted.csv, config.yaml, the SMART BOM (backcalc/bom_engine.py).
"""
import math
import os
import re
import sys

import pandas as pd
import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from bom_engine import BASE_GROUPS, Smart, load_roll_widths, load_smart, norm  # noqa: E402

VER = sys.argv[1] if len(sys.argv) > 1 else "v1"
CI = os.path.join(ROOT, "ctp_inputfiles")
OUTDIR = os.path.join(ROOT, "latest_schedule", VER)
os.makedirs(OUTDIR, exist_ok=True)
OUT = os.path.join(OUTDIR, f"BACK_CALCULATION_{VER}.xlsx")
TOL_OK, TOL_SMALL = 0.03, 0.10
RECIPE = {"MASTER COMPOUND", "SMALL CHEMICAL", "NATURAL RUBBER", "SYNETHIC RUBBER", "CARBON BLACK", "PROCESS OIL",
          "RUBBER CHEMICALS"}
RAW_DESC = RECIPE - {"MASTER COMPOUND"} | {"STEEL TYRE CORD", "NYLON FABRIC", "BEAD WIRE", "MISC.", "ZINCE OXIDE",
                                          "TUBE VALVES &ACCE."}
# DB code -> SMART code where the two name the same material differently
ALIAS = {"CAP66MOTHERROLL": "CAP66", "CAP30MOTHERROLL": "CAP30", "GBCAP20MOTHERROLL": "GBCAP20"}


def verdict(r):
    if r is None or (isinstance(r, float) and (math.isnan(r) or math.isinf(r))):
        return "NOT CHECKABLE"
    if abs(r - 1) <= TOL_OK:
        return "OK"
    if abs(r - 1) <= TOL_SMALL:
        return "SMALL DIFFERENCE (3-10 %)"
    return "DB HIGHER" if r > 1 else "DB LOWER"


# ------------------------------------------------------------------------------------------------ inputs
cfg = yaml.safe_load(open(os.path.join(ROOT, "config.yaml"), encoding="utf-8")) or {}
cb = cfg.get("cutter_bypass") or {}
bom = pd.read_csv(os.path.join(CI, "bom_live.csv"), low_memory=False)
for c in ("Super_parent", "grand_parent", "Parent", "child", "child_Unit", "Parent_unit", "child_description"):
    bom[c] = bom[c].astype(str).str.strip()
bom["child_Unit"] = bom.child_Unit.str.upper()
rt = pd.read_csv(os.path.join(CI, "routing_live.csv"), low_memory=False)
rt["routed_product"] = rt.routed_product.astype(str).str.strip()
links = pd.read_csv(os.path.join(CI, "cutter_bypass_links.csv"))
trans = pd.read_csv(os.path.join(CI, "cutter_bypass_transfer.csv")).set_index("source_item").bypass_transfer_min
ratios = pd.read_csv(os.path.join(CI, "cutter_bypass_stock_ratios.csv"))
bp = pd.read_csv(os.path.join(ROOT, "inputs", "building_plan.csv"))
tyres = bp.groupby(bp.skuCode.astype(str).str.strip()).qty.sum()
stock_conv = os.path.join(ROOT, "inputs", "cutter_bypass_stock_converted.csv")
stock_conv = pd.read_csv(stock_conv) if os.path.exists(stock_conv) else pd.DataFrame()
raw_rt = pd.read_csv(os.path.join(ROOT, "db_pipeline", "raw", "JK_CTP_PPDS__CTP_Master_Routing.csv"), low_memory=False)
raw_rt["routed_product"] = raw_rt.routed_product.astype(str).str.strip()

S = Smart(load_smart(), load_roll_widths())
cut_codes = set(raw_rt[raw_rt.operation_name.astype(str).str.strip().isin(
    [o for s in (cb.get("sections") or {}).values() for o in s.get("operations", [])])].routed_product)
cut_k = {norm(c) for c in cut_codes}
BASE_CODES = set(S.x[S.x.GRP.isin(BASE_GROUPS)].MaterialCodeGroup_IN)
gt_of = bom[bom.child_description.str.upper() == "GREEN TYRES"].drop_duplicates("Super_parent") \
    .set_index("Super_parent").child.to_dict()

# ------------------------------------------------------------------------------------------------ routing rate
# the scheduler's FINAL routing (inputs/routing.csv: DB routing + line-speed file + roll multipliers + overrides +
# the cutting / slitting transfer time) -> the same rates the schedule used
fr = pd.read_csv(os.path.join(ROOT, "inputs", "routing.csv"), low_memory=False)
fr["routed_product"] = fr.routed_product.astype(str).str.strip()
rt1 = fr.sort_values("is_primary", ascending=False).drop_duplicates("routed_product").set_index("routed_product")


def hours(item, qty, unit):
    """Machine hours for qty of item at the scheduler's routing rate / efficiency (no changeovers)."""
    if item not in rt1.index or not qty:
        return None, "", ""
    r = rt1.loc[item]
    pt, u = pd.to_numeric(r.proc_time, errors="coerce"), str(r.proc_time_UOM).strip().upper()
    eff = pd.to_numeric(r.efficiency, errors="coerce")
    eff = eff if eff and eff == eff else 1.0
    mc = str(r.machines)
    if not pt or pt != pt:
        return None, mc, u
    b = pd.to_numeric(r.batch_size, errors="coerce")
    if u in ("M/MIN", "NOS/MIN", "KG/MIN", "CUTS/MIN"):
        h = qty / pt / 60.0
    elif u in ("MIN/BATCH", "SEC/BATCH") and b and b == b:
        h = qty / b * pt / (60.0 if u == "MIN/BATCH" else 3600.0)
    elif u == "SEC":
        h = qty * pt / 3600.0
    else:
        return None, mc, u
    return h / eff, mc, f"{pt:g} {u}" + (f", batch {b:g} kg" if b == b and b else "") + f", eff {eff:g}"


# ------------------------------------------------------------------------------------------------ SMART tyre walk
def smart_tyre(gt):
    """Every item in one SMART tyre -> (qty per tyre, unit): walks GT / carcass rows through cut parts, calendered rolls,
    extrudates down to final compounds, cord, fabric, wire. Weight is kept through cutting (roll m = kg / roll kg/m)."""
    from bom_validation_bridge import smart_components, smart_comp_qty_kg
    acc = {}
    cur = {}

    def add(code, q, u, into=None):
        tgt = cur if into is None else into
        k = ALIAS.get(norm(code), norm(code))
        v = tgt.get(k, [0.0, u, code])
        v[0] += q
        tgt[k] = v

    def merge(src, f=1.0):
        for k, (q, u, code) in src.items():
            v = acc.get(k, [0.0, u, code])
            v[0] += q * (f if u == "KG" else 1.0)      # only weights follow the GT-row weight; metres / pieces do not
            acc[k] = v

    def walk(item, qty, unit, depth=0):
        if depth > 12 or not qty:
            return
        if unit == "KG":                       # an intermediate held in kg (apex profile): split by weight share
            rr = [r for r in S.rows(item) if S.is_base(r) and r.GRP in BASE_GROUPS]
            tot = sum(S.kg_link(item, r, "M")[0] for r in rr)
            for r in rr:
                if tot > 0:
                    add(r.MaterialCodeGroup_IN, qty * S.kg_link(item, r, "M")[0] / tot, "KG")
            return
        for r in S.rows(item):
            kl, _ = S.kg_link(item, r, unit)
            I = r.MaterialCodeGroup_IN
            if S.is_base(r):
                if r.GRP in BASE_GROUPS:
                    add(I, qty * kl, "KG")
                elif r.UNIT == "KG":          # apex profile (AP 0018) without its own BOM: its compound, unnamed
                    add("UNRESOLVED|" + item, qty * kl, "KG")
                continue
            if r.UNIT == "NOS":
                q2, u2 = qty * (r.Peace or 1), "NOS"
            elif r.UNIT == "KG":
                q2, u2 = qty * kl, "KG"
            else:
                k_in = S.kgpm(I)
                if not k_in:
                    continue
                q2, u2 = qty * kl / k_in, "M"
            add(I, q2, u2)
            walk(I, q2, u2, depth + 1)

    for k, rs in smart_components(S, gt).items():
        sq, su, skg = smart_comp_qty_kg(rs)
        code = rs[0].MaterialCodeGroup_IN
        cur.clear()
        if S.eg.get(norm(code)) == "4RC":        # the GT row names the 4RC roll but means the strip cut from it
            kg = skg or sq * sum(x.WeightPerMeter for x in rs)
            k_in = S.kgpm(code)
            if k_in:
                add(code, kg / k_in, "M")
                walk(code, kg / k_in, "M")
            merge(cur)
            continue
        walk(code, sq, su if su in ("M", "NOS") else "NOS")
        # SMART's own weight of this part in this tyre (GT / carcass row) rules where its part BOM disagrees
        # (sidewall / BEF001: 'Peace 2' but the weight of one) -> everything below the part is scaled to it
        deep = sum(q for kk, (q, u, c) in cur.items() if u == "KG" and (c in BASE_CODES or c.startswith("UNRESOLVED|")))
        f = (skg / deep) if (skg and deep and abs(deep / skg - 1) > TOL_OK) else 1.0
        merge(cur, f)
        add(code, sq, su, into=acc)              # the part itself keeps SMART's quantity (Length x Peace)
    return acc


# ------------------------------------------------------------------------------------------------ DB per tyre
routed = set(rt.routed_product)
desc = bom.drop_duplicates("child").set_index("child").child_description.to_dict()
compound = set(bom[bom.child_description.str.upper().isin({"FINAL COMPOUND", "MASTER COMPOUND"})].child) | \
    set(bom[(bom.Parent_unit.str.upper() == "KG") & bom.child_description.str.upper().isin(RECIPE)].Parent)
# DB per tyre of every item = sum of all its rows (phase 1b rule); compounds: only the kg used directly by a part
# (masters under finals follow the recipes, which agree with SMART -- bom_validation 'Recipes')
use = bom[~bom.child_description.str.upper().isin(RAW_DESC) | bom.child.isin(compound)]
use = use[~(use.child.isin(compound) & use.Parent.isin(compound))]
use = use[use.child_description.str.upper() != "GREEN TYRES"]
dbq = use.groupby(["Super_parent", "child", "child_Unit"]).child_quantity.sum().reset_index()

# where each DB item sits: TBM kit (direct child of the green tyre) or a feed of another item
lvl = bom.assign(top=bom.Parent == bom.Super_parent.map(gt_of)).groupby(["Super_parent", "child"]).top.max()
path = links.groupby(["Super_parent", "child"])["_byp_path"].agg(lambda s: " | ".join(sorted(set(map(str, s)))))
consumer = bom.groupby(["Super_parent", "child"]).Parent.agg(lambda s: ", ".join(sorted(set(s)))[:120])

# known DB errors (bom_validation/README.md) -> suggestion text
DUP = {v[0]: v[1] for v in S.dup_info.values()} if hasattr(S, "dup_info") else {}


def suggestion(sku, gt, item, d_, r, unit):
    v = verdict(r)
    if v == "OK":
        return ""
    s = []
    it = norm(item)
    if norm(gt) in (norm("GT 1513 XPC1 MSIL"),) or norm(gt).replace("GT", "CAR", 1) in (norm("CAR 1402 XPC TATA"),):
        s.append("SKU BOM has SMART's repeated block added up (validation #1): rebuild from one copy")
    if d_ in ("FINAL COMPOUND",) and r and r > 1.15 and str(item).startswith(("T", "PCRT")):
        s.append("tread: first compound carries the whole tread weight (validation #2): use WeightPerMeter x length")
    if it in (norm("MONOCHAFER"), "S36"):
        s.append("chafer chain MONOCHAF80 -> BI-MONO001 (NOS) -> MONOCHAFER 0.0317 m/m; geometry 0.264 (validation #3)")
    if it in ("BRTTL", "B333"):
        s.append("BCFT00x used as n pieces with a one-piece BOM (validation #4): BRTTL / B333 under-stated")
    if it in ("BPG01", "B278"):
        s.append("BPGSL01 / BPG01 quantities swapped (validation #5b): B278 ~7.5x")
    if str(d_).upper() == "CALANDARED ROLL" and r and abs(r - 0.5) <= 0.05 and it not in ("BRTTL", norm("MONOCHAFER")):
        s.append("ply / belt used as 2 pieces with a one-piece BOM (validation #8): DB half -- write the piece BOM "
                 "for one piece or the quantity in metres")
    if it in ("EG01", "EG02", "B127", "B1144"):
        s.append("edge strip / wrap counted 2x (validation #5): plant to confirm 1 or 2 strips per edge")
    if unit == "M" and r and abs(r - 0.5) <= 0.03 and (str(d_).upper() in ("SIDEWALL", "BELT EDGE FILLER")
                                                       or it == "BEF001"):
        s.append("pair basis: DB 1 m = left + right pair, SMART 2 pieces; kg the same (validation #6) -- no change")
    if it in ("B606", "B107") and "SP003" in str(consumer.get((sku, item), "")):
        s.append("SP003 kg x2 in the DB; already halved by config master_fixes (validation #7)")
    if not s:
        s.append(f"DB {r:.2f}x SMART: check the DB rows for this item in this SKU (bom_validation 'Link check detail')"
                 if r else "item not in the SMART tyre: check naming / coverage")
    return "; ".join(s)


rows = []
for sp, g in dbq.groupby("Super_parent"):
    if sp not in tyres.index:
        continue
    gt = gt_of.get(sp, "")
    sm = smart_tyre(gt)
    n = float(tyres[sp])
    seen = set()
    for kk in [k for k in sm if k.startswith(norm("UNRESOLVED|"))]:
        q_un, _, code_un = sm.pop(kk)
        part = code_un.split("|", 1)[1]
        cands = [c for c in bom[(bom.Super_parent == sp) & (bom.Parent == part)].child if c in compound and
                 not sm.get(norm(c), [0])[0]]
        if len(cands) == 1:
            k2 = norm(cands[0])
            v = sm.get(k2, [0.0, "KG", cands[0]])
            v[0] += q_un
            sm[k2] = v
    for r in g.itertuples():
        k = ALIAS.get(norm(r.child), norm(r.child))
        seen.add(k)
        sq = sm.get(k)
        smart_q, smart_u = (sq[0], sq[1]) if sq else (None, "")
        if sq and smart_u != r.child_Unit:
            # NOS <-> M (cut parts / treads): compare in metres with the SMART piece length when possible
            pl = S.piece_len_m(sq[2])
            if smart_u == "NOS" and r.child_Unit == "M" and pl:
                smart_q, smart_u = smart_q * pl, "M"
        ratio = (r.child_quantity / smart_q) if (smart_q and smart_u == r.child_Unit) else None
        d_ = desc.get(r.child, "")
        h, mc, rate = hours(r.child, r.child_quantity * n, r.child_Unit)
        hs = hours(r.child, (smart_q or 0) * n, r.child_Unit)[0] if ratio is not None else None
        rows.append({"SKU": sp, "Green tyre": gt, "TBM tyres (building plan)": n, "Item": r.child, "Item type": d_,
                     "Level": "TBM kit (taken by the tyre builder)" if lvl.get((sp, r.child), False) else
                     ("compound" if r.child in compound else "feeds " + str(consumer.get((sp, r.child), ""))),
                     "Reached through bypassed cutter / slitter items": path.get((sp, r.child), ""),
                     "Cutting / slitting transfer added (min)": trans.get(r.child, 0.0) if r.child in trans.index else 0.0,
                     "Total transfer time in routing (min)": pd.to_numeric(rt1.transfer_time_min.get(r.child),
                                                                           errors="coerce") if r.child in rt1.index else None,
                     "Unit": r.child_Unit, "DB qty / tyre": r.child_quantity, "SMART qty / tyre": smart_q,
                     "SMART unit": smart_u, "DB / SMART": ratio, "Result": verdict(ratio) if sq else "NOT IN SMART TYRE",
                     "DB qty in window": r.child_quantity * n,
                     "SMART qty in window": (smart_q or 0) * n if ratio is not None else None,
                     "Machine(s) (DB routing)": mc, "Rate (DB routing)": rate,
                     "Machine hours in window (DB)": h, "Machine hours in window (SMART qty)": hs,
                     "Suggestion": suggestion(sp, gt, r.child, d_, ratio, r.child_Unit) if sq else
                     "not found in the SMART tyre (naming / coverage)"})
    for k, (q, u, code) in sm.items():
        if k in seen or k in cut_k or u == "KG" and code not in compound and not str(code).startswith(("B", "T", "S")):
            continue
        if any(x in str(code) for x in ("NOT APPLICABLE",)):
            continue
        rows.append({"SKU": sp, "Green tyre": gt, "TBM tyres (building plan)": n, "Item": code, "Item type": "",
                     "Level": "only in SMART", "Unit": u, "SMART qty / tyre": q, "SMART unit": u,
                     "Result": "ONLY IN SMART", "SMART qty in window": q * n,
                     "Suggestion": "SMART has it, the DB tyre does not (or under another code)"})
per_sku = pd.DataFrame(rows)

# ------------------------------------------------------------------------------------------------ aggregates
ok = per_sku[per_sku["DB qty in window"].notna()]
items = ok.groupby(["Item", "Unit"]).agg(
    item_type=("Item type", "first"), level=("Level", lambda s: s.mode().iloc[0] if len(s) else ""),
    skus=("SKU", "nunique"), db=("DB qty in window", "sum"), smart=("SMART qty in window", "sum"),
    machines=("Machine(s) (DB routing)", "first"), rate=("Rate (DB routing)", "first"),
    h_db=("Machine hours in window (DB)", "sum"), h_sm=("Machine hours in window (SMART qty)", "sum"),
    transfer=("Cutting / slitting transfer added (min)", "max"),
    through=("Reached through bypassed cutter / slitter items",
             lambda s: " | ".join(sorted({p.split(" > ")[0] for x in s if x for p in str(x).split(" | ")}))[:200]),
    sugg=("Suggestion", lambda s: s[s != ""].mode().iloc[0] if (s != "").any() else "")).reset_index()
cmp_ = ok[ok["SMART qty in window"].notna()].groupby(["Item", "Unit"])[["DB qty in window", "SMART qty in window"]].sum()
items = items.merge(cmp_.rename(columns={"DB qty in window": "db_c", "SMART qty in window": "sm_c"}),
                    left_on=["Item", "Unit"], right_index=True, how="left")
items["DB / SMART (where both exist)"] = items.db_c / items.sm_c
items["Result"] = [verdict(x) for x in items["DB / SMART (where both exist)"]]
items = items.rename(columns={"item_type": "Item type", "level": "Level", "skus": "Planned SKUs", "db": "DB qty in window",
                              "smart": "SMART qty in window", "machines": "Machine(s)", "rate": "Rate (DB routing)",
                              "h_db": "Machine hours (DB)", "h_sm": "Machine hours (SMART qty)",
                              "transfer": "Cutting / slitting transfer (min)",
                              "through": "Bypassed cut items it now replaces", "sugg": "Suggestion (most common)"}) \
    .drop(columns=["db_c", "sm_c"])
items = items.sort_values(["Level", "Item type", "DB qty in window"], ascending=[True, True, False])

bytype = ok.groupby(["Item type", "Unit"]).apply(lambda g: pd.Series({
    "Items": g.Item.nunique(), "DB qty in window": g["DB qty in window"].sum(),
    "SMART qty in window (comparable rows)": g["SMART qty in window"].sum(),
    "DB qty in window (comparable rows)": g.loc[g["SMART qty in window"].notna(), "DB qty in window"].sum(),
    "Rows OK": int((g.Result == "OK").sum()), "Rows differing": int(g.Result.isin(["DB HIGHER", "DB LOWER"]).sum()),
    "Machine hours (DB)": g["Machine hours in window (DB)"].sum()})).reset_index()
bytype["DB / SMART"] = bytype["DB qty in window (comparable rows)"] / bytype["SMART qty in window (comparable rows)"]
bytype["Result"] = [verdict(x) for x in bytype["DB / SMART"]]

# machine load: hours per machine group (DB routing rate), with and without SMART quantities
mach = items.groupby("Machine(s)").agg(items=("Item", "nunique"), h_db=("Machine hours (DB)", "sum"),
                                       h_sm=("Machine hours (SMART qty)", "sum")).reset_index()
mach["hours per day (DB)"] = mach.h_db / 10.0
mach = mach.rename(columns={"items": "Items", "h_db": "Hours in window (DB qty)", "h_sm": "Hours in window (SMART qty)"})

# bypass map
sec = {}
for name, s in (cb.get("sections") or {}).items():
    for op in s.get("operations", []):
        sec[op] = (name, s.get("transfer_min"))
bm = raw_rt[raw_rt.routed_product.isin(cut_codes)].drop_duplicates("routed_product")[
    ["routed_product", "operation_name", "machines"]].copy()
bm["Section"] = bm.operation_name.map(lambda o: sec.get(str(o).strip(), ("", ""))[0])
bm["Transfer min (this step)"] = bm.operation_name.map(lambda o: sec.get(str(o).strip(), ("", ""))[1])
src = ratios.groupby("cut_item").apply(lambda g: "; ".join(f"{a} {b:.4g} {u}/m" for a, b, u in
                                                          zip(g.source, g.per_metre, g.source_unit)))
bm["Now taken directly (source per metre of the cut item)"] = bm.routed_product.map(src)
bm = bm.rename(columns={"routed_product": "Cut / slit item (bypassed)", "operation_name": "Operation",
                        "machines": "Machine(s)"}).sort_values(["Section", "Cut / slit item (bypassed)"])
tr = trans.rename("Total cutting / slitting transfer added (min)").reset_index().rename(
    columns={"source_item": "Source item"})
tr["Item type"] = tr["Source item"].map(desc)

# ------------------------------------------------------------------------------------------------ workbook
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

FILL = {"OK": "C6EFCE", "SMALL": "FFEB9C", "BAD": "FFC7CE", "GREY": "EDEDED"}


def sheet(wb, name, df, note=None):
    ws = wb.create_sheet(name[:31])
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True)
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(6, min(len(df.columns), 12)))
        ws.row_dimensions[1].height = 48
        r0 = 2
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(r0, j, c)
        x.font = Font(bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", fgColor="1F4E78")
        x.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[r0].height = 42
    rc = [j for j, c in enumerate(df.columns, 1) if c == "Result"]
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float):
                v = None if v != v else (round(v, 0) if abs(v) >= 1000 else round(v, 4))
            ws.cell(i, j, v)
        for j in rc:
            s = str(ws.cell(i, j).value).upper()
            f = FILL["OK"] if s.startswith("OK") else FILL["SMALL"] if s.startswith("SMALL") else \
                FILL["BAD"] if ("HIGHER" in s or "LOWER" in s or "ONLY" in s) else FILL["GREY"]
            ws.cell(i, j).fill = PatternFill("solid", fgColor=f)
    for j, c in enumerate(df.columns, 1):
        w = max([len(str(c)) * 0.7] + [len(str(v)) for v in df[c].head(150)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(10, w + 2), 60)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    if len(df):
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(df.columns))}{r0 + len(df)}"


wb = Workbook()
ws = wb.active
ws.title = "Read me"
secs_txt = [f"  - {n}: operations {', '.join(s.get('operations', []))} -> {s.get('transfer_min')} min"
            for n, s in (cb.get("sections") or {}).items()]
lines = [
    (f"BACK CALCULATION {VER}: TBM full kit -> calenders / extruders -> mixing, cutters and slitters bypassed", True),
    ("", False),
    ("WHAT CHANGED vs v133", True),
    ("Every item made on a cutter or slitter (belt cutters 801 / 802, ply cutters 1103 / 1104, TBR body ply 2101, bias "
     "cutter 2001, chafer slitters 2401 / 2402, cap ply slitters 701 / 702, cap ply cutter / slitter 601) is taken out "
     "of the BOM and the routing. Its consumer -- the tyre builder, or the inner-liner calender / bead apexer that "
     "used the strip -- now takes the SOURCE material directly: the 4RC calendered roll, the edge gum, the "
     "mother roll. The per-tyre quantity is exactly what the DB BOM states for that source (DB rows are absolute "
     "per-tyre quantities, so nothing is re-scaled; sheet 'Bypass map').", False),
    ("Back-calculation chain: TBM plan (inputs/building_plan.csv, tyres per SKU) x DB qty per tyre of each kit item "
     "(4RC roll metres, extrudate metres, bead apex pieces ...) -> each item's own DB BOM gives the compound kg -> "
     "the DB routing rate gives machine hours. All numbers are in 'Per SKU back calc' and summed in 'Items (window)'.",
     False),
    ("CUTTING / SLITTING TRANSFER TIME (configurable: config.yaml cutter_bypass.sections.<section>.transfer_min)", True),
    ("The cutting step is not scheduled, but its time is kept as a lag: the source item's routing transfer time is "
     "increased by the minutes of every bypassed step on its way to the consumer (e.g. MONOCHAFER: bias cutter 90 + "
     "chafer slitter 60 = 150 min). The source must therefore be ready that much earlier than the builder needs it. "
     "ASSUMED values (typical tyre-plant time to cut / slit a batch, let it relax, handle and move it to the TBM):",
     False),
] + [(t, False) for t in secs_txt] + [
    ("", False),
    ("SMART COMPARISON", True),
    ("SMART per tyre: the SMART green-tyre / carcass rows (Length x Peace) walked down through cut parts, rolls and "
     "extrudates to compounds, keeping the weight through every cut (roll m = kg / roll kg per m; roll widths from "
     "Master Data 4RC.xlsx). Repeated SMART blocks counted once. OK = within 3 %, SMALL = 3-10 %, else DB HIGHER / "
     "LOWER. 'Suggestion' links each difference to the DB BOM error found in bom_validation (README there).", False),
    ("v1 uses the DB BOM AS IT IS (no corrections), as asked; the differences are reported, not fixed.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    c.font = Font(bold=b)
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 150

sheet(wb, "Bypass map", bm, "Each cut / slit item taken out, the operation and machine it was made on, the transfer "
                            "minutes of that step, and the source(s) the consumer now takes directly (per metre of the "
                            "cut item, median over the tyres that use it -- also used to convert its opening stock).")
sheet(wb, "Transfer time added", tr.sort_values("Total cutting / slitting transfer added (min)"),
      "Minutes added to each source item's routing transfer time (sum of the bypassed steps on the longest chain).")
sheet(wb, "By item type", bytype, "DB vs SMART per item type over the plan window (TBM tyres x qty per tyre).")
sheet(wb, "Items (window)", items, "Every item the TBM kit back-calculates to: DB quantity in the 10-day window, "
                                   "SMART quantity, machine hours from the DB routing rate.")
sheet(wb, "Machine load", mach.sort_values("Hours in window (DB qty)", ascending=False),
      "Machine hours in the window at the scheduler's routing rate (inputs/routing.csv incl. roll multipliers and "
      "efficiency; no changeovers) -- a capacity check, not the schedule. 'per day' = window hours / 10.")
ps = per_sku.sort_values(["TBM tyres (building plan)", "SKU", "Level"], ascending=[False, True, True])
sheet(wb, "Per SKU back calc", ps, "One row per planned SKU and item: DB qty per tyre, SMART qty per tyre, window "
                                   "quantities, machine hours and the suggested correction.")
if len(stock_conv):
    sheet(wb, "Stock converted", stock_conv, "Opening stock of cut / slit items booked as the source items it holds.")
wb.save(OUT)
print("written", OUT, "|", len(per_sku), "rows;", per_sku.Result.value_counts().to_dict())
